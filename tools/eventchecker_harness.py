#!/usr/bin/env python3
"""Lightweight harness for catching repeat regressions in EventInspector.

Run this before pushing changes:

    python3 tools/eventchecker_harness.py

The harness intentionally focuses on the brittle parts of the app:
- update manifest shape
- exact-match parsing contracts
- installation-id state transitions
- package-code mapping
- release payload/source sync for every build target

It does not start the UI or connect to devices.
"""

from __future__ import annotations

import argparse
import copy
import functools
import http.server
import io
import importlib.util
import json
import os
import re
import hashlib
import plistlib
import subprocess
import sys
import tempfile
import threading
import time
import types
import zipfile
from dataclasses import dataclass
from pathlib import Path, PosixPath
from typing import Callable, Iterable, List
from urllib.parse import urlparse

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
CURRENT_RELEASE_VERSION = "2026-10-27-1-2.5.0-83"
CURRENT_RELEASE_BUILD = 83
ROLLBACK_SOURCE_BUILD = 56
# Build marker of the source being prepared.  It runs ahead of
# CURRENT_RELEASE_BUILD (the published manifest payload) until the release gate
# regenerates the manifest from the committed source.
RELEASE_SOURCE_BUILD = 84
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import Log_checker as lc  # noqa: E402
# desktop_app only needs webview after main() starts; keep the harness usable
# in a plain Python environment that has not installed the UI dependency.
sys.modules.setdefault("webview", types.ModuleType("webview"))
import desktop_app as desktop  # noqa: E402


@dataclass
class HarnessResult:
    name: str
    passed: bool
    message: str = ""


def _reset_runtime_state() -> None:
    lc.active_platform = "android"
    lc.connected_devices_info = []
    lc.installation_id_state.clear()
    lc.sdk_check_runtime_state.clear()
    lc.sdk_check_expected_map.clear()
    lc.sdk_check_expected_order.clear()
    lc.sdk_check_search_list.clear()
    lc.sdk_check_current_network.clear()
    lc.is_paused = False
    lc.target_package_name = ""
    lc.package_log_cache.clear()
    lc.active_ios_package_log_buffers.clear()
    lc.active_ios_package_log_frame_buffers.clear()
    lc.active_ios_package_log_stream_modes.clear()
    lc.incomplete_android_event_logs.clear()
    lc.ios_device_missing_polls.clear()
    lc._clear_record_dedup()


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _assert_equal(actual, expected, message: str) -> None:
    if actual != expected:
        raise AssertionError(f"{message}: expected={expected!r} actual={actual!r}")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _payload_path_candidates(manifest_path: Path, item: dict) -> List[Path]:
    rel_path = str(item.get("source_path") or item.get("path", "")).strip()
    if not rel_path:
        return []
    payload_dir = manifest_path.parent
    return [ROOT / rel_path, payload_dir / rel_path]


def _immutable_manifest_payloads(manifest: dict) -> dict[str, bytes]:
    """Load the exact Git blobs named by the manifest payload_ref.

    Update-flow tests must model what the client downloads.  Reading the
    mutable worktree here can mix uncommitted source changes with an older
    manifest hash and create a false repeated-download failure.
    """
    payload_ref = str(manifest.get("payload_ref") or "").strip()
    _assert(payload_ref, "manifest payload_ref is required for update-flow fixtures")
    payloads: dict[str, bytes] = {}
    for item in manifest.get("files") or []:
        rel_path = str(item.get("path") or "").strip()
        if not rel_path:
            continue
        try:
            payloads[rel_path] = subprocess.check_output(
                ["git", "show", f"{payload_ref}:{rel_path}"],
                cwd=str(ROOT),
            )
        except (OSError, subprocess.CalledProcessError) as exc:
            raise AssertionError(
                f"manifest payload blob is unavailable for {rel_path} at {payload_ref}: {exc}"
            ) from exc
    return payloads


def _valid_payload_urls(manifest_path: Path, rel_path: str, payload_path: Path) -> set[str]:
    repo_rel = payload_path.relative_to(ROOT).as_posix()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    payload_ref = str(manifest.get("payload_ref") or "").strip()
    # The repository moved to indiezvn; trucbm URLs remain valid (redirected).
    return {
        url
        for owner in ("indiezvn", "trucbm")
        for url in (
            f"https://github.com/{owner}/Eventchecker/raw/{payload_ref}/{repo_rel}",
            f"https://raw.githubusercontent.com/{owner}/Eventchecker/{payload_ref}/{repo_rel}",
            f"https://cdn.jsdelivr.net/gh/{owner}/Eventchecker@{payload_ref}/{repo_rel}",
        )
    }


def _manifest_paths() -> List[Path]:
    candidates = [ROOT / "Updates_2_5" / "remote_manifest.json"]
    return [path for path in candidates if path.exists()]


def test_manifest_contract() -> None:
    manifests = _manifest_paths()
    _assert(manifests, "No release manifest files found")
    for manifest_path in manifests:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
        _assert("version" in data and str(data["version"]).strip(), f"{manifest_path.name} missing version")
        payload_ref = str(data.get("payload_ref") or "").strip().lower()
        _assert(re.fullmatch(r"[0-9a-f]{40}", payload_ref) is not None, f"{manifest_path.name} payload_ref must be an immutable Git SHA")
        files = data.get("files") or []
        _assert(files, f"{manifest_path.name} missing files list")
        log_checker_files = [item for item in files if str(item.get("path", "")).strip() == "Log_checker.py"]
        _assert(log_checker_files, f"{manifest_path.name} missing Log_checker.py payload")
        for item in files:
            _assert(str(item.get("path", "")).strip(), f"{manifest_path.name} contains file entry without path")
            _assert(str(item.get("url", "")).strip(), f"{manifest_path.name} contains file entry without url")
            _assert(str(item.get("sha256", "")).strip(), f"{manifest_path.name} contains file entry without sha256")
            for url in [item.get("url"), *(item.get("urls") or [])]:
                _assert(payload_ref in str(url), f"{manifest_path.name} payload URL is not pinned: {url}")
                _assert("/main/" not in str(url) and "@main/" not in str(url), f"{manifest_path.name} payload URL is mutable: {url}")


def test_manifest_payload_integrity() -> None:
    manifests = _manifest_paths()
    _assert(manifests, "No release manifest files found")
    for manifest_path in manifests:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
        immutable_payloads = _immutable_manifest_payloads(data)
        for item in data.get("files") or []:
            rel_path = str(item.get("path", "")).strip()
            if not rel_path:
                continue
            candidates = _payload_path_candidates(manifest_path, item)
            payload_path = next((candidate for candidate in candidates if candidate.exists()), None)
            _assert(payload_path is not None, f"{manifest_path.name} payload missing for {rel_path}")

            expected_sha = str(item.get("sha256", "")).strip().lower()
            payload_bytes = immutable_payloads.get(rel_path)
            _assert(payload_bytes is not None, f"{manifest_path.name} Git payload missing for {rel_path}")
            actual_sha = hashlib.sha256(payload_bytes).hexdigest().lower()
            _assert_equal(actual_sha, expected_sha, f"{manifest_path.name} sha mismatch for {rel_path} at payload_ref")

            valid_urls = _valid_payload_urls(manifest_path, rel_path, payload_path)
            url = str(item.get("url", "")).strip()
            _assert(
                url in valid_urls,
                f"{manifest_path.name} url points outside {manifest_path.parent.name} for {rel_path}: {url}",
            )
            for extra_url in item.get("urls") or []:
                extra_url = str(extra_url).strip()
                _assert(
                    extra_url in valid_urls,
                    f"{manifest_path.name} urls entry points outside {manifest_path.parent.name} for {rel_path}: {extra_url}",
                )


def test_canonical_update_channel_contract() -> None:
    """Keep checks on the single v250 channel with immutable payloads."""
    import remote_update as updater

    _assert_equal(updater.CHANNEL_ID, "v250", "current updater channel changed")
    _assert_equal(tuple(updater.KNOWN_CHANNELS), ("v250",), "current updater must have one channel")
    _assert_equal(updater.CONFIG_FILENAME, "remote_update_config_v250.json", "current updater config changed")
    _assert_equal(updater.STATE_FILENAME, "update_state_v250.json", "current updater state changed")
    _assert_equal(updater.UPDATES_DIRNAME, "updates_v250", "current updater payload directory changed")

    config_path = ROOT / "remote_update_config_v250.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    _assert(
        config.get("manifest_url") in updater.DEFAULT_MANIFEST_URLS,
        "current manifest URL must use the canonical main/v250 URL set",
    )
    _assert_equal(
        set(config.get("manifest_urls") or []),
        set(updater.DEFAULT_MANIFEST_URLS),
        "current manifest URL fallbacks changed",
    )
    _assert_equal(config.get("min_interval_sec"), 0, "current updater must check without a stale throttle")

    manifest_path = ROOT / "Updates_2_5" / "remote_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    payload_ref = str(manifest.get("payload_ref") or "").strip().lower()
    _assert(re.fullmatch(r"[0-9a-f]{40}", payload_ref) is not None, "current payload ref must be an immutable Git SHA")
    for item in manifest.get("files") or []:
        urls = [item.get("url"), *(item.get("urls") or [])]
        _assert(
            all(payload_ref in str(url) and "/main/" not in str(url) and "@main/" not in str(url) for url in urls if url),
            f"current release payload must use immutable ref {payload_ref}: {item.get('path')}",
        )


def test_package_code_mapping() -> None:
    expected = {
        "com.indiez.nonogram": "NG",
        "com.indiez.train.miner": "TM",
        "com.indiez.idletycoon.horse.racing": "HR",
        "com.indiez.solitaire.word.card.puzzle": "SW",
        "com.nostel.dot.line.puzzle": "KN",
        "com.nostel.parking.car": "CP",
        "com.afk.idle.cat.food.restaurent": "CR",
        "tap.monster.block.away": "TP",
    }
    for package_id, game_code in expected.items():
        _assert_equal(lc._game_code_for_package(package_id), game_code, f"wrong game code for {package_id}")
    _assert_equal(lc._game_code_for_package("com.example.unknown"), "Unknown", "unknown package should stay Unknown")


def test_package_log_presets_contract() -> None:
    """Package Log Quick Select must come from the editable JSON catalog."""
    preset_path = ROOT / "package_log_presets.json"
    raw = json.loads(preset_path.read_text(encoding="utf-8"))
    presets = lc._sanitize_package_log_presets(raw)
    _assert_equal(len(presets), 8, "package preset catalog entry count changed")
    _assert(
        any(item.get("android", {}).get("id") == "com.indiez.nonogram" for item in presets),
        "package preset catalog lost the NG Android entry",
    )
    _assert(
        any(item.get("ios", {}).get("id") == "CatRestaurant" for item in presets),
        "package preset catalog lost the CatRestaurant iOS entry",
    )

    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8")
    for needle in (
        "package_log_presets.json",
        "PACKAGE_LOG_PRESETS_REMOTE_URLS",
        "@app.get('/api/package-log-presets')",
        "packageQuickSelectList",
        "reloadPackagePresetsBtn?.addEventListener('click', () => loadPackageLogPresets(true))",
    ):
        _assert(needle in source, f"Package Log dynamic catalog contract missing: {needle}")

    for build_script in (
        ROOT / "build" / "macos" / "build_macos.sh",
        ROOT / "build" / "windows" / "build_portable.bat",
        ROOT / "build" / "windows" / "build_windows.bat",
    ):
        _assert(
            "package_log_presets.json" not in build_script.read_text(encoding="utf-8", errors="ignore"),
            f"{build_script.name} must not bundle package_log_presets.json",
        )


def test_installation_id_state_machine() -> None:
    _reset_runtime_state()
    device_id = "device-1"

    lc._update_installation_id_runtime(device_id, package_id="com.indiez.nonogram", installation_id="id-old", force_emit=True)
    _assert_equal(lc.installation_id_state[device_id]["installation_id"], "id-old", "initial installation id not stored")

    lc._update_installation_id_runtime(device_id, package_id="com.indiez.nonogram", installation_id="id-old-2")
    _assert_equal(lc.installation_id_state[device_id]["installation_id"], "id-old-2", "same package should accept the newest installation id")

    lc._update_installation_id_runtime(device_id, package_id="com.nostel.parking.car", installation_id="", force_emit=True)
    _assert_equal(lc.installation_id_state[device_id]["installation_id"], "", "package switch should clear stale installation id")
    _assert_equal(lc.installation_id_state[device_id]["game_code"], "CP", "package switch should refresh game code")

    lc._update_installation_id_runtime(device_id, package_id="com.nostel.parking.car", installation_id="id-new", force_emit=True)
    _assert_equal(lc.installation_id_state[device_id]["installation_id"], "id-new", "new package should accept a fresh installation id")


def test_installation_id_log_parsing() -> None:
    _reset_runtime_state()
    device_id = "device-2"

    original_get_package = lc._get_android_foreground_package
    try:
        lc._get_android_foreground_package = lambda _device_id: "com.indiez.nonogram"
        lc.process_installation_id_log(
            '16:31:37.056 SS S21 Unity 07-02 16:31:37.056  3425  3954 I Unity   : [Firebase] FirebaseInstallationIdPostInitHandler->_DebugInstallationId: {"idTask":{"idTask.IsFaulted":false,"idTask.Result":"d_vP_hhkSc2Z2c1b215kBR"}}',
            device_id,
        )
        _assert_equal(lc.installation_id_state[device_id]["installation_id"], "d_vP_hhkSc2Z2c1b215kBR", "json installation id log not parsed")

        lc._get_android_foreground_package = lambda _device_id: "com.nostel.parking.car"
        lc.process_installation_id_log("18:11:37.095\tUnity\tInstallations id fg88SZWfT6e_L4CxDX04Gp", device_id)
        _assert_equal(lc.installation_id_state[device_id]["installation_id"], "fg88SZWfT6e_L4CxDX04Gp", "plain installation id log not parsed")
        _assert_equal(lc.installation_id_state[device_id]["game_code"], "CP", "installation id log should follow active package")
    finally:
        lc._get_android_foreground_package = original_get_package


def test_sdk_exact_contracts() -> None:
    _reset_runtime_state()
    android_line = "Initializing Firebase Crashlytics 20.0.5"
    ios_line = "[Firebase/Crashlytics] Version 12.12.0"
    _assert("Initializing Firebase Crashlytics 20.0.5" in android_line, "android crashlytics contract changed")
    _assert("[Firebase/Crashlytics] Version 12.12.0" in ios_line, "ios crashlytics contract changed")
    _assert_equal(lc._extract_sdk_comparable_version(android_line), "20.0.5", "android crashlytics version should stay exact")
    _assert_equal(lc._extract_sdk_comparable_version(ios_line), "12.12.0", "ios crashlytics version should stay exact")


def test_cloudx_sdk_adapter_metadata() -> None:
    original_platform = lc.active_platform
    original_devices = lc.connected_devices_info
    original_sdk_active = lc.sdk_check_active
    original_search_list = list(lc.sdk_check_search_list)
    try:
        lc.active_platform = "android"
        lc.connected_devices_info = []
        lc.sdk_check_active = True
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        lc.sdk_check_runtime_state = {}
        lc.sdk_check_current_network = {}

        labels = [
            ("Digital Turbine (fyber) - Cloudx", "Fyber", "1.0.0"),
            ("InMobi - Cloudx", "InMobi", "2.0.0"),
            ("Liftoff Monetization (vungle) - Cloudx", "Vungle", "3.0.0"),
            ("Meta Audience Network - Cloudx", "Meta", "4.0.0"),
            ("Mintegral - Cloudx", "Mintegral", "5.0.0"),
            ("Mobilefuse - Cloudx", "Mobilefuse", "6.0.0"),
            ("Moloco - Cloudx", "Moloco", "7.0.0"),
            ("Pangle - Cloudx", "Pangle", "8.0.0"),
            ("UnityAds - Cloudx", "Unity", "9.0.0"),
            ("Verve / Pubnative - Cloudx", "Pubnative", "10.0.0"),
            ("TaurusX - Cloudx", "TaurusX", "11.0.0"),
        ]

        cloudx_parsed = lc._parse_sdk_expected_line("CloudX\t4.5.0\t4.5.0")
        _assert(cloudx_parsed is not None, "CloudX network label not accepted")
        _assert_equal(cloudx_parsed.get("source"), "", "CloudX must not be treated as a - Cloudx row")
        lc._register_sdk_expected("CloudX", adapter="4.5.0", sdk="4.5.0")
        cloudx_http_log = (
            '08-14 17:18:30.818 D CloudX [CloudXHttpClient] '
            '[{"sdk":{"sdkVersion":"4.5.0","pluginVersion":"unity-4.5.0"}}]'
        )
        _assert(
            lc._process_sdk_cloudx_http_metadata_line(cloudx_http_log, "cloudx-device"),
            "CloudX telemetry log was not parsed",
        )
        cloudx_state = lc.sdk_check_runtime_state["cloudx-device"]["cloudx"]
        _assert_equal(cloudx_state.get("adapter_version"), "4.5.0", "CloudX plugin version was not normalized")
        _assert_equal(cloudx_state.get("sdk_version"), "4.5.0", "CloudX SDK version was not parsed")
        _assert(
            not lc._process_sdk_cloudx_http_metadata_line(
                '[CloudXHttpClientExtra] [{"sdk":{"sdkVersion":"9.9.9","pluginVersion":"unity-9.9.9"}}]',
                "cloudx-device",
            ),
            "CloudX parser must require the exact CloudXHttpClient marker",
        )
        suffix_state = lc.sdk_check_runtime_state["cloudx-device"].get("mintegralcloudx")
        _assert(not suffix_state, "CloudX telemetry must not update a - Cloudx network")

        for display_name, _, expected_version in labels:
            parsed = lc._parse_sdk_expected_line(display_name)
            _assert(parsed is not None, f"Cloudx network label not accepted: {display_name}")
            _assert_equal(parsed.get("source"), "cloudx", f"Cloudx source missing for {display_name}")
            lc._register_sdk_expected(
                display_name,
                adapter=expected_version,
                source=parsed.get("source", ""),
                match_network=parsed.get("match_network", ""),
            )

        for _, actual_name, actual_version in labels:
            line = (
                "*[AdapterMetadataResolver] Discovered adapter metadata: "
                f"network={actual_name}, adapterVersion={actual_version},*"
            )
            lc._process_sdk_cloudx_adapter_metadata_line(line, "cloudx-device")

        _assert_equal(len(lc.sdk_check_expected_map), len(labels) + 1, "Cloudx network count mismatch")
        for display_name, _, expected_version in labels:
            expected_key = lc._normalize_sdk_network_name(display_name)
            actual = lc.sdk_check_runtime_state["cloudx-device"][expected_key]["adapter_version"]
            _assert_equal(actual, expected_version, f"Cloudx adapter version mismatch for {display_name}")

        cloudx_native_versions = {
            "digitalturbinefybercloudx": "8.4.7",
            "inmobicloudx": "11.4.0",
            "liftoffmonetizationvunglecloudx": "7.7.7",
            "metaaudiencenetworkcloudx": "6.22.0",
            "mintegralcloudx": "17.1.71",
            "mobilefusecloudx": "1.11.0",
            "molococloudx": "4.11.0",
            "panglecloudx": "8.2.0.4",
            "unityadscloudx": "4.19.0",
            "vervepubnativecloudx": "3.9.0",
            "taurusxcloudx": "1.18.3",
        }
        cloudx_actual_names = {
            "digitalturbinefybercloudx": "Fyber",
            "inmobicloudx": "InMobi",
            "liftoffmonetizationvunglecloudx": "Vungle",
            "metaaudiencenetworkcloudx": "Meta",
            "mintegralcloudx": "Mintegral",
            "mobilefusecloudx": "Mobilefuse",
            "molococloudx": "Moloco",
            "panglecloudx": "Pangle",
            "unityadscloudx": "Unity",
            "vervepubnativecloudx": "Pubnative",
            "taurusxcloudx": "TaurusX",
        }
        for expected_key, native_version in cloudx_native_versions.items():
            lc.sdk_check_expected_map[expected_key]["sdk"] = native_version

        mintegral_key = lc._normalize_sdk_network_name("Mintegral - Cloudx")
        lc._process_sdk_cloudx_adapter_metadata_line(
            "08-14 15:26:50.991 com.indiez.nonogram 20641 20791 D CloudX "
            "[AdapterMetadataResolver] Discovered adapter metadata: "
            "network=Mintegral, adapterVersion=17.1.71.0, "
            "networkSdkVersion=17.1.71, minimumSdkVersion=4.2.0, "
            "minimumSdkVersionCode=4002000, adapterApiVersion=1, extras=Bundle[]",
            "cloudx-device",
        )
        mintegral_cloudx_state = lc.sdk_check_runtime_state["cloudx-device"][mintegral_key]
        _assert_equal(mintegral_cloudx_state.get("adapter_version"), "17.1.71.0", "Mintegral Cloudx adapter version mismatch")
        _assert_equal(mintegral_cloudx_state.get("native_version"), "17.1.71", "Mintegral Cloudx native version was not parsed")

        for expected_key, native_version in cloudx_native_versions.items():
            display_name = lc.sdk_check_expected_map[expected_key]["display_name"]
            lc._process_sdk_cloudx_adapter_metadata_line(
                f"[AdapterMetadataResolver] Discovered adapter metadata: network={cloudx_actual_names[expected_key]}, "
                f"adapterVersion=1.0.0, networkSdkVersion={native_version}, minimumSdkVersion=1.0.0",
                "cloudx-native-device",
            )
            _assert_equal(
                lc.sdk_check_runtime_state["cloudx-native-device"][expected_key].get("native_version"),
                native_version,
                f"Cloudx native version mismatch for {display_name}",
            )

        emitted = []
        original_emit = lc.socketio.emit
        try:
            lc.connected_devices_info = [{"id": "cloudx-native-device", "name": "CloudX device"}]
            lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
            lc._emit_sdk_check_results()
        finally:
            lc.socketio.emit = original_emit
        emitted_rows = next(payload for event, payload in emitted if event == "update_sdk_check_table")
        native_rows = [row for row in emitted_rows if row.get("display_text", "").startswith("Native Version")]
        _assert_equal(len(native_rows), len(cloudx_native_versions), "Cloudx native result row count changed")
        _assert(all(row.get("status") == "PASSED" for row in native_rows), "Cloudx native result must pass matching versions")

        lc._process_sdk_cloudx_adapter_metadata_line(
            "[AdapterMetadataResolver] Discovered adapter metadata: "
            "network=Mintegral, adapterVersion=17.1.71.0, networkSdkVersion=17.1.70",
            "cloudx-native-device",
        )
        emitted = []
        original_emit = lc.socketio.emit
        try:
            lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
            lc._emit_sdk_check_results()
        finally:
            lc.socketio.emit = original_emit
        emitted_rows = next(payload for event, payload in emitted if event == "update_sdk_check_table")
        mintegral_native_rows = [
            row for row in emitted_rows
            if row.get("display_text", "").startswith("Native Version")
            and "17.1.70" in row.get("display_text", "")
        ]
        _assert_equal(len(mintegral_native_rows), 1, "Mintegral Cloudx native mismatch row missing")
        _assert_equal(mintegral_native_rows[0].get("status"), "FAILED", "wrong Cloudx native version must fail")

        lc._register_sdk_expected("Mintegral", adapter="5.0.0", sdk="17.1.71")
        lc._process_sdk_check_line("IntegrationHelper ----- Mintegral -----", "cloudx-device")
        lc._process_sdk_check_line("IntegrationHelper Adapter Version - 5.0.0", "cloudx-device")
        mintegral_standard_key = lc._normalize_sdk_network_name("Mintegral")
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"][mintegral_standard_key]["adapter_version"],
            "5.0.0",
            "IntegrationHelper Mintegral state was not kept separate",
        )
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"][mintegral_key].get("native_version"),
            "17.1.71",
            "IntegrationHelper Mintegral changed Cloudx native state",
        )

        lc._register_sdk_expected("Meta Audience Network", adapter="1.0.0")
        lc._process_sdk_check_line("IntegrationHelper ----- Meta Audience Network -----", "cloudx-device")
        lc._process_sdk_check_line("IntegrationHelper Adapter Version - 1.0.0", "cloudx-device")
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"]["metaaudiencenetwork"]["adapter_version"],
            "1.0.0",
            "IntegrationHelper state was not kept separate",
        )
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"]["metaaudiencenetworkcloudx"]["adapter_version"],
            "4.0.0",
            "Cloudx state was overwritten by IntegrationHelper",
        )

        _assert_equal(
            lc._match_sdk_expected_key("Liftoff Monetization"),
            "",
            "IntegrationHelper matcher must not fall through to Cloudx-only entries",
        )
        lc._register_sdk_expected("Liftoff Monetization (vungle)", adapter="12.0.0")
        liftoff_key = lc._normalize_sdk_network_name("Liftoff Monetization (vungle)")
        _assert_equal(
            lc._match_sdk_expected_key("Liftoff Monetization"),
            liftoff_key,
            "Liftoff Monetization alias did not match the list label",
        )
        _assert_equal(
            lc._match_sdk_expected_key("Liftoff Monetize"),
            liftoff_key,
            "Liftoff Monetize alias did not match the list label",
        )
        lc._process_sdk_check_line("IntegrationHelper ----- Liftoff Monetize -----", "cloudx-device")
        lc._process_sdk_check_line("IntegrationHelper Adapter Version - 12.0.0", "cloudx-device")
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"][liftoff_key]["adapter_version"],
            "12.0.0",
            "IntegrationHelper Liftoff alias did not update the standard entry",
        )
        _assert_equal(
            lc.sdk_check_runtime_state["cloudx-device"]["liftoffmonetizationvunglecloudx"].get("adapter_version"),
            "3.0.0",
            "IntegrationHelper Liftoff alias overwrote Cloudx state",
        )
    finally:
        lc.active_platform = original_platform
        lc.connected_devices_info = original_devices
        lc.sdk_check_active = original_sdk_active
        lc.sdk_check_search_list[:] = original_search_list
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        lc.sdk_check_runtime_state = {}
        lc.sdk_check_current_network = {}


def test_max_sdk_logs() -> None:
    original_platform = lc.active_platform
    original_devices = lc.connected_devices_info
    original_sdk_active = lc.sdk_check_active
    original_search_list = list(lc.sdk_check_search_list)
    original_expected_map = lc.sdk_check_expected_map
    original_expected_order = lc.sdk_check_expected_order
    original_runtime_state = lc.sdk_check_runtime_state
    original_current_network = lc.sdk_check_current_network
    original_emit = lc.socketio.emit
    try:
        lc.active_platform = "android"
        lc.connected_devices_info = []
        lc.sdk_check_active = True
        lc.sdk_check_search_list = []
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        lc.sdk_check_runtime_state = {}
        lc.sdk_check_current_network = {}

        c193_lines = lc._load_sdk_check_presets()["C-193-Android"]["lines"]
        max_expected = {}
        for line in c193_lines:
            if " - MAX" not in line:
                continue
            parsed = lc._parse_sdk_expected_line(line)
            _assert(parsed is not None, f"MAX entry cannot be parsed: {line}")
            _assert_equal(parsed.get("source"), "max", f"MAX entry source missing: {line}")
            max_expected[parsed["network"]] = parsed

        original_emit = lc.socketio.emit
        lc.socketio.emit = lambda *_args, **_kwargs: None
        try:
            lc.sdk_check({"text": "\n".join(c193_lines)})
        finally:
            lc.socketio.emit = original_emit

        core_key = lc._normalize_sdk_network_name("MAX / AppLovin - MAX")
        failed_core_line = (
            "09-03 14:38:39.685 com.indiez.nonogram 7428 8001 D AppLovinSdk "
            "[TaskInitializeSdk] AppLovin SDK 99.9.9 initialization failed"
        )
        _assert(not lc._process_sdk_max_line(failed_core_line, "max-device"), "failed MAX core init must be ignored")
        _assert(core_key not in lc.sdk_check_runtime_state.get("max-device", {}), "failed MAX core init created state")

        original_emit = lc.socketio.emit
        lc.socketio.emit = lambda *_args, **_kwargs: None
        try:
            core_line = (
                "09-03 14:38:39.685\tcom.indiez.nonogram\t7428\t8001\tD\tAppLovinSdk\t"
                "[TaskInitializeSdk] AppLovin SDK 13.6.4 initialization succeeded in 94ms"
            )
            lc._process_sdk_check_line(core_line, "max-device")
        finally:
            lc.socketio.emit = original_emit

        core_state = lc.sdk_check_runtime_state["max-device"][core_key]
        _assert_equal(core_state.get("sdk_version"), "13.6.4", "MAX core SDK version was not parsed")
        _assert_equal(core_state.get("adapter_version"), "", "MAX core line must update only one SDK version")

        started_core_line = core_line.replace("initialization succeeded", "initialization started").replace("13.6.4", "99.9.9")
        _assert(not lc._process_sdk_max_line(started_core_line, "max-device"), "MAX core parser accepted a non-succeeded line")
        _assert_equal(core_state.get("sdk_version"), "13.6.4", "non-succeeded MAX core line overwrote the version")

        facebook_line = (
            "09-03 13:58:14.758\tcom.indiez.nonogram\t24597\t25054\tD\tAppLovinSdk\t"
            "[HealthEventsReporter] Reporting signal_collection_success with extra parameters "
            "collection [{duration_ms=0, ad_format=REWARDED, ad_unit_id=08072f2e4324b79d, "
            "adapter_version=6.22.0.0, network_name=FACEBOOK_NETWORK, "
            "adapter_class=com.applovin.mediation.adapters.FacebookMediationAdapter}]"
        )
        original_emit = lc.socketio.emit
        lc.socketio.emit = lambda *_args, **_kwargs: None
        try:
            lc._process_sdk_check_line(facebook_line, "max-device")
        finally:
            lc.socketio.emit = original_emit

        facebook_key = lc._normalize_sdk_network_name("Meta Audience Network - MAX")
        facebook_state = lc.sdk_check_runtime_state["max-device"][facebook_key]
        _assert_equal(facebook_state.get("adapter_version"), "6.22.0.0", "FACEBOOK_NETWORK adapter version was not parsed")
        _assert_equal(facebook_state.get("observed_network_name"), "FACEBOOK_NETWORK", "MAX network name was not retained")
        _assert_equal(
            lc._sdk_result_status(facebook_state.get("adapter_version"), max_expected["Meta Audience Network - MAX"]["adapter"]),
            "PASSED",
            "FACEBOOK_NETWORK adapter version did not compare against the preset",
        )

        alias_cases = [
            ("IRONSOURCE_NETWORK", "LevelPlay / ironSource - MAX"),
            ("BIDMACHINE_NETWORK", "BidMachine - MAX"),
            ("BIGO_NETWORK", "Bigo Ads - MAX"),
            ("CHARTBOOST_NETWORK", "Chartboost - MAX"),
            ("DT_EXCHANGE", "Digital Turbine (fyber) - MAX"),
            ("GOOGLE_AD_MANAGER_NETWORK", "Google (AdMob and Ad Manager) - MAX"),
            ("INMOBI_NETWORK", "InMobi - MAX"),
            ("LINE_NETWORK", "LINE Ads - MAX"),
            ("LIFTOFF_NETWORK", "Liftoff Monetization (vungle) - MAX"),
            ("MINTEGRAL_NETWORK", "Mintegral - MAX"),
            ("MOBILEFUSE_NETWORK", "Mobilefuse - MAX"),
            ("MOLOCO_NETWORK", "Moloco - MAX"),
            ("OGURY_NETWORK", "Ogury - MAX"),
            ("PANGLE_NETWORK", "Pangle (Tiktok) - MAX"),
            ("PUBMATIC_NETWORK", "PubMatic (OpenWrap) - MAX"),
            ("UNITY_ADS_NETWORK", "UnityAds - MAX"),
            ("VERVE_NETWORK", "Verve / Pubnative - MAX"),
            ("YANDEX_NETWORK", "Yandex - MAX"),
            ("YSO_NETWORK", "YSO - MAX"),
            ("ASCENDX_NETWORK", "Ascendx - MAX"),
            ("MATICOO_NETWORK", "Yeahmobi/ Maticoo - MAX"),
            ("TAURUSX_NETWORK", "TaurusX - MAX"),
            ("PRADO_NETWORK", "Prado - MAX"),
            ("BIGO_BIDDING", "Bigo Ads - MAX"),
            ("BIGO_NATIVE_BIDDING", "Bigo Ads - MAX"),
            ("CHARTBOOST_BIDDING", "Chartboost - MAX"),
            ("BIDMACHINE_BIDDING", "BidMachine - MAX"),
            ("FYBER_BIDDING", "Digital Turbine (fyber) - MAX"),
            ("INMOBI_NATIVE_BIDDING", "InMobi - MAX"),
            ("MINTEGRAL_NATIVE_BIDDING", "Mintegral - MAX"),
            ("TIKTOK_NATIVE_BIDDING", "Pangle (Tiktok) - MAX"),
            ("OGURY_PRESAGE_BIDDING", "Ogury - MAX"),
            ("VUNGLE_NATIVE_BIDDING", "Liftoff Monetization (vungle) - MAX"),
            ("MOLOCO_NATIVE_BIDDING", "Moloco - MAX"),
        ]
        for observed_name, display_name in alias_cases:
            expected = max_expected.get(display_name)
            if not expected:
                continue
            version = expected["adapter"]
            _assert(version, f"MAX adapter preset is empty: {display_name}")
            metadata_line = (
                "[HealthEventsReporter] Reporting signal_collection_success with extra parameters "
                f"collection [{{adapter_version={version}, network_name={observed_name}, "
                "adapter_class=com.applovin.mediation.adapters.TestMediationAdapter}}]"
            )
            _assert(lc._process_sdk_max_line(metadata_line, "alias-device"), f"MAX metadata did not update: {display_name}")
            expected_key = lc._normalize_sdk_network_name(display_name)
            actual = lc.sdk_check_runtime_state["alias-device"][expected_key].get("adapter_version")
            _assert_equal(actual, version, f"MAX alias mapping failed for {observed_name}")
            _assert_equal(lc._sdk_result_status(actual, version), "PASSED", f"MAX version compare failed for {display_name}")
    finally:
        lc.active_platform = original_platform
        lc.connected_devices_info = original_devices
        lc.sdk_check_active = original_sdk_active
        lc.sdk_check_search_list[:] = original_search_list
        lc.sdk_check_expected_map = original_expected_map
        lc.sdk_check_expected_order = original_expected_order
        lc.sdk_check_runtime_state = original_runtime_state
        lc.sdk_check_current_network = original_current_network
        lc.socketio.emit = original_emit


def test_ios_max_sdk_search_only() -> None:
    """iOS MAX search must be isolated from the legacy iOS log reader."""
    original_platform = lc.active_platform
    original_sdk_active = lc.sdk_check_active
    original_search_list = list(lc.sdk_check_search_list)
    original_expected_map = lc.sdk_check_expected_map
    original_expected_order = lc.sdk_check_expected_order
    original_runtime_state = lc.sdk_check_runtime_state
    original_current_network = lc.sdk_check_current_network
    original_pending = dict(lc.sdk_max_ios_pending_lines)
    original_core_pending = dict(lc.sdk_max_ios_core_pending_lines)
    original_emit = lc.socketio.emit
    try:
        lc.active_platform = "ios"
        lc.sdk_check_active = False
        lc.sdk_check_search_list = []
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        lc.sdk_check_runtime_state = {}
        lc.sdk_check_current_network = {}
        lc.sdk_max_ios_pending_lines = {}
        lc.sdk_max_ios_core_pending_lines = {}
        lc.socketio.emit = lambda *_args, **_kwargs: None

        lc.sdk_check({
            "text": (
                "Ads Network\tAdapter\tNative\n"
                "MAX / AppLovin - MAX\t\t13.6.4\n"
                "Yandex - MAX\t8.5.0.1\t"
            )
        })

        core_line = (
            "Sep 18 17:10:58 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
            "[TaskInitializeSdk] AppLovin SDK 13.6.4 initialization succeeded"
        )
        lc._process_sdk_check_line(core_line, "ios-device")

        max_lines = [
            (
                "Sep 18 17:10:59 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
                "[AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting signal_collection_success "
                "with extra parameters {"
            ),
            '    "adapter_version" = "8.5.0.1";',
            '    "network_name" = "YANDEX_BIDDING";',
            "}",
        ]
        for line in max_lines:
            lc._process_sdk_check_line(line, "ios-device")

        core_key = lc._normalize_sdk_network_name("MAX / AppLovin - MAX")
        core_state = lc.sdk_check_runtime_state["ios-device"][core_key]
        _assert_equal(core_state.get("sdk_version"), "13.6.4", "iOS MAX core SDK version was not parsed")

        yandex_key = lc._normalize_sdk_network_name("Yandex - MAX")
        yandex_state = lc.sdk_check_runtime_state["ios-device"][yandex_key]
        _assert_equal(yandex_state.get("adapter_version"), "8.5.0.1", "iOS MAX adapter version was not parsed")
        _assert_equal(
            yandex_state.get("observed_network_name"),
            "YANDEX_BIDDING",
            "iOS MAX network name was not parsed",
        )

        core_block_lines = [
            "Sep 18 18:29:05 iPhone-11-pro PixelArt(AppLovinSDK)[2628] <Notice>: [AppLovinSdk] DEBUG [AppLovinSdk]",
            "========================================",
            "=====AppLovin SDK=====",
            "===SDK Versions===",
            "Version: 13.6.4",
        ]
        for line in core_block_lines:
            lc._process_sdk_check_line(line, "ios-core-device")
        core_block = lc.sdk_check_runtime_state["ios-core-device"][core_key]
        _assert_equal(
            core_block.get("sdk_version"),
            "13.6.4",
            "iOS MAX AppLovin SDK version block was not parsed",
        )

        source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
        _assert(
            "process_ios_package_log_stream_chunk(device_id, raw_chunk)" in source_text,
            "iOS reader must route raw framing to Package Log",
        )
        _assert(
            "active_ios_package_log_stream_modes" in source_text
            and "for raw_line in raw_chunk.splitlines():" in source_text,
            "iOS Package Log must detect newline streams instead of assuming NUL framing",
        )
        _assert("pending_record" not in source_text, "MAX SDK change must not add a shared iOS log buffer")
    finally:
        lc.active_platform = original_platform
        lc.sdk_check_active = original_sdk_active
        lc.sdk_check_search_list[:] = original_search_list
        lc.sdk_check_expected_map = original_expected_map
        lc.sdk_check_expected_order = original_expected_order
        lc.sdk_check_runtime_state = original_runtime_state
        lc.sdk_check_current_network = original_current_network
        lc.sdk_max_ios_pending_lines = original_pending
        lc.sdk_max_ios_core_pending_lines = original_core_pending
        lc.socketio.emit = original_emit


def test_audiomob_open_measurement_sdk_log() -> None:
    """AudioMob Android logs expose the SDK version inside Open Measurement text."""
    original_platform = lc.active_platform
    original_devices = lc.connected_devices_info
    original_sdk_active = lc.sdk_check_active
    original_search_list = list(lc.sdk_check_search_list)
    original_expected_map = lc.sdk_check_expected_map
    original_expected_order = lc.sdk_check_expected_order
    original_runtime_state = lc.sdk_check_runtime_state
    original_current_network = lc.sdk_check_current_network
    original_emit = lc.socketio.emit
    try:
        lc.active_platform = "android"
        lc.connected_devices_info = [{"id": "audio-device", "name": "AudioMob device"}]
        lc.sdk_check_active = True
        lc.sdk_check_search_list = []
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        lc.sdk_check_runtime_state = {}
        lc.sdk_check_current_network = {}
        emitted = []
        lc.socketio.emit = lambda event, payload=None, **_kwargs: emitted.append((event, payload))

        parsed = lc._parse_sdk_expected_line("AudioMob\t10.2.3")
        _assert(parsed is not None, "AudioMob preset row cannot be parsed")
        lc._register_sdk_expected(
            parsed["network"],
            adapter=parsed["adapter"],
            sdk=parsed["sdk"],
            log_search=parsed["log_search"],
        )

        log_line = (
            "09-28 10:36:05.716\tcom.indiez.nonogram\t3646\t-\tI\tAudiomob\t"
            "Open Measurement started (SDK: 10.2.3)."
        )
        lc._process_sdk_check_line(log_line, "audio-device")

        key = lc._normalize_sdk_network_name("AudioMob")
        state = lc.sdk_check_runtime_state["audio-device"][key]
        _assert_equal(state.get("adapter_version"), "10.2.3", "AudioMob Open Measurement SDK version was not captured")
        _assert_equal(
            lc._sdk_result_status(state.get("adapter_version"), parsed["adapter"]),
            "PASSED",
            "AudioMob Open Measurement SDK version did not compare against the preset",
        )
        rows = next(payload for event, payload in reversed(emitted) if event == "update_sdk_check_table")
        _assert(
            any(row.get("status") == "PASSED" and "10.2.3" in row.get("display_text", "") for row in rows),
            "AudioMob SDK check did not emit a PASSED result",
        )
    finally:
        lc.active_platform = original_platform
        lc.connected_devices_info = original_devices
        lc.sdk_check_active = original_sdk_active
        lc.sdk_check_search_list[:] = original_search_list
        lc.sdk_check_expected_map = original_expected_map
        lc.sdk_check_expected_order = original_expected_order
        lc.sdk_check_runtime_state = original_runtime_state
        lc.sdk_check_current_network = original_current_network
        lc.socketio.emit = original_emit


def test_sdk_preset_switch_does_not_race_stop_start() -> None:
    """Changing presets must replace the running check with one atomic start."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    apply_block = source.split("function applySdkCheckPreset(name) {", 1)[1].split(
        "renderSdkCheckPresetOptions();", 1
    )[0]
    _assert(
        "socket.emit('start_sdk_check', {text: input.value});" in apply_block,
        "preset selection must start the newly selected SDK check",
    )
    _assert(
        "socket.emit('stop_sdk_check');" not in apply_block,
        "preset selection must not race a stop event against the replacement start",
    )


def test_clear_all_resets_sdk_check_state() -> None:
    """Clear All must leave SDK Check ready for a fresh preset selection."""
    original_active = lc.sdk_check_active
    original_search_list = lc.sdk_check_search_list
    original_input_list = lc.sdk_check_input_list
    original_expected_map = lc.sdk_check_expected_map
    original_runtime_state = lc.sdk_check_runtime_state
    original_current_network = lc.sdk_check_current_network
    original_expected_order = lc.sdk_check_expected_order
    original_ios_pending = lc.sdk_max_ios_pending_lines
    original_ios_core_pending = lc.sdk_max_ios_core_pending_lines
    original_emit = lc.socketio.emit
    try:
        lc.sdk_check_active = True
        lc.sdk_check_search_list = [{"network": "Old SDK"}]
        lc.sdk_check_input_list = [{"type": "label", "display_name": "Old SDK"}]
        lc.sdk_check_expected_map = {"oldsdk": {"display_name": "Old SDK"}}
        lc.sdk_check_runtime_state = {"device": {"oldsdk": {"adapter_version": "1.0.0"}}}
        lc.sdk_check_current_network = {"device": "Old SDK"}
        lc.sdk_check_expected_order = ["oldsdk"]
        lc.sdk_max_ios_pending_lines = {"device": "pending"}
        lc.sdk_max_ios_core_pending_lines = {"device": "pending"}
        lc.socketio.emit = lambda *_args, **_kwargs: None

        lc.cl()

        _assert(not lc.sdk_check_active, "Clear All must stop the old SDK Check session")
        _assert(not lc.sdk_check_search_list, "Clear All left the old SDK search list")
        _assert(not lc.sdk_check_input_list, "Clear All left the old SDK input list")
        _assert(not lc.sdk_check_expected_map, "Clear All left the old SDK preset map")
        _assert(not lc.sdk_check_runtime_state, "Clear All left old SDK runtime results")
        _assert(not lc.sdk_check_current_network, "Clear All left the old SDK current network")
        _assert(not lc.sdk_check_expected_order, "Clear All left the old SDK preset order")
        _assert(not lc.sdk_max_ios_pending_lines, "Clear All left pending iOS MAX SDK data")
        _assert(not lc.sdk_max_ios_core_pending_lines, "Clear All left pending iOS MAX core data")
    finally:
        lc.sdk_check_active = original_active
        lc.sdk_check_search_list = original_search_list
        lc.sdk_check_input_list = original_input_list
        lc.sdk_check_expected_map = original_expected_map
        lc.sdk_check_runtime_state = original_runtime_state
        lc.sdk_check_current_network = original_current_network
        lc.sdk_check_expected_order = original_expected_order
        lc.sdk_max_ios_pending_lines = original_ios_pending
        lc.sdk_max_ios_core_pending_lines = original_ios_core_pending
        lc.socketio.emit = original_emit


def test_platform_reconnect_contract() -> None:
    """A server restart must restore the persisted platform on Socket.IO reconnect."""
    rendered = lc.app.test_client().get("/").get_data(as_text=True)
    _assert("socket.on('connect'" in rendered, "platform reconnect handler is missing")
    _assert("reset: false" in rendered, "platform reconnect must not clear runtime state")
    _assert("socket.emit('set_platform'" in rendered, "platform reconnect does not resync the backend")


def test_platform_bootstrap_sdk_state_contract() -> None:
    """Initial platform bootstrap must not hit the SDK Check temporal dead zone."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    running_declaration = source.find("let sdkCheckRunning = false;")
    reset_function = source.find("function resetSdkCheckUiState(clearInput = false)")
    _assert(running_declaration >= 0, "SDK Check running state declaration is missing")
    _assert(
        running_declaration < reset_function,
        "SDK Check running state must be declared before platform bootstrap calls resetSdkCheckUiState",
    )
    reset_block = source[reset_function:source.find("function resetRuntimeUiForPlatformSwitch", reset_function)]
    _assert(
        "typeof sdkCheckRunning" not in reset_block,
        "SDK Check reset must not use a temporal-dead-zone typeof guard",
    )
    bootstrap_block = source[source.find("platformBtn?.addEventListener('click', showPlatformModal);"):source.find("// --- Tab Logic ---")]
    _assert(
        "window.setTimeout(() =>" in bootstrap_block,
        "platform bootstrap must wait until all UI state declarations are initialized",
    )


def test_device_filter_reset_contract() -> None:
    """A stale device selection must not hide logs after a clear/platform switch."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert("function resetSelectedDeviceFilter()" in source, "device filter reset helper is missing")
    _assert(
        "resetSelectedDeviceFilter();\n            syncPlatformUi();" in source,
        "runtime reset must return the device filter to All Devices",
    )
    _assert(
        "resetRuntimeUiForPlatformSwitch();\n            socket.emit('set_platform'" in source,
        "platform switching must clear the stale device selection before resync",
    )
    _assert(
        "else {\n                resetSelectedDeviceFilter();\n            }" in source,
        "device discovery must reset a selection whose device disappeared",
    )


def test_platform_switch_drops_stale_device_status_contract() -> None:
    """A late discovery result from the old platform must be ignored."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert("platform_generation = 0" in source, "platform generation state is missing")
    _assert(
        "generation != platform_generation or active_platform != platform" in source,
        "device manager must invalidate an old discovery cycle",
    )
    _assert('"platform": platform' in source, "device status payload must identify its platform")
    _assert(
        "if (!statusPlatform || statusPlatform !== activePlatform) return;" in source,
        "frontend must reject untagged or stale device status from another platform",
    )
    set_platform_block = source.split("@socketio.on('set_platform')", 1)[1].split(
        "@socketio.on('toggle_record')", 1
    )[0]
    _assert("platform_changed" in set_platform_block, "backend platform transition must detect a stale platform")
    _assert(
        "or platform_changed" in set_platform_block,
        "backend platform transition must reset stale device state even when reset=false",
    )


def test_ios_reader_singleton_contract() -> None:
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    for marker in (
        "reader_thread = threading.current_thread()",
        "current_reader = active_ios_log_readers.get(device_id)",
        "if thread and thread.is_alive():",
        "thread.join(timeout=2)",
        "readers_to_stop = []",
        "for did, reason in readers_to_stop:",
    ):
        _assert(marker in source, f"iOS reader singleton guard is missing: {marker}")
    _assert(
        "active_ios_log_readers.get(device_id) is thread" in source,
        "iOS reader shutdown must not remove a replacement reader",
    )


def test_ios_transport_fallback_contract() -> None:
    """iOS discovery must ignore trusted network-paired tidevice rows."""
    original_resolve = lc._resolve_ios_tool
    original_run = lc.subprocess.run
    outputs = [
        """UDID                                      SerialNumber    NAME        MarketName             ProductVersion    ConnType
00008030-0009398422F3C02E  V149104TF2      iPad gen 9   iPad (9th generation)  18.5  ConnectionType.NETWORK
eb6b13cc453f5a53ef07ff7149858a8635d18c10  F17VXQA8JCLH    iPhone X     iPhone X               16.7.12  ConnectionType.NETWORK
""",
        """UDID                                      SerialNumber    NAME        MarketName             ProductVersion    ConnType
00008030-0009398422F3C02E  V149104TF2      iPad gen 9   iPad (9th generation)  18.5  ConnectionType.USB
eb6b13cc453f5a53ef07ff7149858a8635d18c10  F17VXQA8JCLH    iPhone X     iPhone X               16.7.12  ConnectionType.NETWORK
""",
    ]
    try:
        lc._resolve_ios_tool = lambda name: "/fake/tidevice" if name == "tidevice" else None
        lc.subprocess.run = lambda *args, **kwargs: types.SimpleNamespace(stdout=outputs.pop(0))

        network_ids = lc._list_ios_device_ids()
        _assert_equal(
            network_ids,
            [],
            "network-paired iOS devices must not be discovered",
        )
        usb_ids = lc._list_ios_device_ids()
        _assert_equal(
            usb_ids,
            ["00008030-0009398422F3C02E"],
            "USB discovery must ignore network-paired rows",
        )
        source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
        _assert(
            "selected_ids = ids or network_ids" not in source,
            "iOS discovery must not fall back to network-paired IDs",
        )
    finally:
        lc._resolve_ios_tool = original_resolve
        lc.subprocess.run = original_run


def test_ios_discovery_grace_contract() -> None:
    """A transient empty tidevice poll must not flap the connected-device UI."""
    device_id = "00008120-001C0D3C0C61A01E"
    original_missing = dict(lc.ios_device_missing_polls)
    try:
        lc.ios_device_missing_polls.clear()
        _assert_equal(
            lc._stabilize_ios_device_ids({device_id}, set()),
            {device_id},
            "observed iOS device was not retained",
        )
        _assert_equal(
            lc._stabilize_ios_device_ids(set(), {device_id}),
            {device_id},
            "first empty discovery poll incorrectly disconnected iOS device",
        )
        _assert_equal(
            lc._stabilize_ios_device_ids(set(), {device_id}),
            {device_id},
            "second empty discovery poll incorrectly disconnected iOS device",
        )
        _assert_equal(
            lc._stabilize_ios_device_ids(set(), {device_id}),
            set(),
            "device was retained after the discovery grace window",
        )
        _assert_equal(
            lc._stabilize_ios_device_ids({device_id}, set()),
            {device_id},
            "device was not rediscovered after a real disconnect",
        )
    finally:
        lc.ios_device_missing_polls.clear()
        lc.ios_device_missing_polls.update(original_missing)


def test_ios_reader_liveness_contract() -> None:
    """A streaming iOS reader survives discovery misses; silent stale ones stop; stalled ones restart."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    manager = source.split("def device_manager():", 1)[1].split("def _append_package_log_row", 1)[0]
    _assert("observed_ids = _observe_ios_device_ids()" in manager, "discovery must go through the throttled helper")
    _assert("_ios_readers_to_stop(observed_ids, time.time())" in manager, "reader decisions must use the tested helper")
    _assert("did not in ids" not in manager, "iOS discovery must not stop a live reader on one missing row")
    _assert("ids.difference_update(stale_reader_ids)" in manager, "stale iOS devices must be removed from the UI")
    _assert("set(active_ios_log_readers) - stale_reader_ids" in manager, "live iOS readers must stay visible during discovery grace")
    _assert("set(active_ios_log_processes) - stale_reader_ids" in manager, "live iOS processes must stay visible during discovery grace")
    _assert('env=dict(os.environ, PYTHONUNBUFFERED="1")' in source, "tidevice output must be unbuffered")

    class FakeProc:
        def __init__(self, code=None):
            self.code = code

        def poll(self):
            return self.code

    class FakeThread:
        def __init__(self, alive=True):
            self.alive = alive

        def is_alive(self):
            return self.alive

    state_names = ("active_ios_log_readers", "active_ios_log_processes", "active_ios_log_started_at", "active_ios_log_last_seen", "ios_device_missing_polls")
    saved = {name: dict(getattr(lc, name)) for name in state_names}
    saved_cache = dict(lc.ios_discovery_cache)
    saved_list = lc._list_ios_device_ids
    now = 10_000.0
    try:
        for name in state_names:
            getattr(lc, name).clear()

        def reader(did, started, last_seen, proc_code=None, alive=True, misses=0):
            lc.active_ios_log_readers[did] = FakeThread(alive)
            lc.active_ios_log_processes[did] = FakeProc(proc_code)
            lc.active_ios_log_started_at[did] = started
            lc.active_ios_log_last_seen[did] = last_seen
            if misses:
                lc.ios_device_missing_polls[did] = misses

        grace = lc.IOS_DEVICE_MISSING_GRACE_POLLS + 1
        reader("streaming-missing", now - 600, now - 2, misses=grace)    # discovery lost it, logs still flowing
        reader("silent-missing", now - 600, now - 60, misses=grace)      # unplugged: no logs, not discovered
        reader("stalled", now - 600, now - 45)                           # connected but silent 45 s
        reader("quiet", now - 600, now - 12)                             # short quiet period
        reader("just-started", now - 5, now - 5)                         # still warming up
        reader("dead", now - 600, now - 1, proc_code=1)                  # process exited
        lc.active_ios_log_processes["orphan"] = FakeProc(None)           # process without a reader thread
        observed = {"stalled", "quiet", "just-started", "dead", "orphan"}
        with lc.lock:
            stops, stale = lc._ios_readers_to_stop(observed, now)
        reasons = {did: reason for did, reason in stops}
        _assert_equal(stale, {"silent-missing"}, "only the silent, undiscovered reader is stale")
        _assert("streaming-missing" not in reasons, "a reader still receiving logs must survive discovery misses")
        _assert_equal(reasons.get("silent-missing"), "missing from discovery", "an unplugged silent reader must stop")
        _assert(str(reasons.get("stalled", "")).startswith("stalled: no logs for 45s"), f"a stuck reader must restart: {reasons}")
        _assert("quiet" not in reasons and "just-started" not in reasons, "short quiet periods must not restart a reader")
        _assert_equal(reasons.get("dead"), "dead", "an exited reader must restart")
        _assert_equal(reasons.get("orphan"), "orphan process", "an orphan process must stop")

        # Discovery is skipped while every reader streams, and runs again when one goes quiet.
        for name in state_names:
            getattr(lc, name).clear()
        calls = []
        lc._list_ios_device_ids = lambda: calls.append(1) or ["phone"]
        lc.ios_discovery_cache.update(at=0.0, ids=set())
        reader("phone", now - 600, now - 1)
        _assert_equal(lc._observe_ios_device_ids(now), {"phone"}, "first discovery must run")
        _assert_equal(lc._observe_ios_device_ids(now + 5), {"phone"}, "cached discovery result is wrong")
        _assert_equal(len(calls), 1, "discovery must be skipped while the reader streams")
        _assert_equal(lc._observe_ios_device_ids(now + lc.IOS_DISCOVERY_SLOW_INTERVAL_SECONDS + 1), {"phone"}, "slow discovery result is wrong")
        _assert_equal(len(calls), 2, "discovery must still run every IOS_DISCOVERY_SLOW_INTERVAL_SECONDS")
        lc.active_ios_log_last_seen["phone"] = now - 30
        lc._observe_ios_device_ids(now + lc.IOS_DISCOVERY_SLOW_INTERVAL_SECONDS + 2)
        _assert_equal(len(calls), 3, "discovery must run right away when a reader goes quiet")
    finally:
        lc._list_ios_device_ids = saved_list
        lc.ios_discovery_cache.clear()
        lc.ios_discovery_cache.update(saved_cache)
        for name in state_names:
            getattr(lc, name).clear()
            getattr(lc, name).update(saved[name])


def test_android_logcat_reconnect_cursor_contract() -> None:
    """Continuous Android streams must start near now, not at the ring-buffer head."""
    command = lc._android_logcat_stream_command("android-device")
    _assert("-T" in command, "Android logcat reader has no start cursor")
    cursor_index = command.index("-T") + 1
    cursor = command[cursor_index]
    _assert(re.fullmatch(r"\d{2}-\d{2} \d{2}:\d{2}:\d{2}\.000", cursor), "invalid Android logcat cursor")
    package_command = lc._android_logcat_stream_command("android-device", pid="1234")
    _assert("--pid=1234" in package_command, "Package Log lost the PID filter")
    _assert(package_command[cursor_index - 1] == "-T", "Package Log lost the start cursor")


def test_package_consumer_replacement_guard() -> None:
    """A stopped/replaced Package Log process must not append its tail."""
    original_platform = lc.active_platform
    original_target = lc.target_package_name
    original_rows = list(lc.package_log_cache)
    original_processes = dict(lc.active_logcat_processes)
    original_dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
    original_dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
    try:
        lc.active_platform = "android"
        lc.target_package_name = "com.example.current"
        lc.package_log_cache.clear()
        lc.active_logcat_processes.clear()
        lc._clear_record_dedup("package")
        old_process = type("FakeProcess", (), {"stdout": io.StringIO(
            "09-29 14:00:00.000  1  1 I Unity: stale row\n"
        )})()
        replacement_process = type("FakeProcess", (), {"stdout": io.StringIO()})()
        lc.active_logcat_processes["android-device"] = replacement_process
        lc.package_log_consumer("android-device", old_process)
        _assert_equal(len(lc.package_log_cache), 0, "stale Package Log consumer appended after replacement")
    finally:
        lc.active_platform = original_platform
        lc.target_package_name = original_target
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_rows)
        lc.active_logcat_processes.clear()
        lc.active_logcat_processes.update(original_processes)
        for name in lc.record_dedup_seen:
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_order[name].clear()
        for name, values in original_dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(original_dedup_order[name])


def test_clear_does_not_reopen_replay_dedup_contract() -> None:
    """Clear All must not erase the replay index while live readers continue."""
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    clear_block = source.split("@socketio.on('clear_all_logs')", 1)[1].split(
        "@socketio.on('start_validation')", 1
    )[0]
    platform_reset_block = source.split("def _reset_runtime_for_platform_switch():", 1)[1].split(
        "@socketio.on('set_platform')", 1
    )[0]
    _assert("_clear_record_dedup()" not in clear_block, "Clear All reopens exact replay duplicates")
    _assert("_clear_record_dedup()" not in platform_reset_block, "platform reset reopens exact replay duplicates")
    _assert(
        "Queue(maxsize=PACKAGE_LOG_DB_QUEUE_MAX_ITEMS)" in source,
        "Package Log database queue is not bounded",
    )


def test_ios_package_frame_buffer_has_memory_guard() -> None:
    """An unterminated iOS relay frame must not grow without bound."""
    original_platform = lc.active_platform
    original_frames = dict(lc.active_ios_package_log_frame_buffers)
    original_modes = dict(lc.active_ios_package_log_stream_modes)
    try:
        lc.active_platform = "ios"
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_stream_modes.clear()
        lc.process_ios_package_log_stream_chunk(
            "ios-device",
            "x" * (lc.MAX_IOS_PACKAGE_LOG_BUFFER_CHARS + 1),
        )
        _assert(
            len(lc.active_ios_package_log_frame_buffers.get("ios-device", ""))
            <= lc.MAX_IOS_PACKAGE_LOG_BUFFER_CHARS,
            "unterminated iOS Package Log frame exceeded its memory limit",
        )
    finally:
        lc.active_platform = original_platform
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_frame_buffers.update(original_frames)
        lc.active_ios_package_log_stream_modes.clear()
        lc.active_ios_package_log_stream_modes.update(original_modes)


def test_android_tracking_stream_reassembles_wrapped_records() -> None:
    """Wrapped Unity TrackingService JSON must reach every Android event tab."""
    original_platform = lc.active_platform
    original_paused = lc.is_paused
    original_validator_active = lc.validator_active
    original_emit = lc.socketio.emit
    original_callback_rows = list(lc.callback_ad_logs)
    original_validator_rows = list(lc.validator_results)
    original_specific_rows = list(lc.event_log_cache)
    original_specific_results = list(lc.specific_event_results)
    original_buffer = dict(lc.incomplete_android_event_logs)
    original_dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
    original_dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
    try:
        lc.active_platform = "android"
        lc.is_paused = False
        lc.validator_active = True
        lc.socketio.emit = lambda *_args, **_kwargs: None
        lc.callback_ad_logs.clear()
        lc.validator_results.clear()
        lc.event_log_cache.clear()
        lc.specific_event_results.clear()
        lc.incomplete_android_event_logs.clear()
        lc._clear_record_dedup()

        first = (
            '09-28 14:43:16.087  6927  7783 I Unity : [Tracking] '
            'TrackingService->Track: {"EventName":"ad_impression",'
        )
        second = (
            '"params":{"ad_platform":"ironsource","ad_network":"pangle",'
            '"ad_format":"banner","value":7.428e-05}}'
        )
        lc._process_android_runtime_line(first, "android-device")
        _assert_equal(len(lc.event_log_cache), 0, "incomplete Android event was emitted too early")
        lc._process_android_runtime_line(second, "android-device")

        _assert_equal(len(lc.callback_ad_logs), 1, "wrapped Android event missing from Callback & Ads")
        _assert_equal(len(lc.validator_results), 1, "wrapped Android event missing from Validator")
        _assert_equal(len(lc.event_log_cache), 1, "wrapped Android event missing from Specific Validator")
        _assert_equal(lc.callback_ad_logs[-1]["event_name"], "ad_impression", "wrapped event name changed")
    finally:
        lc.active_platform = original_platform
        lc.is_paused = original_paused
        lc.validator_active = original_validator_active
        lc.socketio.emit = original_emit
        lc.callback_ad_logs.clear()
        lc.callback_ad_logs.extend(original_callback_rows)
        lc.validator_results.clear()
        lc.validator_results.extend(original_validator_rows)
        lc.event_log_cache.clear()
        lc.event_log_cache.extend(original_specific_rows)
        lc.specific_event_results.clear()
        lc.specific_event_results.extend(original_specific_results)
        lc.incomplete_android_event_logs.clear()
        lc.incomplete_android_event_logs.update(original_buffer)
        for name in lc.record_dedup_seen:
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_order[name].clear()
        for name, values in original_dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(original_dedup_order[name])


def test_ios_tracking_stream_reassembles_repeated_header_wrap() -> None:
    """A wrapped iOS TrackingService event with a repeated header is recorded."""
    original_platform = lc.active_platform
    original_paused = lc.is_paused
    original_validator_active = lc.validator_active
    original_emit = lc.socketio.emit
    original_validator_rows = list(lc.validator_results)
    original_specific_rows = list(lc.event_log_cache)
    original_specific_results = list(lc.specific_event_results)
    original_buffer = dict(lc.incomplete_ios_event_logs)
    original_dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
    original_dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
    try:
        lc.active_platform = "ios"
        lc.is_paused = False
        lc.validator_active = True
        lc.socketio.emit = lambda *_args, **_kwargs: None
        lc.validator_results.clear()
        lc.event_log_cache.clear()
        lc.specific_event_results.clear()
        lc.incomplete_ios_event_logs.clear()
        lc._clear_record_dedup()

        first = (
            'Sep 29 13:47:22 Galaxy-iphone CatRestaurant(UnityFramework)[71830] <Notice>: '
            '[Tracking] TrackingService->Track: '
            '{"EventName":"loading_to_home","params":{"time":4.379315,'
            '"level":3,"m'
        )
        second = (
            'Sep 29 13:47:22 Galaxy-iphone CatRestaurant(UnityFramework)[71830] <Notice>: '
            '[Tracking] TrackingService->Track: ode_game":"null"}}'
        )

        lc._process_ios_runtime_line(first, "ios-device")
        _assert_equal(len(lc.validator_results), 0, "incomplete iOS event was emitted too early")
        lc._process_ios_runtime_line(second, "ios-device")

        _assert_equal(len(lc.validator_results), 1, "wrapped iOS event missing from Validator")
        _assert_equal(len(lc.event_log_cache), 1, "wrapped iOS event missing from Specific Validator")
        _assert_equal(
            lc.validator_results[-1]["event_name"],
            "loading_to_home",
            "wrapped iOS event name changed",
        )
        _assert_equal(
            lc.validator_results[-1]["details"].count("mode_game"),
            1,
            "wrapped iOS event JSON was not joined at the split point",
        )
    finally:
        lc.active_platform = original_platform
        lc.is_paused = original_paused
        lc.validator_active = original_validator_active
        lc.socketio.emit = original_emit
        lc.validator_results.clear()
        lc.validator_results.extend(original_validator_rows)
        lc.event_log_cache.clear()
        lc.event_log_cache.extend(original_specific_rows)
        lc.specific_event_results.clear()
        lc.specific_event_results.extend(original_specific_results)
        lc.incomplete_ios_event_logs.clear()
        lc.incomplete_ios_event_logs.update(original_buffer)
        for name in lc.record_dedup_seen:
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_order[name].clear()
        for name, values in original_dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(original_dedup_order[name])


class _PackageLogStateGuard:
    """Snapshot and restore the Package Log globals a protocol test touches."""

    NAMES = (
        "active_platform", "target_package_name", "active_package_log_session_id",
        "package_log_seq", "package_log_stream_id", "package_log_last_emitted_seq",
    )

    def __enter__(self):
        self.values = {name: getattr(lc, name) for name in self.NAMES}
        self.rows = list(lc.package_log_cache)
        self.dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
        self.dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
        self.emit = lc.socketio.emit
        self.emitted = []
        lc.socketio.emit = lambda event, payload=None, **kw: self.emitted.append((event, payload, kw.get("to")))
        lc.package_log_cache.clear()
        lc._drain_package_log_db_queue()
        lc._clear_record_dedup("package")
        lc.active_platform = "android"
        lc.active_package_log_session_id = None
        return self

    def __exit__(self, *exc):
        lc.socketio.emit = self.emit
        lc._drain_package_log_db_queue()
        for name, value in self.values.items():
            setattr(lc, name, value)
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(self.rows)
        for name in lc.record_dedup_seen:
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_order[name].clear()
        for name, values in self.dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(self.dedup_order[name])
        return False

    def events(self, name):
        return [payload for event, payload, _to in self.emitted if event == name]


def _start_package_stream_for_test(package_id="com.harness.package"):
    with lc.lock:
        lc.target_package_name = package_id
        return lc._bump_package_log_stream_locked()


def test_package_log_delta_sends_only_new_rows() -> None:
    """Live rows reach the UI once, in seq order, in chunks, without the raw copy."""
    with _PackageLogStateGuard() as guard:
        stream_id = _start_package_stream_for_test()
        total = lc.PACKAGE_LOG_DELTA_CHUNK_ROWS * 2 + 500
        for index in range(total):
            lc._append_package_log_row("harness-device", f"raw line {index}", "10-20 10:00:00.000",
                                       "10:00:00.000", "I", "Tag", f"message {index}", False, "123", "456")
        _assert_equal(lc._emit_pending_package_log_rows(), total, "emitter did not send every new row")
        deltas = guard.events("package_log_delta")
        _assert_equal(len(deltas), 3, "rows were not split into ordered chunks")
        seqs = [row["seq"] for payload in deltas for row in payload["rows"]]
        _assert_equal(len(seqs), total, "a delta row was skipped")
        _assert_equal(seqs, list(range(seqs[0], seqs[0] + total)), "delta seq order has a gap")
        _assert(all(payload["stream"] == stream_id for payload in deltas), "delta carries the wrong stream")
        first_row = deltas[0]["rows"][0]
        _assert("log" not in first_row, "delta rows must not carry the duplicate raw log")
        _assert_equal((first_row["pid"], first_row["tid"]), ("123", "456"), "pid/tid missing from the delta row")
        _assert_equal(lc._emit_pending_package_log_rows(), 0, "an already sent row was emitted again")
        _assert_equal(len(guard.events("package_log_delta")), 3, "no-op pass emitted a delta")


def test_package_log_reset_and_sync_protocol() -> None:
    """Reconnecting clients replay missed rows; a new stream or an old gap is reported."""
    with _PackageLogStateGuard() as guard:
        stream_id = _start_package_stream_for_test()
        for index in range(30):
            lc._append_package_log_row("harness-device", f"sync raw {index}", message=f"sync {index}")
        lc._emit_pending_package_log_rows()
        first_seq = lc.package_log_cache[0]["seq"]

        def sync(data):
            guard.emitted.clear()
            with lc.app.test_request_context("/"):
                lc.request.sid = "harness-sid"
                lc.handle_package_log_sync(data)
            return guard.emitted

        emitted = sync({"stream": stream_id, "since_seq": first_seq + 19})
        rows = [row for event, payload, to in emitted if event == "package_log_delta" for row in payload["rows"]]
        _assert_equal([row["seq"] for row in rows], list(range(first_seq + 20, first_seq + 30)), "sync replayed the wrong rows")
        _assert(all(to == "harness-sid" for _event, _payload, to in emitted), "sync reply must go only to the requesting client")
        _assert(all(payload["sync"] for event, payload, _to in emitted if event == "package_log_delta"), "sync reply is not marked")

        emitted = sync({"stream": stream_id - 1, "since_seq": first_seq + 25})
        _assert_equal(emitted[0][0], "package_log_reset", "a client on an old stream must be reset first")
        rows = [row for event, payload, _to in emitted if event == "package_log_delta" for row in payload["rows"]]
        _assert_equal(len(rows), 30, "a reset client must receive the whole current stream")

        for _ in range(5):
            lc.package_log_cache.popleft()
        emitted = sync({"stream": stream_id, "since_seq": first_seq})
        delta = next(payload for event, payload, _to in emitted if event == "package_log_delta")
        _assert(delta["gap"], "rows older than the cache window must be reported as a gap")

        lc._emit_package_log_reset()
        reset = guard.events("package_log_reset")[-1]
        _assert_equal(reset["stream"], stream_id + 1, "reset did not start a new stream")
        lc._append_package_log_row("harness-device", "after reset", message="after reset")
        _assert_equal(lc.package_log_cache[-1]["stream"], stream_id + 1, "new rows must follow the new stream")


def test_package_log_raw_and_db_seq_contract() -> None:
    """Raw logs are fetched on demand; SQLite rows store the same seq."""
    with _PackageLogStateGuard():
        _start_package_stream_for_test()
        lc.active_package_log_session_id = 987654
        lc._append_package_log_row("harness-device", "RAW-ONLY-PAYLOAD full record", message="short message")
        seq = lc.package_log_cache[-1]["seq"]
        queued = lc.package_log_db_queue.get_nowait()
        _assert_equal(len(queued), 11, "DB row tuple must match the 11 INSERT columns")
        _assert_equal(queued[-1], seq, "DB row must carry the live seq")
        response = lc.app.test_client().get(f"/api/package-log/raw?seq={seq}")
        _assert_equal(response.status_code, 200, "raw API must find a live row")
        _assert_equal(response.get_json()["raw_log"], "RAW-ONLY-PAYLOAD full record", "raw API returned the wrong record")

    with tempfile.TemporaryDirectory() as temp_dir:
        legacy_db = os.path.join(temp_dir, "legacy.sqlite3")
        conn = lc.sqlite3.connect(legacy_db)
        conn.execute(
            "CREATE TABLE package_log_entries (id INTEGER PRIMARY KEY AUTOINCREMENT, session_id INTEGER NOT NULL, "
            "created_at REAL NOT NULL, time_display TEXT, device_id TEXT, device_name TEXT, level TEXT, tag TEXT, "
            "message TEXT, raw_log TEXT, is_error INTEGER NOT NULL DEFAULT 0)"
        )
        conn.execute("INSERT INTO package_log_entries (session_id, created_at, message) VALUES (1, 1.0, 'old row')")
        conn.commit()
        conn.close()
        original_path = lc.PACKAGE_LOG_DB_PATH
        try:
            lc.PACKAGE_LOG_DB_PATH = legacy_db
            lc._init_package_log_db()
            lc._init_package_log_db()
        finally:
            lc.PACKAGE_LOG_DB_PATH = original_path
        conn = lc.sqlite3.connect(legacy_db)
        columns = {row[1] for row in conn.execute("PRAGMA table_info(package_log_entries)")}
        kept = conn.execute("SELECT message FROM package_log_entries").fetchall()
        conn.close()
        _assert("seq" in columns, "existing history DB was not migrated with a seq column")
        _assert_equal(kept, [("old row",)], "migration must keep existing history rows")


def _legacy_free_package_db(path: str) -> None:
    original_path = lc.PACKAGE_LOG_DB_PATH
    try:
        lc.PACKAGE_LOG_DB_PATH = path
        lc._init_package_log_db()
    finally:
        lc.PACKAGE_LOG_DB_PATH = original_path


def test_package_log_db_writer_survives_write_errors() -> None:
    """A failing SQLite write must be retried, not end the writer thread."""
    original_path = lc.PACKAGE_LOG_DB_PATH
    original_connect = lc._get_package_db_connection
    original_dropped = lc.package_log_db_dropped_rows
    original_error = lc.package_log_db_last_error
    with tempfile.TemporaryDirectory() as temp_dir:
        db_path = os.path.join(temp_dir, "writer.sqlite3")
        _legacy_free_package_db(db_path)
        failures = {"left": 2, "seen": []}

        def flaky_connect():
            if failures["left"] > 0:
                failures["left"] -= 1
                raise lc.sqlite3.OperationalError("database is locked")
            return original_connect()

        lc._drain_package_log_db_queue()
        lc.package_log_db_writer_stop.clear()
        writer = None
        try:
            lc.PACKAGE_LOG_DB_PATH = db_path
            lc._get_package_db_connection = flaky_connect
            total = lc.PACKAGE_LOG_DB_BATCH_ROWS * 2 + 37
            for index in range(total):
                lc._enqueue_package_log_db_row((1, 1.0, "t", "d", "dn", "I", "Tag", f"m{index}", f"r{index}", 0, index + 1))
            writer = threading.Thread(target=lc._package_log_db_writer, daemon=True)
            writer.start()
            deadline = time.time() + 20
            stored = 0
            while time.time() < deadline:
                if lc.package_log_db_last_error:
                    failures["seen"].append(lc.package_log_db_last_error)
                conn = lc.sqlite3.connect(db_path)
                stored = conn.execute("SELECT COUNT(*) FROM package_log_entries").fetchone()[0]
                conn.close()
                if stored >= total:
                    break
                time.sleep(0.2)
            _assert_equal(stored, total, "writer lost rows after transient SQLite errors")
            _assert(failures["seen"], "writer did not expose the SQLite error to the UI state")
            _assert_equal(lc.package_log_db_last_error, "", "writer did not clear the error after recovering")
            conn = lc.sqlite3.connect(db_path)
            seqs = [row[0] for row in conn.execute("SELECT seq FROM package_log_entries ORDER BY id")]
            conn.close()
            _assert_equal(seqs, list(range(1, total + 1)), "writer changed the row order while retrying")
        finally:
            lc.package_log_db_writer_stop.set()
            if writer is not None:
                writer.join(timeout=5)
            lc.package_log_db_writer_stop.clear()
            lc._get_package_db_connection = original_connect
            lc.PACKAGE_LOG_DB_PATH = original_path
            lc.package_log_db_dropped_rows = original_dropped
            lc.package_log_db_last_error = original_error
            lc._drain_package_log_db_queue()


def test_package_log_older_rows_and_db_status() -> None:
    """Scroll-back pages come from SQLite in seq order; DB problems reach the UI."""
    original_path = lc.PACKAGE_LOG_DB_PATH
    with tempfile.TemporaryDirectory() as temp_dir:
        db_path = os.path.join(temp_dir, "older.sqlite3")
        _legacy_free_package_db(db_path)
        conn = lc.sqlite3.connect(db_path)
        conn.executemany(
            "INSERT INTO package_log_entries (session_id, created_at, time_display, device_id, device_name, level, tag, "
            "message, raw_log, is_error, seq) VALUES (?, 1.0, '10:00:00.000', 'dev', 'Dev', 'I', 'Tag', ?, ?, 0, ?)",
            [(7, f"message {seq}", f"raw {seq}", seq) for seq in range(1, 5001)]
            + [(8, f"other {seq}", f"raw other {seq}", seq) for seq in range(5001, 5101)],
        )
        conn.commit()
        conn.close()
        try:
            lc.PACKAGE_LOG_DB_PATH = db_path
            client = lc.app.test_client()
            page = client.get("/api/package-log/older?before_seq=3001&limit=1000&session_id=7").get_json()
            seqs = [row["seq"] for row in page["rows"]]
            _assert_equal(seqs, list(range(2001, 3001)), "older page must be the rows just before before_seq, oldest first")
            _assert(page["has_more"], "older page must report that more rows exist")
            _assert("raw_log" not in page["rows"][0], "older rows must not carry the raw log")
            last = client.get("/api/package-log/older?before_seq=501&limit=1000&session_id=7").get_json()
            _assert_equal([row["seq"] for row in last["rows"]], list(range(1, 501)), "first page of the session is wrong")
            _assert(not last["has_more"], "the start of the session must end scroll-back")
            other = client.get("/api/package-log/older?before_seq=5101&limit=10&session_id=7").get_json()
            _assert(all(row["message"].startswith("message") for row in other["rows"]), "older rows leaked from another session")
        finally:
            lc.PACKAGE_LOG_DB_PATH = original_path

    with _PackageLogStateGuard() as guard:
        original_dropped = lc.package_log_db_dropped_rows
        original_error = lc.package_log_db_last_error
        original_reported = lc._package_log_db_status_reported
        try:
            lc._package_log_db_status_reported = (lc.package_log_db_dropped_rows, lc.package_log_db_last_error)
            _assert(not lc._emit_package_log_db_status_if_changed(), "unchanged DB status must not be re-sent")
            lc.package_log_db_dropped_rows += 3
            lc.package_log_db_last_error = "OperationalError: database is locked"
            _assert(lc._emit_package_log_db_status_if_changed(), "DB status change was not sent")
            status = guard.events("package_log_db_status")[-1]
            _assert_equal(status["error"], "OperationalError: database is locked", "DB error missing from status")
            _start_package_stream_for_test()
            lc.active_package_log_session_id = 4242
            lc._emit_package_log_reset()
            _assert_equal(guard.events("package_log_reset")[-1]["session_id"], 4242, "reset must name the live session")
        finally:
            lc.package_log_db_dropped_rows = original_dropped
            lc.package_log_db_last_error = original_error
            lc._package_log_db_status_reported = original_reported


class _TempPackageDb:
    """Point Log_checker at a throwaway history DB (and retention file)."""

    def __enter__(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.temp_dir.name, "history.sqlite3")
        self.original = (lc.PACKAGE_LOG_DB_PATH, lc.active_package_log_session_id, lc.target_package_name)
        lc.PACKAGE_LOG_DB_PATH = self.path
        lc.active_package_log_session_id = None
        lc.target_package_name = ""
        return self

    def __exit__(self, *exc):
        lc.PACKAGE_LOG_DB_PATH, lc.active_package_log_session_id, lc.target_package_name = self.original
        self.temp_dir.cleanup()
        return False

    def connect(self):
        return lc.sqlite3.connect(self.path)


def _add_history_session(db, package_id, started_at, rows, status="stopped", with_row_count=True):
    conn = db.connect()
    cur = conn.execute(
        "INSERT INTO package_log_sessions (package_id, started_at, ended_at, status, row_count) VALUES (?, ?, ?, ?, ?)",
        (package_id, started_at, started_at + 60, status, rows if with_row_count else None),
    )
    session_id = cur.lastrowid
    conn.executemany(
        "INSERT INTO package_log_entries (session_id, created_at, message, raw_log, is_error, seq) VALUES (?, ?, ?, ?, 0, ?)",
        [(session_id, started_at + index, f"{package_id} {index}", "x" * 50, index + 1) for index in range(rows)],
    )
    conn.commit()
    conn.close()
    return session_id


def test_package_log_stuck_sessions_and_row_counts() -> None:
    """Startup closes crashed sessions; row counts are backfilled and kept by the writer."""
    with _TempPackageDb() as db:
        conn = db.connect()
        conn.execute(
            "CREATE TABLE package_log_sessions (id INTEGER PRIMARY KEY AUTOINCREMENT, package_id TEXT NOT NULL, "
            "started_at REAL NOT NULL, ended_at REAL, status TEXT NOT NULL DEFAULT 'running')"
        )
        conn.execute(
            "CREATE TABLE package_log_entries (id INTEGER PRIMARY KEY AUTOINCREMENT, session_id INTEGER NOT NULL, "
            "created_at REAL NOT NULL, time_display TEXT, device_id TEXT, device_name TEXT, level TEXT, tag TEXT, "
            "message TEXT, raw_log TEXT, is_error INTEGER NOT NULL DEFAULT 0)"
        )
        conn.execute("INSERT INTO package_log_sessions (package_id, started_at, status) VALUES ('crashed', 100.0, 'running')")
        conn.executemany(
            "INSERT INTO package_log_entries (session_id, created_at, message) VALUES (1, ?, 'row')",
            [(100.0 + index,) for index in range(7)],
        )
        conn.execute("INSERT INTO package_log_sessions (package_id, started_at, status) VALUES ('empty', 200.0, 'running')")
        conn.commit()
        conn.close()

        lc._init_package_log_db()
        conn = db.connect()
        crashed = conn.execute("SELECT status, ended_at, row_count FROM package_log_sessions WHERE id = 1").fetchone()
        empty = conn.execute("SELECT status, ended_at FROM package_log_sessions WHERE id = 2").fetchone()
        has_index = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'index' AND name = 'idx_package_log_entries_session_seq'"
        ).fetchone()
        conn.close()
        _assert_equal(crashed[:2], ("interrupted", 106.0), "crashed session was not closed at its last row")
        _assert_equal(empty, ("interrupted", 200.0), "empty crashed session was not closed at its start")
        _assert(crashed[2] is None, "legacy session must wait for the backfill")
        _assert(has_index, "a small history DB must get the scroll-back index at startup")
        _assert_equal(lc._backfill_package_log_row_counts(), 2, "backfill did not fill the legacy sessions")
        conn = db.connect()
        _assert_equal(conn.execute("SELECT row_count FROM package_log_sessions WHERE id = 1").fetchone()[0], 7, "backfilled row count is wrong")
        conn.close()

        session_id = lc._start_package_log_session("live")
        lc._drain_package_log_db_queue()
        lc.package_log_db_writer_stop.clear()
        writer = threading.Thread(target=lc._package_log_db_writer, daemon=True)
        try:
            for index in range(lc.PACKAGE_LOG_DB_BATCH_ROWS + 3):
                lc._enqueue_package_log_db_row((session_id, 1.0, "t", "d", "dn", "I", "Tag", f"m{index}", f"r{index}", 0, index + 1))
            writer.start()
            deadline = time.time() + 15
            count = 0
            while time.time() < deadline:
                conn = db.connect()
                count = conn.execute("SELECT row_count FROM package_log_sessions WHERE id = ?", (session_id,)).fetchone()[0]
                conn.close()
                if count == lc.PACKAGE_LOG_DB_BATCH_ROWS + 3:
                    break
                time.sleep(0.2)
            _assert_equal(count, lc.PACKAGE_LOG_DB_BATCH_ROWS + 3, "writer did not keep the session row count")
        finally:
            lc.package_log_db_writer_stop.set()
            writer.join(timeout=5)
            lc.package_log_db_writer_stop.clear()
            lc._drain_package_log_db_queue()


def test_package_log_sessions_api_pages() -> None:
    with _TempPackageDb() as db:
        lc._init_package_log_db()
        for index in range(120):
            _add_history_session(db, f"pkg{index}", 1000.0 + index, 1)
        client = lc.app.test_client()
        first = client.get("/api/package-log/sessions?limit=50").get_json()
        _assert_equal((len(first["sessions"]), first["total"], first["has_more"]), (50, 120, True), "first session page is wrong")
        _assert_equal(first["sessions"][0]["package_id"], "pkg119", "sessions must be newest first")
        last = client.get("/api/package-log/sessions?limit=50&offset=100").get_json()
        _assert_equal((len(last["sessions"]), last["has_more"]), (20, False), "last session page is wrong")
        default = client.get("/api/package-log/sessions").get_json()
        _assert_equal(len(default["sessions"]), 120, "the history list must no longer stop at 50 sessions")


def test_package_log_retention_policy() -> None:
    """Retention deletes only old, finished sessions, and only after one confirmation."""
    with _TempPackageDb() as db:
        lc._init_package_log_db()
        now = time.time()
        day = 86400
        old_a = _add_history_session(db, "old-a", now - 30 * day, 40)
        old_b = _add_history_session(db, "old-b", now - 20 * day, 30)
        recent = _add_history_session(db, "recent", now - 2 * day, 20)
        running_old = _add_history_session(db, "running-old", now - 40 * day, 10, status="running")
        live_old = _add_history_session(db, "live-old", now - 50 * day, 10)
        lc.active_package_log_session_id = live_old

        refused = lc._apply_package_log_retention()
        _assert_equal(refused, {"ok": False, "error": "confirmation_required"}, "retention ran without the first confirmation")
        client = lc.app.test_client()
        _assert_equal(client.post("/api/package-log/retention/apply", json={}).status_code, 400, "apply API must require confirm")
        info = client.get("/api/package-log/maintenance").get_json()
        _assert_equal((info["cleanup"]["sessions"], info["cleanup"]["rows"]), (2, 70), "cleanup preview is wrong")
        _assert(not info["policy"]["confirmed"], "policy must start unconfirmed")

        result = client.post("/api/package-log/retention/apply", json={"confirm": True}).get_json()
        _assert_equal((result["deleted_sessions"], result["deleted_rows"]), (2, 70), "retention deleted the wrong amount")
        conn = db.connect()
        remaining = {row[0] for row in conn.execute("SELECT id FROM package_log_sessions")}
        orphan_rows = conn.execute(
            "SELECT COUNT(*) FROM package_log_entries WHERE session_id NOT IN (SELECT id FROM package_log_sessions)"
        ).fetchone()[0]
        conn.close()
        _assert_equal(remaining, {recent, running_old, live_old}, "retention touched a protected or recent session")
        _assert_equal(orphan_rows, 0, "retention left rows of deleted sessions behind")
        _assert(lc._load_package_log_retention()["confirmed"], "the confirmation was not remembered")
        _assert(lc._apply_package_log_retention()["ok"], "confirmed retention must run without asking again")

        original_usage = lc._package_log_db_usage
        try:
            lc._package_log_db_usage = lambda _conn: {"file_bytes": 10 * 1024 ** 3, "used_bytes": 10 * 1024 ** 3, "free_bytes": 0}
            lc.active_package_log_session_id = None
            conn = lc._get_package_db_connection()
            try:
                plan = lc._package_log_retention_plan(conn, lc._load_package_log_retention())
            finally:
                conn.close()
            reasons = [(item["id"], item["reason"]) for item in plan["sessions"]]
            _assert(all(item_id != running_old for item_id, _ in reasons), "size policy chose a running session")
            _assert_equal(reasons[0], (live_old, "age"), "an old finished session must go first")
            _assert(any(reason == "size" for _, reason in reasons), "over-size history did not select sessions by size")
        finally:
            lc._package_log_db_usage = original_usage


def test_package_log_compact_api() -> None:
    with _TempPackageDb() as db:
        lc._init_package_log_db()
        conn = db.connect()
        conn.execute("DROP INDEX IF EXISTS idx_package_log_entries_session_seq")
        conn.commit()
        conn.close()
        _add_history_session(db, "compact", time.time(), 50)
        client = lc.app.test_client()
        _assert_equal(client.post("/api/package-log/compact", json={}).status_code, 400, "compact must require confirm")
        lc.target_package_name = "com.capturing.now"
        _assert_equal(client.post("/api/package-log/compact", json={"confirm": True}).status_code, 409, "compact must refuse while capturing")
        lc.target_package_name = ""
        response = client.post("/api/package-log/compact", json={"confirm": True})
        _assert_equal(response.status_code, 200, "compact failed on an idle history")
        conn = db.connect()
        has_index = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'index' AND name = 'idx_package_log_entries_session_seq'"
        ).fetchone()
        rows = conn.execute("SELECT COUNT(*) FROM package_log_entries").fetchone()[0]
        conn.close()
        _assert(has_index, "compact did not build the scroll-back index")
        _assert_equal(rows, 50, "compact changed the stored rows")


class _SpecificStateGuard:
    """Snapshot and restore Specific Validator globals; capture its emits."""

    NAMES = (
        "specific_event_results", "specific_event_pending", "specific_event_seq",
        "specific_event_name_filters", "specific_event_params_filters", "is_paused", "active_platform",
    )

    def __enter__(self):
        self.values = {name: getattr(lc, name) for name in self.NAMES}
        self.cache = list(lc.event_log_cache)
        self.dedup_seen = set(lc.record_dedup_seen["specific"])
        self.dedup_order = list(lc.record_dedup_order["specific"])
        self.emit = lc.socketio.emit
        self.emitted = []
        lc.socketio.emit = lambda event, payload=None, **kw: self.emitted.append((event, payload, kw.get("to")))
        lc.event_log_cache.clear()
        lc._clear_record_dedup("specific")
        lc.specific_event_results = []
        lc.specific_event_pending = []
        lc.is_paused = False
        lc.active_platform = "android"
        return self

    def __exit__(self, *exc):
        lc.socketio.emit = self.emit
        for name, value in self.values.items():
            setattr(lc, name, value)
        lc.event_log_cache.clear()
        lc.event_log_cache.extend(self.cache)
        lc.record_dedup_seen["specific"].clear()
        lc.record_dedup_seen["specific"].update(self.dedup_seen)
        lc.record_dedup_order["specific"].clear()
        lc.record_dedup_order["specific"].extend(self.dedup_order)
        return False

    def events(self, name):
        return [payload for event, payload, _to in self.emitted if event == name]


def _feed_specific_events(count, start=0):
    params = {f"param_{k}": f"value_{k}" for k in range(8)}
    for index in range(start, start + count):
        name = "ad_load" if index % 3 else "session_start_customs"
        line = (
            f"Oct 20 16:43:{index % 60:02d} iPhone-12 TrainIdle(UnityFramework)[5706] <Notice>: [Tracking] "
            "TrackingService->Track: " + json.dumps({"EventName": name, "params": dict(params, seq=index)})
        )
        event_name, event_params, json_string = lc.find_and_parse_event(line)
        lc.cache_specific_event_log(event_name, event_params, json_string, line, "ios-device")


def test_specific_validator_delta_matches_full_rebuild() -> None:
    """Rows sent as deltas must be exactly the rows a full rebuild produces.

    Same validation, same formatted JSON HTML and the untouched raw log, so
    the parsed-JSON column can still be compared with the raw log column.
    """
    with _SpecificStateGuard() as guard:
        lc.specific_event_name_filters = ["ad_", "session_"]
        lc.specific_event_params_filters = []
        total = lc.event_log_cache.maxlen + 120
        delta_rows = []
        min_seq = 0
        for chunk_start in range(0, total, 97):
            _feed_specific_events(min(97, total - chunk_start), start=chunk_start)
            _assert(not guard.events("update_specific_event_table"), "a new event must not resend the whole table")
            lc._flush_specific_event_delta()
        for payload in guard.events("specific_event_delta"):
            delta_rows.extend(payload["rows"])
            min_seq = payload["min_seq"]
        kept_delta_rows = [row for row in delta_rows if row["seq"] >= min_seq]

        guard.emitted.clear()
        lc._apply_specific_filter_and_emit()
        full_rows = guard.events("update_specific_event_table")[-1]
        _assert_equal(len(full_rows), lc.event_log_cache.maxlen, "full rebuild row count changed")
        _assert_equal(kept_delta_rows, full_rows, "delta rows differ from the full rebuild rows")
        _assert_equal(lc.specific_event_results, full_rows, "incremental server table differs from a rebuild")
        _assert_equal(min_seq, lc.event_log_cache[0]["seq"], "delta must report the oldest cached row")
        raw_by_seq = {item["seq"]: item["log"] for item in lc.event_log_cache}
        _assert(all(row["raw_log"] == raw_by_seq[row["seq"]] for row in full_rows), "raw log must stay byte-identical")
        _assert(all("json-" in row["details"] or "<" in row["details"] for row in full_rows), "formatted JSON column missing")


def test_specific_validator_filters_clear_and_snapshot() -> None:
    with _SpecificStateGuard() as guard:
        lc.specific_event_name_filters = ["ad_"]
        _feed_specific_events(30)
        _assert(all(row["event_name"].startswith("ad_") for row in lc.specific_event_pending), "name filter not applied to new rows")
        lc.specific_event_name_filters = ["session_"]
        lc._apply_specific_filter_and_emit()
        _assert_equal(lc.specific_event_pending, [], "a full rebuild must absorb pending rows")
        full = guard.events("update_specific_event_table")[-1]
        _assert(full and all(row["event_name"].startswith("session_") for row in full), "filter change did not rebuild")
        _assert_equal(lc._flush_specific_event_delta(), 0, "rows inside the rebuild were sent again")

        guard.emitted.clear()
        with lc.app.test_request_context("/"):
            lc.request.sid = "harness-sid"
            lc.refresh()
        snapshot = [(payload, to) for event, payload, to in guard.emitted if event == "update_specific_event_table"]
        _assert(snapshot and snapshot[-1][1] == "harness-sid", "refresh must send the current table to that client")
        _assert_equal(snapshot[-1][0], lc.specific_event_results, "refresh snapshot differs from the server table")

        _feed_specific_events(5, start=100)
        lc.cl() if hasattr(lc, "cl") else None
        _assert_equal((lc.specific_event_pending, list(lc.event_log_cache)), ([], []), "Clear must drop pending rows")


def test_specific_validator_delta_ui_contract() -> None:
    html = lc.app.test_client().get("/").get_data(as_text=True)
    for needle in (
        "function createLiveTable(opts)",
        "tbody.innerHTML = filtered.map(opts.rowHtml).join('');",
        "tbody.insertAdjacentHTML('beforeend', appendRows.map(opts.rowHtml).join(''));",
        "var LIVE_TABLE_LARGE_INTERVAL_MS = 1000;",
        "function specificRowHtml(res)",
        "socket.on('specific_event_delta', (payload) => specificLive.applyDelta(payload));",
        "if (tabName === 'Specific' && specificLive) specificLive.onTabShown();",
    ):
        _assert(needle in html, f"Specific Validator delta UI contract missing: {needle}")
    row_template = html.split("function specificRowHtml(res) {", 1)[1].split("`;", 1)[0]
    for column in ('class="py-2 px-3 details-cell text-sm">${res.details}</td>',
                   "${escapeHTML(res.raw_log || '')}",
                   "data-json='${escapeAttribute(res.json_data)}'>View JSON</button>"):
        _assert(column in row_template, f"Specific row template lost a column: {column}")


class _ValidatorStateGuard:
    NAMES = ("validator_pending", "validator_seq", "validator_active", "is_paused", "active_platform",
             "default_params", "required_params")

    def __enter__(self):
        self.values = {name: getattr(lc, name) for name in self.NAMES}
        self.rows = list(lc.validator_results)
        self.dedup_seen = set(lc.record_dedup_seen["validator"])
        self.dedup_order = list(lc.record_dedup_order["validator"])
        self.emit = lc.socketio.emit
        self.emitted = []
        lc.socketio.emit = lambda event, payload=None, **kw: self.emitted.append((event, payload, kw.get("to")))
        lc.validator_results.clear()
        lc._clear_record_dedup("validator")
        lc.validator_pending = []
        lc.validator_active = True
        lc.is_paused = False
        lc.active_platform = "android"
        return self

    def __exit__(self, *exc):
        lc.socketio.emit = self.emit
        for name, value in self.values.items():
            setattr(lc, name, value)
        lc.validator_results.clear()
        lc.validator_results.extend(self.rows)
        lc.record_dedup_seen["validator"].clear()
        lc.record_dedup_seen["validator"].update(self.dedup_seen)
        lc.record_dedup_order["validator"].clear()
        lc.record_dedup_order["validator"].extend(self.dedup_order)
        return False

    def events(self, name):
        return [payload for event, payload, _to in self.emitted if event == name]


def test_validator_delta_matches_server_table() -> None:
    """Validator rows reach the UI once, unchanged, in order; old rows are reported by min_seq."""
    with _ValidatorStateGuard() as guard:
        lc.default_params = ["param_0", "missing_param"]
        total = lc.validator_results.maxlen + 150
        sent = []
        min_seq = 0
        params = {f"param_{k}": f"value_{k}" for k in range(8)}
        for index in range(total):
            line = (
                f"Oct 20 16:43:{index % 60:02d} iPhone-12 TrainIdle(UnityFramework)[5706] <Notice>: [Tracking] "
                "TrackingService->Track: " + json.dumps({"EventName": "session_start_customs", "params": dict(params, seq=index)})
            )
            event_name, event_params, json_string = lc.find_and_parse_event(line)
            lc.process_event_validator_log(event_name, event_params, json_string, line, "ios-device")
            if index % 113 == 0:
                lc._flush_validator_delta()
        lc._flush_validator_delta()
        _assert(not guard.events("update_validator_table"), "a new event must not resend the whole Validator table")
        for payload in guard.events("validator_delta"):
            sent.extend(payload["rows"])
            min_seq = payload["min_seq"]
        kept = [row for row in sent if row["seq"] >= min_seq]
        _assert_equal(kept, list(lc.validator_results), "Validator deltas differ from the server table")
        _assert_equal(min_seq, lc.validator_results[0]["seq"], "delta must report the oldest kept row")
        _assert(all("Missing" in row["details"] and row["status"] == "FAILED" for row in kept), "validation result changed")
        _assert(all(row["raw_log"].startswith("Oct 20 16:43:") for row in kept), "raw log column changed")

        guard.emitted.clear()
        with lc.app.test_request_context("/"):
            lc.request.sid = "harness-sid"
            lc.refresh()
        snapshot = [(payload, to) for event, payload, to in guard.emitted if event == "update_validator_table"]
        _assert(snapshot and snapshot[-1][1] == "harness-sid", "refresh must send the Validator table to that client")
        _assert_equal(snapshot[-1][0], list(lc.validator_results), "refresh snapshot differs from the server table")

        lc.validator_pending.append({"seq": -1})
        lc.cl()
        _assert_equal((lc.validator_pending, list(lc.validator_results)), ([], []), "Clear must drop pending Validator rows")


def test_validator_delta_ui_contract() -> None:
    html = lc.app.test_client().get("/").get_data(as_text=True)
    for needle in (
        "function validatorRowHtml(res)",
        "function renderValidatorTable(results) {",
        "validatorLive.renderFull(results);",
        "if (validatorLive.applyDelta(payload)) updateDefaultEventStatus(validator_results_cache);",
        "validatorLive.setAll(validator_results_cache);",
        "if (tabName === 'Validator' && validatorLive) validatorLive.onTabShown();",
        "if (el.textContent !== text) el.textContent = text;",
        "function removeLeadingRows(tbody, count) {",
    ):
        _assert(needle in html, f"Validator delta UI contract missing: {needle}")
    row_template = html.split("function validatorRowHtml(res) {", 1)[1].split("`;", 1)[0]
    for column in ("${res.details}", "${escapeHTML(res.raw_log || '')}", "View JSON"):
        _assert(column in row_template, f"Validator row template lost a column: {column}")


class _LiveTablesGuard:
    """Snapshot/restore the Revenue and live-table globals; capture emits."""

    def __enter__(self):
        self.tables = {name: list(rows) for name, rows in lc.LIVE_TABLE_ROWS.items()}
        self.adrevenue = list(lc.adrevenue_logs)
        self.values = {name: getattr(lc, name) for name in (
            "adrevenue_rendered", "adrevenue_pending", "adrevenue_seq", "live_table_seq",
            "adrevenue_default_params", "adrevenue_source_params", "active_platform", "is_paused")}
        self.pending = {name: list(rows) for name, rows in lc.LIVE_TABLE_PENDING.items()}
        self.dedup = {name: (set(lc.record_dedup_seen[name]), list(lc.record_dedup_order[name])) for name in ("adrevenue",)}
        self.emit = lc.socketio.emit
        self.emitted = []
        lc.socketio.emit = lambda event, payload=None, **kw: self.emitted.append((event, payload, kw.get("to")))
        for rows in lc.LIVE_TABLE_ROWS.values():
            rows.clear()
        lc.adrevenue_logs.clear()
        lc.adrevenue_rendered = []
        lc.adrevenue_pending = []
        lc._clear_live_table_pending()
        lc._clear_record_dedup("adrevenue")
        lc.active_platform = "android"
        lc.is_paused = False
        return self

    def __exit__(self, *exc):
        lc.socketio.emit = self.emit
        for name, rows in self.tables.items():
            lc.LIVE_TABLE_ROWS[name].clear()
            lc.LIVE_TABLE_ROWS[name].extend(rows)
        lc.adrevenue_logs.clear()
        lc.adrevenue_logs.extend(self.adrevenue)
        for name, value in self.values.items():
            setattr(lc, name, value)
        for name, rows in self.pending.items():
            lc.LIVE_TABLE_PENDING[name] = rows
        for name, (seen, order) in self.dedup.items():
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_seen[name].update(seen)
            lc.record_dedup_order[name].clear()
            lc.record_dedup_order[name].extend(order)
        return False

    def events(self, name):
        return [payload for event, payload, _to in self.emitted if event == name]


def test_revenue_delta_matches_full_rebuild() -> None:
    """Revenue rows built one at a time equal a full rebuild, row for row."""
    with _LiveTablesGuard() as guard:
        lc.adrevenue_default_params = ["adRevenue", "currency", "missing_param"]
        lc.adrevenue_source_params = {"appsflyer": ["custom_parameters"]}
        total = lc.adrevenue_logs.maxlen + 75
        for index in range(total):
            kind = index % 4
            raw = f"Oct 21 09:00:{index % 60:02d} iPhone-12 TrainIdle[5706] <Notice>: revenue row {index}"
            if kind == 0:
                lc._record_adrevenue_log("ios-device", "appmetrica", "ad_revenue",
                                         {"adRevenue": index, "currency": "USD", "payload": {"x": index}}, raw)
            elif kind == 1:
                lc._record_adrevenue_log("ios-device", "appsflyer", "af_ad_revenue",
                                         {"ad_network": {"payload": {"custom_parameters": {"k": index}}}}, raw)
            elif kind == 2:
                lc._record_adrevenue_log("ios-device", "adjust", "Adjust adrevenue", {"revenue": index}, raw,
                                         json.dumps({"revenue": index}), "response", skip_validation=True)
            else:
                lc._record_adrevenue_log("ios-device", "adjust", "Adjust Config",
                                         {"IsSendingInBackgroundEnabled": True, "IsCostDataInAttributionEnabled": index % 8 == 3,
                                          "IsAdServicesEnabled": True, "IsIdfaReadingEnabled": True}, raw, skip_validation=True)
            if index % 211 == 0:
                lc._flush_adrevenue_delta()
        lc._flush_adrevenue_delta()
        _assert(not guard.events("update_adrevenue_table"), "a new Revenue row must not resend the whole table")
        sent = [row for payload in guard.events("adrevenue_delta") for row in payload["rows"]]
        min_seq = guard.events("adrevenue_delta")[-1]["min_seq"]
        kept = [row for row in sent if row["seq"] >= min_seq]
        guard.emitted.clear()
        lc._apply_adrevenue_filter_and_emit()
        full = guard.events("update_adrevenue_table")[-1]
        _assert_equal(len(full), lc.adrevenue_logs.maxlen, "Revenue rebuild row count changed")
        _assert_equal(kept, full, "Revenue delta rows differ from a full rebuild")
        _assert_equal(lc.adrevenue_rendered, full, "server Revenue table differs from a full rebuild")
        _assert_equal({row["status"] for row in full}, {"PASSED", "FAILED", "INFO"}, "Revenue statuses not exercised")


def test_live_tables_delta_snapshot_and_clear() -> None:
    """Load Ads, Load Ads Ext, Callback and Price Rotation send new rows only, unchanged and in order."""
    with _LiveTablesGuard() as guard:
        for name, rows in lc.LIVE_TABLE_ROWS.items():
            total = rows.maxlen + 40
            for index in range(total):
                lc._live_table_append(name, {"device_id": "dev", "device_name": "Dev", "raw_log": f"{name} raw {index}",
                                             "type": "Ad Event", "event_name": "onAdLoadFailed", "ad_format": "interstitial"})
                if index % 157 == 0:
                    lc._flush_live_table(name)
            lc._flush_live_table(name)
            deltas = guard.events(lc.LIVE_TABLE_DELTA_EVENTS[name])
            sent = [row for payload in deltas for row in payload["rows"]]
            kept = [row for row in sent if row["seq"] >= deltas[-1]["min_seq"]]
            _assert_equal(kept, list(rows), f"{name} deltas differ from the server table")
            _assert_equal(len(sent), total, f"{name} lost or repeated a row")
        _assert(not any(event.startswith("update_") for event, _p, _t in guard.emitted), "a new row must not resend a whole table")

        guard.emitted.clear()
        with lc.app.test_request_context("/"):
            lc.request.sid = "harness-sid"
            lc.refresh()
        snapshots = {event: (payload, to) for event, payload, to in guard.emitted
                     if event in ("update_load_ads", "update_load_ads_ext", "update_callback_ad_table",
                                  "update_price_rotation_table", "update_adrevenue_table") and to == "harness-sid"}
        _assert_equal(len(snapshots), 5, "refresh must send every live table to that client")
        _assert_equal(snapshots["update_callback_ad_table"][0], list(lc.callback_ad_logs), "callback snapshot differs")

        lc._live_table_append("callback", {"device_id": "dev", "raw_log": "pending row"})
        lc.cl()
        _assert(not any(lc.LIVE_TABLE_PENDING.values()), "Clear must drop pending live-table rows")


def test_live_tables_ui_contract() -> None:
    html = lc.app.test_client().get("/").get_data(as_text=True)
    for needle in (
        "var adRevenueLive = createLiveTable({",
        "var callbackLive = createLiveTable({",
        "var priceRotationLive = createLiveTable({",
        "var loadAdsLive = simpleTableLive('loadAdsTableBody', 'tabContentLoadAds'",
        "var loadAdsExtLive = simpleTableLive('loadAdsExtTableBody', 'tabContentLoadAdsExt'",
        "socket.on('adrevenue_delta', (payload) => adRevenueLive.applyDelta(payload));",
        "socket.on('callback_ad_delta', (payload) => callbackLive.applyDelta(payload));",
        "socket.on('price_rotation_delta', (payload) => priceRotationLive.applyDelta(payload));",
        "socket.on('load_ads_delta', (payload) => loadAdsLive.applyDelta(payload));",
        "socket.on('load_ads_ext_delta', (payload) => loadAdsExtLive.applyDelta(payload));",
        "if (tabName === 'AdRevenue' && adRevenueLive) adRevenueLive.onTabShown();",
        "if (tabName === 'LoadAdsExt' && loadAdsExtLive) loadAdsExtLive.onTabShown();",
        "const visible = opts.makeFilter ? opts.makeFilter() : opts.visible;",
    ):
        _assert(needle in html, f"live-table UI contract missing: {needle}")
    for fn, columns in (
        ("function adRevenueRowHtml(res) {", ("${res.details}", "escapeHTML(res.raw_log", "View JSON")),
        ("function callbackRowHtml(res) {", ("escapeHTML(res.raw_log", "View JSON")),
        ("function priceRotationRowHtml(res) {", ("escapeHTML(res.details || '')", "escapeHTML(res.raw_log || '')")),
    ):
        template = html.split(fn, 1)[1].split("\n        }", 1)[0]
        for column in columns:
            _assert(column in template, f"{fn} lost a column: {column}")


def test_repository_moved_to_indiezvn() -> None:
    """Update, payload and preset URLs use the company repo first and keep trucbm as fallback."""
    import remote_update as updater
    import importlib.util as _ilu
    spec = _ilu.spec_from_file_location("sync_update_manifest", ROOT / "tools" / "sync_update_manifest.py")
    sync = _ilu.module_from_spec(spec)
    spec.loader.exec_module(sync)
    config = json.loads((ROOT / "remote_update_config_v250.json").read_text(encoding="utf-8"))
    services = (ROOT / "services_checker" / "app.py").read_text(encoding="utf-8", errors="ignore")
    groups = {
        "remote_update.DEFAULT_MANIFEST_URLS": list(updater.DEFAULT_MANIFEST_URLS),
        "remote_update.DEFAULT_FILE_URL_BASES": list(updater.DEFAULT_FILE_URL_BASES),
        "remote_update_config_v250.json": [config["manifest_url"]] + list(config["manifest_urls"]),
        "Log_checker.DEFAULT_REMOTE_MANIFEST_URLS": [lc.DEFAULT_REMOTE_MANIFEST_URL] + list(lc.DEFAULT_REMOTE_MANIFEST_URLS),
        "Log_checker.SDK_CHECK_PRESETS_REMOTE_URLS": list(lc.SDK_CHECK_PRESETS_REMOTE_URLS),
        "Log_checker.PACKAGE_LOG_PRESETS_REMOTE_URLS": list(lc.PACKAGE_LOG_PRESETS_REMOTE_URLS),
        "sync_update_manifest.PAYLOAD_URL_BASES": list(sync.PAYLOAD_URL_BASES),
    }
    for name, urls in groups.items():
        _assert("/indiezvn/Eventchecker" in urls[0], f"{name} must try the indiezvn repo first: {urls[0]}")
        owners = ["indiezvn" if "/indiezvn/" in u else "trucbm" if "/trucbm/" in u else "?" for u in urls]
        _assert("?" not in owners, f"{name} has an unknown repository: {urls}")
        _assert("trucbm" in owners, f"{name} must keep a trucbm fallback")
        first_old = owners.index("trucbm")
        github_new = [u for u in urls if "/indiezvn/" in u and "jsdelivr" not in u]
        _assert(all(urls.index(u) < first_old for u in github_new), f"{name}: every GitHub indiezvn URL must come before trucbm")
    _assert('SERVICES_CHECKER_GITHUB_REPO = "indiezvn/Eventchecker"' in services, "Services Checker presets must read the indiezvn repo")
    for doc in ("README_SETUP_TEAM.md", "README_SETUP_MAC.md", "README_SETUP_WINDOWS.md", "CODEX_SETUP_CHECKLIST.md"):
        _assert("github.com/trucbm/Eventchecker" not in (ROOT / doc).read_text(encoding="utf-8"), f"{doc} still clones the old repository")

    # If the company repo cannot be reached, the next (trucbm) URL is used.
    calls = []
    original_download = updater._download
    original_delays = updater.DOWNLOAD_RETRY_DELAYS
    try:
        def fake_download(url, timeout):
            calls.append(url)
            if "/indiezvn/" in url:
                raise OSError("indiezvn unreachable")
            return b'{"version": "ok"}'
        updater._download = fake_download
        updater.DOWNLOAD_RETRY_DELAYS = (0, 0)
        data, used = updater._download_first(updater.DEFAULT_MANIFEST_URLS, 5)
        _assert("/trucbm/" in used and data, f"updater did not fall back to trucbm: {used}")
        _assert("/indiezvn/" in calls[0], "updater must try indiezvn first")
    finally:
        updater._download = original_download
        updater.DOWNLOAD_RETRY_DELAYS = original_delays


def test_sdk_check_results_text_size() -> None:
    """SDK Check results use the same small text as the other result tables."""
    text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert('<pre class="font-mono ${sizeClass} leading-5" style="margin: 0; white-space: pre-wrap;">' in text, "SDK Check rows must have an explicit small text size")
    _assert("const sizeClass = res.status === 'LABEL' ? 'text-sm' : 'text-xs';" in text, "SDK Check label/row sizes changed")
    _assert('<pre style="font-family: monospace; margin: 0;' not in text, "unsized monospace <pre> is back in SDK Check")


def test_full_window_layout() -> None:
    """The app shell spans the whole window and the visible tab fills the height below the tab bar."""
    text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert('<div class="container mx-auto' not in text, "the width-capped container is back")
    _assert('<div id="appShell" class="w-full max-w-none p-3 sm:p-4 lg:p-6 h-full flex flex-col">' in text, "app shell must be full width")
    _assert('<div id="tabScrollArea" class="flex-grow overflow-auto">' in text, "tab area id missing")
    _assert('#tabScrollArea > [id^="tabContent"]:not(.hidden) { flex: 1 0 auto; display: flex; flex-direction: column; }' in text, "visible tab must fill the height")
    _assert('#tabScrollArea > [id^="tabContent"]:not(.hidden) > div:last-child { flex: 1 0 auto; }' in text, "last card of the visible tab must stretch")


def test_specific_filter_boxes_compact() -> None:
    """Specific Validator filter boxes have a fixed height, level with the last Source Filter option."""
    text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    for id_ in ("specificEventInput", "specificParamInput"):
        _assert(f'<textarea id="{id_}" class="w-full h-[4.25rem] resize-none p-2 border rounded-md shadow-sm"' in text, f"{id_} must keep the compact fixed height")


def test_ui_zoom_setting_contract() -> None:
    """Text zoom is saved server-side and rendered into the page; the default is 85%."""
    original_data_dir = os.environ.get("EVENTINSPECTOR_DATA_DIR")
    with tempfile.TemporaryDirectory() as temp_dir:
        os.environ["EVENTINSPECTOR_DATA_DIR"] = temp_dir
        try:
            settings_path = os.path.join(temp_dir, "ui_settings.json")
            client = lc.app.test_client()
            html = client.get("/").get_data(as_text=True)
            _assert_equal(lc.UI_DEFAULT_ZOOM, 0.85, "default text size must be 85%")
            _assert('data-ui-zoom="0.85" style="font-size: 13.6px;"' in html, "a new machine must start at the 85% default")
            _assert('title="Reset to the default size (Ctrl/Cmd 0)">85%</button>' in html, "the reset button must show the current size")
            for needle in ('id="uiZoomOut"', 'id="uiZoomReset"', 'id="uiZoomIn"', "function stepUiZoom(direction)",
                           "fetch('/api/ui-settings'", "if (e.key === '-' || e.key === '_')",
                           "var UI_DEFAULT_ZOOM = 0.85;", "applyUiZoom(UI_DEFAULT_ZOOM)"):
                _assert(needle in html, f"zoom UI contract missing: {needle}")
            _assert_equal(client.post("/api/ui-settings", json={"zoom": 0.33}).status_code, 400, "zoom outside the levels must be refused")
            _assert_equal(client.post("/api/ui-settings", json={"zoom": "x"}).status_code, 400, "non-numeric zoom must be refused")
            _assert('<div id="appHeader" class="bg-white rounded-xl shadow-md p-4 mb-4 flex-shrink-0">' in html, "header must be visible by default")
            _assert('>Hide Top</button>' in html and "function applyTopHidden(hidden)" in html, "Hide Top toggle is missing")
            # Hide Top alone must not pin the zoom: a later default must still apply.
            _assert_equal(client.post("/api/ui-settings", json={"hide_top": True}).get_json(), {"ok": True, "zoom": 0.85, "hide_top": True}, "hide_top was not saved")
            with open(settings_path, encoding="utf-8") as handle:
                _assert_equal(json.load(handle), {"hide_top": True}, "only the values the user changed may be stored")
            _assert_equal(client.post("/api/ui-settings", json={"zoom": 1.1}).get_json(), {"ok": True, "zoom": 1.1, "hide_top": True}, "zoom must save without resetting hide_top")
            _assert_equal(client.post("/api/ui-settings", json={"hide_top": "yes"}).status_code, 400, "non-boolean hide_top must be refused")
            _assert_equal(client.post("/api/ui-settings", json={}).status_code, 400, "an empty settings update must be refused")
            html = client.get("/").get_data(as_text=True)
            _assert('data-ui-zoom="1.1" style="font-size: 17.6px;"' in html, "saved zoom is not rendered into the page")
            _assert('flex-shrink-0 hidden">' in html and '>Show Top</button>' in html, "saved hide_top is not rendered into the page")
            with open(settings_path, "w", encoding="utf-8") as handle:
                handle.write("{broken")
            html = client.get("/").get_data(as_text=True)
            _assert('data-ui-zoom="0.85"' in html, "a broken settings file must fall back to the default")
        finally:
            if original_data_dir is None:
                os.environ.pop("EVENTINSPECTOR_DATA_DIR", None)
            else:
                os.environ["EVENTINSPECTOR_DATA_DIR"] = original_data_dir


def test_package_log_delta_ui_contract() -> None:
    html = lc.app.test_client().get("/").get_data(as_text=True)
    for needle in (
        "socket.on('package_log_delta'",
        "socket.on('package_log_reset'",
        "socket.emit('package_log_sync'",
        "if (l && l.seq) return 'seq:' + l.seq;",
    ):
        _assert(needle in html, f"Package Log delta UI contract missing: {needle}")
    _assert("socket.on('package_log_cache'" not in html, "the full-snapshot Package Log listener must be gone")
    for needle in (
        "const PACKAGE_STORE_MAX_ROWS = 200000;",
        'class="package-log-spacer"',
        "function renderPackageViewport(",
        "function packageRowCellTexts(row)",
        "for (let i = minIdx; i <= maxIdx; i++) selectedPackageRowKeys.add(getPackageRowKey(packageView[i]));",
        'id="packageNewRowsBtn"',
        "#packageLogTableBody tr.package-log-row td {",
        "socket.on('package_log_db_status'",
        "function loadOlderPackageRows(",
        "if (scrolledUp && container.scrollTop < packageRowHeight * 30) loadOlderPackageRows();",
        "setPackageOlderStatus(packageOlderExhausted && packageHistoryPrefix.length ? 'Đã tới đầu session' : '', 2500);",
        "if (text && autoHideMs > 0) packageOlderStatusTimer = setTimeout(() => setPackageOlderStatus(''), autoHideMs);",
        "#packageLogTableBody tr.package-log-row { height: 18px; }",
        "function clearPackageLiveSelection()",
        'id="clearPackageFiltersBtn"',
        "document.getElementById('clearPackageFiltersBtn')?.addEventListener('click', () => {",
        "if (target.closest('button, input, select, textarea, label, a, [role=\"button\"], .resizer')) return;",
        "if (filterChanged && selectedPackageRowKeys.size && revealSelectedPackageRow()) return;",
        "function revealSelectedPackageRow()",
        "window.addEventListener('blur', endPackageRowDrags);",
        "if (e.buttons === 0) {",
        "let packageRowHeight = 18;",
        "function releasePackageHistoryPrefix(",
        'id="searchPackageSessionBtn"',
        'id="packageDbWarning"',
        'id="packageHistoryMaintenance"',
        "fetch('/api/package-log/maintenance')",
        "fetch(`/api/package-log/sessions?limit=500&platform=${packageHistoryPlatform}`)",
        "fetch('/api/package-log/retention/apply'",
        "fetch('/api/package-log/compact'",
    ):
        _assert(needle in html, f"virtualized Package Log contract missing: {needle}")
    for removed in ("maxDomRows", "function getPackageSourceLogs", "packageUiMaxRows"):
        _assert(removed not in html, f"the fixed-size Package Log DOM window is back: {removed}")


def test_main_script_has_unique_top_level_names() -> None:
    """A duplicate top-level `function`/`let`/`const` aborts the whole UI script.

    The harness does not execute JavaScript, so guard the cheapest fatal
    mistake directly: one name declared twice at the main script's top level.
    """
    html = lc.app.test_client().get("/").get_data(as_text=True)
    names = re.findall(r"^ {8}(?:async\s+)?(?:function\s*\*?\s*|let\s+|const\s+|var\s+|class\s+)([A-Za-z_$][\w$]*)", html, re.M)
    duplicates = sorted({name for name in names if names.count(name) > 1})
    _assert(not duplicates, f"main UI script declares these names more than once: {duplicates}")


def test_android_package_log_ignores_global_pause() -> None:
    """Package Log's own pause control must not be disabled by global Pause."""
    original_platform = lc.active_platform
    original_paused = lc.is_paused
    original_target = lc.target_package_name
    original_rows = list(lc.package_log_cache)
    original_dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
    original_dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
    try:
        lc.active_platform = "android"
        lc.is_paused = True
        lc.target_package_name = "com.nostel.parking.car"
        lc.package_log_cache.clear()
        lc._clear_record_dedup("package")
        package_line = (
            '09-28 14:43:16.087  6927  7783 I Unity : '
            '[Tracking] TrackingService->Track: {"EventName":"ad_impression"}\n'
        )
        fake_process = type("FakeProcess", (), {"stdout": io.StringIO(package_line)})()
        lc.package_log_consumer("android-device", fake_process)
        _assert_equal(len(lc.package_log_cache), 1, "Android Package Log stopped by global Pause")
    finally:
        lc.active_platform = original_platform
        lc.is_paused = original_paused
        lc.target_package_name = original_target
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_rows)
        for name in lc.record_dedup_seen:
            lc.record_dedup_seen[name].clear()
            lc.record_dedup_order[name].clear()
        for name, values in original_dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(original_dedup_order[name])


def test_ios_package_log_preserves_multiline_records() -> None:
    """Package Log keeps iOS continuation lines while other tabs stay line-based."""
    original_platform = lc.active_platform
    original_target = lc.target_package_name
    original_paused = lc.is_paused
    original_cache = list(lc.package_log_cache)
    original_buffers = dict(lc.active_ios_package_log_buffers)
    try:
        lc.active_platform = "ios"
        lc.target_package_name = "PixelArt"
        lc.is_paused = False
        lc.package_log_cache.clear()
        lc._clear_record_dedup("package")
        lc.active_ios_package_log_buffers.clear()

        record_lines = [
            (
                "Sep 18 17:10:59 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
                "[AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting signal_collection_success "
                "with extra parameters {"
            ),
            '    "adapter_version" = "8.5.0.1";',
            '    "network_name" = "YANDEX_BIDDING";',
            "}",
        ]
        for line in record_lines:
            log_obj = lc._normalize_ios_log_line(line, "ios-device")
            lc.process_ios_package_log_stream_line("ios-device", log_obj)

        _assert_equal(len(lc.package_log_cache), 0, "iOS package record should wait for its record boundary")

        next_record = (
            "Sep 18 17:11:00 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
            "next record"
        )
        lc.process_ios_package_log_stream_line(
            "ios-device",
            lc._normalize_ios_log_line(next_record, "ios-device"),
        )
        _assert_equal(len(lc.package_log_cache), 1, "completed iOS package record was not emitted")
        row = lc.package_log_cache[-1]
        _assert('"adapter_version" = "8.5.0.1";' in row["log"], "adapter_version continuation was dropped")
        _assert('"network_name" = "YANDEX_BIDDING";' in row["message"], "network_name continuation was dropped")
        _assert_equal(row["tag"], "PixelArt", "iOS package process tag was not preserved")

        lc._flush_ios_package_log_buffer("ios-device")
        _assert_equal(len(lc.package_log_cache), 2, "last iOS package record was not flushed")
    finally:
        lc.active_platform = original_platform
        lc.target_package_name = original_target
        lc.is_paused = original_paused
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_cache)
        lc.active_ios_package_log_buffers.clear()
        lc.active_ios_package_log_buffers.update(original_buffers)


def test_ios_package_log_tidevice_nul_framing() -> None:
    """Package Log consumes tidevice's NUL-delimited multiline records intact."""
    original_platform = lc.active_platform
    original_target = lc.target_package_name
    original_paused = lc.is_paused
    original_cache = list(lc.package_log_cache)
    original_buffers = dict(lc.active_ios_package_log_frame_buffers)
    try:
        lc.active_platform = "ios"
        lc.target_package_name = "PixelArt"
        lc.is_paused = False
        lc.package_log_cache.clear()
        lc._clear_record_dedup("package")
        lc.active_ios_package_log_frame_buffers.clear()

        first_record = (
            "Sep 18 17:10:59 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
            "[AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting signal_collection_success "
            "with extra parameters {\n"
            '    "adapter_version" = "8.5.0.1";\n'
            '    "network_name" = "YANDEX_BIDDING";\n'
            "}"
        )
        second_record = (
            "Sep 18 17:11:00 iPhone-11-pro PixelArt(AppLovinSDK)[2014] <Notice>: "
            "next record"
        )

        # The first record is completed by NUL; the second stays buffered until
        # the next record boundary or reader shutdown.
        lc.process_ios_package_log_stream_chunk(
            "ios-device",
            first_record + "\x00" + second_record,
        )
        _assert_equal(len(lc.package_log_cache), 1, "NUL-framed iOS package record was not emitted")
        row = lc.package_log_cache[-1]
        _assert('"adapter_version" = "8.5.0.1";' in row["log"], "NUL-framed adapter_version was dropped")
        _assert('"network_name" = "YANDEX_BIDDING";' in row["message"], "NUL-framed network_name was dropped")

        lc._flush_ios_package_log_frame_buffer("ios-device")
        _assert_equal(len(lc.package_log_cache), 2, "NUL-framed final package record was not flushed")
    finally:
        lc.active_platform = original_platform
        lc.target_package_name = original_target
        lc.is_paused = original_paused
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_cache)
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_frame_buffers.update(original_buffers)


def test_ios_package_log_tidevice_newline_stream_does_not_merge_kernel() -> None:
    """A newline tidevice stream must not merge kernel records into the app row."""
    original_platform = lc.active_platform
    original_target = lc.target_package_name
    original_paused = lc.is_paused
    original_cache = list(lc.package_log_cache)
    original_buffers = dict(lc.active_ios_package_log_buffers)
    original_frame_buffers = dict(lc.active_ios_package_log_frame_buffers)
    original_modes = dict(lc.active_ios_package_log_stream_modes)
    try:
        lc.active_platform = "ios"
        lc.target_package_name = "PixelArt"
        lc.is_paused = False
        lc.package_log_cache.clear()
        lc._clear_record_dedup("package")
        lc.active_ios_package_log_buffers.clear()
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_stream_modes.clear()

        stream = (
            "Sep 22 16:56:59 Galaxy-iphone PixelArt(AppLovinSDK)[42702] <Notice>: app event {\n"
            '    "adapter_version" = "4.10.0";\n'
            "}\n"
            "Sep 22 16:57:01 Galaxy-iphone kernel[0] <Notice>: IO80211ControllerMonitor::setAMPDUstat unhandled PixelArt marker\n"
            "Sep 22 16:57:01 Galaxy-iphone akd[8532] <Error>: Attestation map does not contain the cert attestation.\n"
        )
        lc.process_ios_package_log_stream_chunk("ios-device", stream)

        _assert_equal(
            lc.active_ios_package_log_stream_modes.get("ios-device"),
            "line",
            "newline iOS stream was misclassified as NUL-framed",
        )
        _assert_equal(len(lc.package_log_cache), 1, "app record was not emitted at the next record boundary")
        row = lc.package_log_cache[-1]
        _assert_equal(row["tag"], "PixelArt", "newline app row tag was not preserved")
        _assert("kernel[0]" not in row["log"], "kernel record was merged into the app row")
        _assert("akd[8532]" not in row["log"], "akd record was merged into the app row")

        lc._flush_ios_package_log_buffer("ios-device")
        _assert_equal(len(lc.package_log_cache), 1, "non-app trailing record was emitted into Package Log")
    finally:
        lc.active_platform = original_platform
        lc.target_package_name = original_target
        lc.is_paused = original_paused
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_cache)
        lc.active_ios_package_log_buffers.clear()
        lc.active_ios_package_log_buffers.update(original_buffers)
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_frame_buffers.update(original_frame_buffers)
        lc.active_ios_package_log_stream_modes.clear()
        lc.active_ios_package_log_stream_modes.update(original_modes)


def _live_rows_sent(emitted, table):
    """Rows a live table delivered to the UI: flush its queue, then read the deltas."""
    lc._flush_live_table(table)
    event = lc.LIVE_TABLE_DELTA_EVENTS[table]
    return [row for name, payload in emitted if name == event for row in payload["rows"]]


def test_ios_max_load_ads_from_package_frame_contract() -> None:
    """A complete tidevice package frame must also feed the MAX Load Ads parser."""
    original_platform = lc.active_platform
    original_target = lc.target_package_name
    original_paused = lc.is_paused
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_cache = list(lc.package_log_cache)
    original_rows = list(lc.load_ads_ext_events)
    original_unique = set(lc.unique_load_ads_ext)
    original_buffers = dict(lc.active_ios_package_log_frame_buffers)
    original_max_buffer = dict(lc.incomplete_ios_max_load_ads_logs)
    original_emit = lc.socketio.emit
    original_send_to_sheet = lc.send_to_sheet
    emitted = []
    sheet_rows = []
    try:
        lc.active_platform = "ios"
        lc.target_package_name = "PixelArt"
        lc.is_paused = False
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": "NG381"})
        lc.package_log_cache.clear()
        lc._clear_record_dedup("package")
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.active_ios_package_log_frame_buffers.clear()
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()
        lc.send_to_sheet = lambda *args: sheet_rows.append(args)

        frame = (
            "Sep 18 17:17:10 iPhone-11-pro PixelArt(AppLovinSDK)[2628] <Notice>: "
            "[AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting "
            "mediated_ad_viewability_impression_called with extra parameters {\n"
            '    "ad_format" = BANNER;\n'
            '    "adapter_class" = ALChartboostMediationAdapter;\n'
            '    "adapter_version" = "9.14.0.1";\n'
            '    "network_name" = "CHARTBOOST_NETWORK";\n'
            "}"
        )
        lc.process_ios_package_log_stream_chunk("ios-max-frame", frame + "\x00")

        rows = list(lc.load_ads_ext_events)
        _assert_equal(len(rows), 1, "complete tidevice MAX frame was not recorded in Load Ads")
        _assert_equal(rows[0].get("ad_network"), "CHARTBOOST_NETWORK", "package-frame MAX network was not parsed")
        _assert_equal(rows[0].get("ad_format"), "BANNER", "package-frame MAX format was not parsed")
        _assert_equal(len(_live_rows_sent(emitted, "load_ads_ext")), 1, "package-frame MAX row was not emitted")
        _assert_equal(sheet_rows[0][4:], ("LoadAdsExt", "MAX"), "package-frame MAX sheet target was wrong")
    finally:
        lc.active_platform = original_platform
        lc.target_package_name = original_target
        lc.is_paused = original_paused
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_cache)
        lc.load_ads_ext_events.clear()
        lc.load_ads_ext_events.extend(original_rows)
        lc.unique_load_ads_ext.clear()
        lc.unique_load_ads_ext.update(original_unique)
        lc.active_ios_package_log_frame_buffers.clear()
        lc.active_ios_package_log_frame_buffers.update(original_buffers)
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.incomplete_ios_max_load_ads_logs.update(original_max_buffer)
        lc.socketio.emit = original_emit
        lc.send_to_sheet = original_send_to_sheet


def test_sdk_base_name_matching() -> None:
    """Standard names may share a base label; the trailing - MAX rows stay separate."""
    original_expected_map = lc.sdk_check_expected_map
    original_expected_order = lc.sdk_check_expected_order
    try:
        lc.sdk_check_expected_map = {}
        lc.sdk_check_expected_order = []
        for line in (
            "LevelPlay / ironSource\t9.6.0.0\t9.6.0",
            "MAX / AppLovin\t\t13.6.4",
            "LINE Ads/FiveAd\t5.8.0.0\t3.1.1",
            "ironSource - MAX\t9.6.0.0.1",
            "MAX / AppLovin - MAX\t\t13.6.4",
            "LINE Ads/FiveAd - MAX\t3.1.1.1",
        ):
            parsed = lc._parse_sdk_expected_line(line)
            _assert(parsed is not None, f"SDK row cannot be parsed: {line}")
            lc._register_sdk_expected(
                parsed["network"],
                adapter=parsed.get("adapter", ""),
                sdk=parsed.get("sdk", ""),
                source=parsed.get("source", ""),
                match_network=parsed.get("match_network", ""),
            )

        _assert_equal(
            lc._match_sdk_expected_key("IronSource"),
            lc._normalize_sdk_network_name("LevelPlay / ironSource"),
            "IronSource must match the standard LevelPlay / ironSource row",
        )
        _assert_equal(
            lc._match_sdk_expected_key("AppLovin"),
            lc._normalize_sdk_network_name("MAX / AppLovin"),
            "AppLovin must match the standard MAX / AppLovin row",
        )
        _assert_equal(
            lc._match_sdk_expected_key("Line"),
            lc._normalize_sdk_network_name("LINE Ads/FiveAd"),
            "Line must match the standard LINE Ads/FiveAd row",
        )
        _assert_equal(
            lc._match_sdk_expected_key("IronSource - MAX"),
            lc._normalize_sdk_network_name("ironSource - MAX"),
            "IronSource - MAX must match only the MAX row",
        )
        _assert_equal(
            lc._match_sdk_expected_key("AppLovin - MAX"),
            lc._normalize_sdk_network_name("MAX / AppLovin - MAX"),
            "AppLovin - MAX must match only the MAX row",
        )
        _assert_equal(
            lc._match_sdk_expected_key("Line - MAX"),
            lc._normalize_sdk_network_name("LINE Ads/FiveAd - MAX"),
            "Line - MAX must match only the LINE Ads/FiveAd MAX row",
        )
        _assert(
            lc._match_sdk_expected_key("AppLovin") != lc._normalize_sdk_network_name("MAX / AppLovin - MAX"),
            "standard AppLovin must not match the trailing - MAX row",
        )
    finally:
        lc.sdk_check_expected_map = original_expected_map
        lc.sdk_check_expected_order = original_expected_order


def test_sdk_check_preset_contract() -> None:
    presets = lc._load_sdk_check_presets()
    _assert("C-193-Android" in presets, "C-193 Android SDK preset is missing")
    _assert("C-180-Android" in presets, "C-180 Android SDK preset is missing")
    _assert("C-180-iOS" in presets, "C-180 iOS SDK preset is missing")
    c193_android = presets["C-193-Android"]
    c193_android_lines = c193_android.get("lines") or []
    _assert_equal(c193_android.get("platform"), "android", "C-193 Android preset platform changed")
    _assert_equal(len(c193_android_lines), 54, "C-193 Android preset line count changed")
    for required_line in (
        "AppLovin\t5.9.0\t13.6.4",
        "Chartboost\t5.10.0\t9.14.1",
        "Google (AdMob and Ad Manager)\t5.9.0\t25.5.0",
        "MAX / AppLovin - MAX\t\t13.6.4",
        "AudioMob\t10.2.3",
        "AppMetrica SDK\t\t8.5.1",
    ):
        _assert(required_line in c193_android_lines, f"C-193 Android preset entry is missing: {required_line}")
    for line in c193_android_lines[1:]:
        parsed = lc._parse_sdk_expected_line(line)
        _assert(parsed is not None, f"C-191 Android entry cannot be parsed: {line}")

    c192_ios = presets.get("C-192-iOS") or {}
    c192_ios_lines = c192_ios.get("lines") or []
    _assert_equal(c192_ios.get("platform"), "ios", "C-192 iOS preset platform changed")
    _assert_equal(len(c192_ios_lines), 55, "C-192 iOS preset line count changed")
    _assert_equal(c192_ios_lines[0], "Ads Network\tAdapter\tNative", "C-192 iOS preset header changed")
    for required_line in (
        "MAX / AppLovin\t\t13.6.4",
        "Yandex\tSkip\tSkip",
        "ironSource - MAX\t9.6.0.0.1",
        "MAX / AppLovin - MAX\t\t13.6.4",
        "Yandex - MAX\t8.5.0.1",
        "Voodoo - MAX\t3.17.1.1",
        "Voodoo\t5.5.0.0\t3.17.1",
        "Firebase Crashlytics\t\t12.18.0",
    ):
        _assert(required_line in c192_ios_lines, f"C-192 iOS preset entry is missing: {required_line}")
    for line in c192_ios_lines[1:]:
        parsed = lc._parse_sdk_expected_line(line)
        _assert(parsed is not None, f"C-192 iOS entry cannot be parsed: {line}")

    c192_bright = c192_ios.get("bright") or {}
    c192_bright_identifiers = c192_bright.get("skadnetwork_identifiers") or []
    _assert_equal(c192_bright.get("version"), "1.605.415", "C-192 iOS BrightData version changed")
    _assert_equal(len(c192_bright_identifiers), 290, "C-192 iOS Bright/SKAds preset identifier count changed")
    _assert_equal(c192_bright_identifiers[0], "22mmun2rn5.skadnetwork", "C-192 iOS Bright/SKAds first identifier changed")
    _assert_equal(c192_bright_identifiers[-1], "zq492l623r.skadnetwork", "C-192 iOS Bright/SKAds last identifier changed")
    _assert("6lz2ygh3q6.adattributionkit" in c192_bright_identifiers, "C-192 iOS AdAttributionKit identifier is missing")
    bright_compare = lc._compare_bright_skad_preset(
        "1.605.415",
        [c192_bright_identifiers[0], "extra-network.skadnetwork", c192_bright_identifiers[0]],
        preset_name="C-192-iOS",
        preset=c192_ios,
    )
    _assert_equal(bright_compare.get("passed_count"), 2, "Bright/SKAds compare must mark version and one identifier as PASSED")
    _assert_equal(bright_compare.get("total_count"), 291, "Bright/SKAds compare row count changed")
    _assert_equal(bright_compare.get("extra_identifiers"), ["extra-network.skadnetwork"], "Bright/SKAds extra identifier warning changed")
    _assert(any(row.get("status") == "MISSING" for row in bright_compare.get("rows", [])), "Bright/SKAds compare must report missing identifiers")
    c180_ios_preset = presets["C-180-iOS"]
    _assert_equal(c180_ios_preset.get("platform"), "ios", "C-180 iOS preset platform changed")
    c180_ios_lines = c180_ios_preset.get("lines") or []
    _assert_equal(len(c180_ios_lines), 31, "C-180 iOS preset line count changed")
    _assert_equal(c180_ios_lines[0], "Ads Network\tAdapter\tSDK", "C-180 iOS preset header changed")
    _assert("ironSource\t9.4.2.1\t9.4.2" in c180_ios_lines, "C-180 iOS ironSource entry changed")
    _assert("Appsflyer\tremoved\tremoved" in c180_ios_lines, "C-180 iOS Appsflyer entry changed")
    _assert("Firebase Crashlytics\t\t12.15.0" in c180_ios_lines, "C-180 iOS Firebase Crashlytics entry changed")
    _assert(all("http://" not in line and "https://" not in line for line in c180_ios_lines), "C-180 iOS preset must not contain links")
    c180_preset = presets["C-180-Android"]
    _assert_equal(c180_preset.get("platform"), "android", "C-180 preset platform changed")
    c180_lines = c180_preset.get("lines") or []
    _assert_equal(len(c180_lines), 33, "C-180 preset line count changed")
    _assert_equal(c180_lines[0], "Ads Network\tAdapter\tNative", "C-180 preset header changed")
    _assert("Appsflyer\t\tRemoved" in c180_lines, "C-180 Appsflyer entry changed")
    _assert(not any(line.startswith("AdQuality\t") for line in c180_lines), "C-180 AdQuality entry must remain absent")
    _assert(all("http://" not in line and "https://" not in line for line in c180_lines), "C-180 preset must not contain links")
    for line in c180_lines[1:]:
        parsed = lc._parse_sdk_expected_line(line)
        _assert(parsed is not None, f"C-180 Android entry cannot be parsed: {line}")
    for line in c180_ios_lines[1:]:
        parsed = lc._parse_sdk_expected_line(line)
        _assert(parsed is not None, f"C-180 iOS entry cannot be parsed: {line}")

    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert("sdkCheckPresetPanel" in source_text, "SDK preset panel is missing")
    _assert("applySdkCheckPreset" in source_text, "SDK preset action is missing")
    _assert("/api/sdk-check-presets" in source_text, "remote SDK preset endpoint is missing")
    _assert("loadSdkCheckPresetsFromGit" in source_text, "SDK preset remote load action is missing")
    _assert("sdkCheckInput" in source_text, "manual SDK input fallback is missing")
    _assert(source_text.count('id="reloadSdkCheckPresetsBtn"') == 1, "SDK preset reload button must exist exactly once")
    _assert(source_text.count('id="reloadBrightSkadPresetsBtn"') == 1, "Bright/SKAds preset reload button must exist exactly once")
    _assert("clean_lines" in source_text and '"lines": clean_lines' in source_text, "empty SDK presets must remain valid")
    _assert('"skadnetwork_identifiers"' in source_text and "_compare_bright_skad_preset" in source_text, "Bright/SKAds preset comparison support is missing")
    _assert("_merge_local_bright_preset_data" not in source_text, "SDK preset loader must not merge stale local Bright/SKAds data")
    _assert("return presets, \"github\"" in source_text, "SDK preset loader must return the live Git payload")
    _assert('"Accept-Encoding": "identity"' in source_text, "SDK preset fetch must bypass stale compressed cache")
    _assert("def _fetch_sdk_check_presets(force_remote=False):" in source_text, "SDK preset force refresh support is missing")
    _assert("refresh_requested = request.args.get(\"refresh\"" in source_text, "SDK preset refresh query is missing")
    _assert("force_remote=refresh_requested" in source_text, "SDK preset endpoint must honor forced refresh")
    _assert("main/sdk_check_presets.json" in source_text, "SDK preset main fallback URL is missing")


def test_rendered_sdk_preset_javascript_contract() -> None:
    response = lc.app.test_client().get("/")
    _assert_equal(response.status_code, 200, "local UI route must render")
    html = response.get_data(as_text=True)
    expected_join = r"input.value = lines.join('\n');"
    broken_join = "input.value = lines.join('" + "\n" + "');"
    _assert(expected_join in html, "rendered SDK preset JavaScript must keep the newline escape")
    _assert(broken_join not in html, "rendered SDK preset JavaScript contains a literal newline inside a string")
    _assert(
        html.index('id="sdkCheckPresetPanel"') < html.index('id="startSdkCheckBtn"'),
        "SDK preset selector must remain above Start Checking",
    )
    _assert('id="reloadSdkCheckPresetsBtn"' in html, "rendered SDK preset reload button is missing")
    _assert('id="reloadBrightSkadPresetsBtn"' in html, "rendered Bright/SKAds preset reload button is missing")
    _assert("loadSdkCheckPresetsFromGit();" in html, "SDK presets must reload when the app UI starts")
    _assert("loadSdkCheckPresetsFromGit(true);" in html, "preset Reload button must force a fresh GitHub request")
    _assert("refreshQuery = force ? '&refresh=1'" in html, "SDK preset Reload must request a remote refresh")
    _assert("reloadSdkCheckPresetsBtn" in html and "loadSdkCheckPresetsFromGit" in html, "reload button handler is missing")
    _assert("reloadBrightSkadPresetsBtn" in html and "Reloading..." in html, "Bright/SKAds reload button handler is missing")


def test_default_ad_event_contract() -> None:
    normal_events = {
        "interstitial": [
            "ad_load",
            "ad_load_failed",
            "ad_load_succeeded",
            "ad_ready",
            "ad_not_ready",
            "ad_impression",
            "ad_show_succeeded",
            "ad_show_failed",
            "ad_clicked",
            "ad_closed",
        ],
        "rewarded": [
            "ad_load",
            "ad_load_failed",
            "ad_load_succeeded",
            "ad_ready",
            "ad_not_ready",
            "ad_impression",
            "ad_show_succeeded",
            "ad_show_failed",
            "ad_clicked",
            "ad_closed",
            "ad_rewarded",
        ],
        "banner": [
            "ad_load",
            "ad_load_failed",
            "ad_load_succeeded",
            "ad_impression",
            "ad_show_succeeded",
            "ad_show_failed",
            "ad_clicked",
        ],
        "mrec": [
            "ad_load",
            "ad_load_failed",
            "ad_load_succeeded",
            "ad_impression",
            "ad_show_succeeded",
            "ad_show_failed",
            "ad_clicked",
        ],
    }
    simple_events = {
        "audioads": {
            "odeeo": [
                "ad_load_succeeded",
                "ad_ready",
                "ad_not_ready",
                "ad_impression",
                "ad_show_succeeded",
                "ad_show_failed",
                "ad_clicked",
                "ad_closed",
                "ad_rewarded",
            ],
            "audiomob": [
                "ad_load",
                "ad_load_failed",
                "ad_impression",
                "ad_show_succeeded",
                "ad_show_failed",
                "ad_clicked",
            ],
        },
        "inplayads": {
            "adverty": ["ad_impression", "ad_show_succeeded", "ad_clicked"],
            "gadsme": ["ad_load_failed", "ad_load_succeeded", "ad_impression", "ad_show_succeeded", "ad_clicked"],
        },
    }

    workbook = Workbook()
    workbook.active.title = "events"

    def add_normal_sheet(provider: str, platform_value=None) -> None:
        worksheet = workbook.create_sheet(provider)
        platform_value = platform_value or provider
        for column, ad_format in zip((2, 6, 10, 14), lc.DEFAULT_AD_EVENT_FORMATS):
            worksheet.merge_cells(
                start_row=2,
                start_column=column,
                end_row=2,
                end_column=column + 2,
            )
            worksheet.cell(2, column).value = lc.DEFAULT_AD_EVENT_FORMAT_LABELS[ad_format]
            row = 3
            for event_name in normal_events[ad_format]:
                worksheet.cell(row, column).value = event_name
                worksheet.cell(row + 1, column + 1).value = "ad_platform"
                worksheet.cell(row + 1, column + 2).value = platform_value
                worksheet.cell(row + 2, column + 1).value = "ad_format"
                worksheet.cell(row + 2, column + 2).value = ad_format
                row += 3

    add_normal_sheet("Ironsource", "ironsource")
    add_normal_sheet("MAX", "max")
    add_normal_sheet("Cloudx")
    add_normal_sheet("Ascendx")

    def add_simple_sheet(category: str) -> None:
        worksheet = workbook.create_sheet(lc.DEFAULT_AD_EVENT_SIMPLE_CATEGORY_LABELS[category])
        for column, provider in zip((2, 6), lc.DEFAULT_AD_EVENT_SIMPLE_PROVIDERS[category]):
            worksheet.merge_cells(
                start_row=2,
                start_column=column,
                end_row=2,
                end_column=column + 2,
            )
            worksheet.cell(2, column).value = provider
            row = 3
            for event_name in simple_events[category][provider]:
                worksheet.cell(row, column).value = event_name
                worksheet.cell(row + 1, column + 1).value = "ad_platform"
                worksheet.cell(row + 1, column + 2).value = provider
                worksheet.cell(row + 2, column + 1).value = "ad_format"
                worksheet.cell(row + 2, column + 2).value = "audio_rewarded" if category == "audioads" else "in_play"
                row += 3

    add_simple_sheet("audioads")
    add_simple_sheet("inplayads")

    parsed_normal = lc._read_default_ad_event_config(workbook)
    expected_normal = {
        ad_format: {
            provider: list(normal_events[ad_format])
            for provider in lc.DEFAULT_AD_EVENT_PROVIDERS
        }
        for ad_format in lc.DEFAULT_AD_EVENT_FORMATS
    }
    _assert_equal(parsed_normal, expected_normal, "normal Default Ad Event sheets parsed incorrectly")
    parsed_simple = lc._read_default_ad_event_simple_config(workbook)
    _assert_equal(parsed_simple, simple_events, "AudioAds/InplayAds sheets parsed incorrectly")

    empty_workbook = Workbook()
    empty_workbook.active.title = "events"
    empty_normal = lc._read_default_ad_event_config(empty_workbook)
    empty_simple = lc._read_default_ad_event_simple_config(empty_workbook)
    _assert(
        not any(event_names for providers in empty_normal.values() for event_names in providers.values()),
        "missing provider sheets must not create normal ad-event rows",
    )
    _assert(
        not any(event_names for providers in empty_simple.values() for event_names in providers.values()),
        "missing AudioAds/InplayAds sheets must not create simple ad-event rows",
    )

    original_config = lc.default_ad_event_config
    original_simple_config = lc.default_ad_event_simple_config
    original_hits = set(lc.default_ad_event_hits)
    original_simple_hits = set(lc.default_ad_event_simple_hits)
    original_clients = set(lc.default_ad_event_clients)
    original_paused = lc.is_paused
    original_emit = lc.socketio.emit
    emitted = []
    try:
        lc.default_ad_event_config = parsed_normal
        lc.default_ad_event_simple_config = parsed_simple
        lc.default_ad_event_hits.clear()
        lc.default_ad_event_simple_hits.clear()
        lc.default_ad_event_clients.clear()
        lc.default_ad_event_clients.add("harness-client")
        lc.is_paused = False
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))

        line = (
            'Unity : TrackingService->Track: '
            '{"eventName":"ad_impression","e":{"ad_platform":"max","ad_format":"Rewarded"}}'
        )
        event_name, actual_params, _json_string = lc.find_and_parse_event(line)
        _assert_equal(event_name, "ad_impression", "generic ad event was not parsed")
        _assert_equal(
            actual_params,
            {"ad_platform": "max", "ad_format": "Rewarded"},
            "generic ad event parameters were not parsed",
        )
        lc._record_default_ad_event_hit(event_name, actual_params, "device-normal")
        lc._record_default_ad_event_hit(
            "ad_clicked",
            {"ad_platform": "Audiomob", "ad_format": "audio_skippable"},
            "device-simple",
        )
        lc._record_default_ad_event_hit(
            "ad_rewarded",
            {"ad_platform": "Cloudx", "ad_format": "rewarded"},
            "device-hidden-provider",
        )

        payload = lc._default_ad_event_payload()
        _assert_equal(
            payload["hits"],
            [{
                "device_id": "device-normal",
                "provider": "max",
                "ad_format": "rewarded",
                "event_name": "ad_impression",
            }],
            "MAX ad-event match did not use max ad_platform and ad_format",
        )
        _assert(
            not any(hit.get("provider") in {"cloudx", "ascendx"} for hit in payload["hits"]),
            "temporarily hidden providers must not create visible coverage hits",
        )
        _assert_equal(
            payload["simple_hits"],
            [{
                "device_id": "device-simple",
                "category": "audioads",
                "provider": "audiomob",
                "event_name": "ad_clicked",
            }],
            "simple ad-event match did not use ad_platform-only matching",
        )
        _assert_equal(len(emitted), 2, "only matching ad events should emit coverage updates")

        catalog = lc._default_ad_event_catalog_payload()
        _assert_equal(
            [provider["key"] for provider in catalog["providers"]],
            ["ironsource", "max"],
            "Default Ad Events must expose Ironsource and MAX only",
        )
        _assert_equal(
            [provider["label"] for provider in catalog["providers"]],
            ["Ironsource", "MAX"],
            "Default Ad Events provider labels changed",
        )

        rendered = lc.app.test_client().get("/")
        _assert_equal(rendered.status_code, 200, "Default Ad Events UI route must render")
        html = rendered.get_data(as_text=True)
        for marker in (
            ">Default Ad Events</button>",
            "defaultAdEventMatrix",
            "defaultAdEventHitKey",
            "defaultAdEventSimpleHitKey",
            "Odeeo",
            "Audiomob",
            "Adverty",
            "Gadsme",
            "MAX",
            "Simple providers are intentionally rendered as separate blocks",
        ):
            _assert(marker in html, f"Default Ad Events UI contract is missing: {marker}")
    finally:
        lc.default_ad_event_config = original_config
        lc.default_ad_event_simple_config = original_simple_config
        lc.default_ad_event_hits.clear()
        lc.default_ad_event_hits.update(original_hits)
        lc.default_ad_event_simple_hits.clear()
        lc.default_ad_event_simple_hits.update(original_simple_hits)
        lc.default_ad_event_clients.clear()
        lc.default_ad_event_clients.update(original_clients)
        lc.is_paused = original_paused
        lc.socketio.emit = original_emit


def test_ios_default_ad_event_tracking_record() -> None:
    original_config = lc.default_ad_event_config
    original_simple_config = lc.default_ad_event_simple_config
    original_hits = set(lc.default_ad_event_hits)
    original_simple_hits = set(lc.default_ad_event_simple_hits)
    original_clients = set(lc.default_ad_event_clients)
    original_paused = lc.is_paused
    original_emit = lc.socketio.emit
    emitted = []
    try:
        lc.default_ad_event_config = lc._new_default_ad_event_config()
        lc.default_ad_event_config["interstitial"]["max"] = ["ad_load"]
        lc.default_ad_event_simple_config = lc._new_default_ad_event_simple_config()
        lc.default_ad_event_hits.clear()
        lc.default_ad_event_simple_hits.clear()
        lc.default_ad_event_clients.clear()
        lc.is_paused = False
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))

        line = (
            "Sep 21 13:54:01 iPhone-11-pro PixelArt(UnityFramework)[1639] <Notice>: "
            '[Tracking] TrackingService->Track: '
            '{"EventName":"ad_load","params":{"ad_platform":"max",'
            '"ad_format":"interstitial","ad_unit_id":"5a0e30d201088d93",'
            '"ad_placement_name":"null","ab_test_ironsource":"null",'
            '"ab_test_firebase":"null","audience_name":"default",'
            '"event_time":1789973641848,"session_id":1789973631660,'
            '"first_open_time":1789973631667,"session_count":1,'
            '"game_start_count":0,"level":-1,"map":1,"mode_game":"null",'
            '"level_design":"a"}}'
        )
        event_name, actual_params, _json_string = lc.find_and_parse_event(line)
        _assert_equal(event_name, "ad_load", "iOS TrackingService event name was not parsed")
        _assert_equal(actual_params["ad_platform"], "max", "iOS ad platform was not parsed")
        _assert_equal(actual_params["ad_format"], "interstitial", "iOS ad format was not parsed")

        # Coverage must not depend on the Default Ad Events tab being open when
        # the device emits the event.
        lc._record_default_ad_event_hit(event_name, actual_params, "ios-device")
        _assert_equal(
            lc._default_ad_event_payload()["hits"],
            [{
                "device_id": "ios-device",
                "provider": "max",
                "ad_format": "interstitial",
                "event_name": "ad_load",
            }],
            "iOS TrackingService ad_load did not create MAX coverage",
        )
        _assert_equal(len(emitted), 1, "iOS default ad-event coverage update was not emitted")
    finally:
        lc.default_ad_event_config = original_config
        lc.default_ad_event_simple_config = original_simple_config
        lc.default_ad_event_hits.clear()
        lc.default_ad_event_hits.update(original_hits)
        lc.default_ad_event_simple_hits.clear()
        lc.default_ad_event_simple_hits.update(original_simple_hits)
        lc.default_ad_event_clients.clear()
        lc.default_ad_event_clients.update(original_clients)
        lc.is_paused = original_paused
        lc.socketio.emit = original_emit


def test_exact_duplicate_record_dedup_across_ios_tabs() -> None:
    """An identical iOS source record must produce one row per tab cache."""
    original_platform = lc.active_platform
    original_paused = lc.is_paused
    original_validator_active = lc.validator_active
    original_target = lc.target_package_name
    original_session = lc.active_package_log_session_id
    original_emit = lc.socketio.emit
    original_callback_rows = list(lc.callback_ad_logs)
    original_validator_rows = list(lc.validator_results)
    original_specific_rows = list(lc.event_log_cache)
    original_adrevenue_rows = list(lc.adrevenue_logs)
    original_price_rows = list(lc.price_rotation_logs)
    original_package_rows = list(lc.package_log_cache)
    original_dedup_seen = {name: set(values) for name, values in lc.record_dedup_seen.items()}
    original_dedup_order = {name: list(values) for name, values in lc.record_dedup_order.items()}
    try:
        lc.active_platform = "ios"
        lc.is_paused = False
        lc.validator_active = True
        lc.target_package_name = "PixelArt"
        lc.active_package_log_session_id = None
        lc.callback_ad_logs.clear()
        lc.validator_results.clear()
        lc.event_log_cache.clear()
        lc.adrevenue_logs.clear()
        lc.price_rotation_logs.clear()
        lc.package_log_cache.clear()
        lc._clear_record_dedup()
        lc.socketio.emit = lambda *_args, **_kwargs: None

        line = (
            "Sep 21 13:54:01 iPhone-11-pro PixelArt(UnityFramework)[1639] <Notice>: "
            '[Tracking] TrackingService->Track: {"EventName":"ad_load",'
            '"params":{"ad_platform":"max","ad_format":"interstitial"}}'
        )
        event_name, params, json_string = lc.find_and_parse_event(line)
        lc.process_callback_and_ad_event_log(line, "ios-device", event_name, params, json_string)
        lc.process_callback_and_ad_event_log(line, "ios-device", event_name, params, json_string)
        _assert_equal(len(lc.callback_ad_logs), 1, "duplicate iOS callback/ad event row was stored")

        lc.process_event_validator_log(event_name, params, json_string, line, "ios-device")
        lc.process_event_validator_log(event_name, params, json_string, line, "ios-device")
        _assert_equal(len(lc.validator_results), 1, "duplicate iOS validator row was stored")

        lc.cache_specific_event_log(event_name, params, json_string, line, "ios-device")
        lc.cache_specific_event_log(event_name, params, json_string, line, "ios-device")
        _assert_equal(len(lc.event_log_cache), 1, "duplicate iOS specific-event row was stored")

        lc._append_package_log_row("ios-device", line)
        lc._append_package_log_row("ios-device", line)
        _assert_equal(len(lc.package_log_cache), 1, "duplicate iOS Package Log row was stored")
    finally:
        lc.active_platform = original_platform
        lc.is_paused = original_paused
        lc.validator_active = original_validator_active
        lc.target_package_name = original_target
        lc.active_package_log_session_id = original_session
        lc.callback_ad_logs.clear()
        lc.callback_ad_logs.extend(original_callback_rows)
        lc.validator_results.clear()
        lc.validator_results.extend(original_validator_rows)
        lc.event_log_cache.clear()
        lc.event_log_cache.extend(original_specific_rows)
        lc.adrevenue_logs.clear()
        lc.adrevenue_logs.extend(original_adrevenue_rows)
        lc.price_rotation_logs.clear()
        lc.price_rotation_logs.extend(original_price_rows)
        lc.package_log_cache.clear()
        lc.package_log_cache.extend(original_package_rows)
        lc._clear_record_dedup()
        for name, values in original_dedup_seen.items():
            lc.record_dedup_seen[name].update(values)
            lc.record_dedup_order[name].extend(original_dedup_order[name])
        lc.socketio.emit = original_emit


def test_installation_id_copy_contract() -> None:
    for source_path in (ROOT / "Log_checker.py",):
        source_text = source_path.read_text(encoding="utf-8", errors="ignore")
        _assert("async function copyTextToClipboard(text)" in source_text, f"copy helper missing in {source_path}")
        _assert("await navigator.clipboard.writeText(value);" in source_text, f"Clipboard API path missing in {source_path}")
        _assert("document.execCommand('copy')" in source_text, f"Windows clipboard fallback missing in {source_path}")
        _assert(
            "if (!await copyTextToClipboard(installationId))" in source_text,
            f"Installation ID button is not using the cross-platform copy helper in {source_path}",
        )
        _assert(
            "await navigator.clipboard.writeText(installationId);" not in source_text,
            f"Installation ID button still bypasses the cross-platform copy helper in {source_path}",
        )
        _assert(
            'id="installationIdValue" type="text" readonly' in source_text,
            f"Installation ID must be exposed as a selectable readonly text field in {source_path}",
        )
        _assert(
            "installationIdValueEl.value = installationId;" in source_text,
            f"Installation ID text field is not updated through its value in {source_path}",
        )
        _assert(
            "installationIdValueEl.select();" in source_text,
            f"Installation ID text field does not select its value on focus in {source_path}",
        )
        _assert(
            ".installation-id-value" in source_text and "user-select: text !important" in source_text,
            f"Installation ID text field is not explicitly selectable in {source_path}",
        )

    rendered = lc.app.test_client().get("/").get_data(as_text=True)
    _assert("async function copyTextToClipboard(text)" in rendered, "rendered copy helper is missing")
    _assert("document.execCommand('copy')" in rendered, "rendered Windows clipboard fallback is missing")
    _assert('id="installationIdValue" type="text" readonly' in rendered, "rendered Installation ID field is not selectable")
    _assert("id=\"clearUpdateCacheBtn\"" not in rendered, "Clear Cache button must stay removed from the release UI")


def test_sdk_failed_groups_sort_first() -> None:
    groups = [
        [{"status": "PASSED"}, {"status": "PASSED"}],
        [{"status": "FAILED"}],
        [{"status": "FOUND"}],
        [{"status": "PASSED"}, {"status": "FAILED"}],
    ]
    ordered = sorted(groups, key=lc._sdk_result_group_sort_key)
    _assert(ordered[0][0]["status"] == "FAILED", "failed SDK group must be first")
    _assert(ordered[1][1]["status"] == "FAILED", "mixed failed SDK group must remain before passed groups")
    _assert(ordered[-1][0]["status"] == "FOUND", "found-only SDK group must remain after passed groups")
    _assert_equal(lc._sdk_result_status("", "1.18.3.0"), "FAILED", "missing actual with expected version must fail")
    _assert_equal(lc._sdk_result_status("", ""), "NOT_FOUND", "empty expected version must be not found")
    _assert_equal(lc._sdk_result_status("NOT FOUND", "Removed"), "PASSED", "removed SDK with missing actual must pass")
    _assert_equal(lc._sdk_result_status("MISSING", "Removed"), "PASSED", "removed SDK with missing marker must pass")


def test_release_build_marker() -> None:
    text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert(
        f"v2.5.0({RELEASE_SOURCE_BUILD})" in text,
        f"Log_checker.py must be prepared for v2.5.0({RELEASE_SOURCE_BUILD})",
    )


def test_preset_revert_protection_contract() -> None:
    script = (ROOT / "tools" / "backup_latest.sh").read_text(encoding="utf-8", errors="ignore")
    policy = (ROOT / "BACKUP_POLICY.md").read_text(encoding="utf-8", errors="ignore")
    protected_paths = (
        "sdk_check_presets.json",
        "package_log_presets.json",
        "services_checker/apk_check_presets.json",
        "services_checker/gradle_check_presets.json",
        "services_checker/gradle_lib_mapping.json",
        "services_checker/podfile_check_presets.json",
        "services_checker/podfile_lib_mapping.json",
        "services_checker/manifest_check_presets.json",
    )
    snapshot_index = script.index("snapshot_protected_presets")
    reset_index = script.index('git reset --hard "$commit"')
    restore_index = script.index("restore_protected_presets", reset_index)
    _assert(snapshot_index < reset_index < restore_index, "preset files must be restored after an app-code revert")
    for path in protected_paths:
        _assert(path in script, f"revert protection is missing for {path}")
        _assert(path in policy, f"preset policy is missing for {path}")


def test_rewarded_bidding_filter_contract() -> None:
    """Package Log quick filters: separate Android/iOS lists, each filter = Tag / Message 1 / Message 2 boxes."""
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    expected_names = ["Integrationhelper", "Appmetrica", "Adjust", "Bidding", "Applovinsdk", "Event tracking", "Ads Event"]
    for platform in ("android", "ios"):
        defaults = lc.DEFAULT_PACKAGE_QUICK_FILTERS[platform]
        _assert_equal([f["name"] for f in defaults], expected_names, f"{platform} default quick filters changed")
        bidding = next(f for f in defaults if f["id"] == "bidding")
        _assert_equal((bidding["tag"], bidding["message1"], bidding["message2"]), ("", "[ad,", "bidding"), f"{platform} Bidding must be Message 1 [ad, + Message 2 bidding")
        ads = next(f for f in defaults if f["id"] == "ads_event")
        _assert_equal((ads["tag"], ads["message1"], ads["message2"]), ("", "TrackingService->Track:", '"EventName":"ad'), f"{platform} Ads Event default is wrong")
        tracking = next(f for f in defaults if f["id"] == "event_tracking")
        _assert_equal((tracking["message1"], tracking["message2"]), ("TrackingService->Track:", '"EventName":"'), f"{platform} Event tracking default is wrong")
        _assert(not lc._validate_package_quick_filters([dict(f) for f in defaults], platform)[1], f"{platform} defaults must pass validation")
        applovin = next(f for f in defaults if f["id"] == "applovinsdk")
        _assert_equal((applovin["tag"], applovin["message1"]), ("", "applovinsdk"), f"{platform} Applovinsdk must search the whole log")
    _assert(all(f["tag"] == "" for f in lc.DEFAULT_PACKAGE_QUICK_FILTERS["ios"]), "iOS defaults must not use the tag")
    _assert_equal(next(f for f in lc.DEFAULT_PACKAGE_QUICK_FILTERS["android"] if f["id"] == "adjust")["tag"], "adjust", "Android Adjust must filter by tag")
    _assert_equal(next(f for f in lc.DEFAULT_PACKAGE_QUICK_FILTERS["ios"] if f["id"] == "adjust")["message1"], "adjust]", "iOS Adjust must search adjust]")
    _assert(source_text.count('value="rewardedcap"') == 1, "RewardedCap filter must exist exactly once")
    _assert("case_sensitive" not in source_text and "caseSensitive" not in source_text, "quick filters have no case-sensitive option")
    for needle in (
        'id="packageQuickFilterList"', 'id="editPackageQuickFiltersBtn"', 'id="packageQuickFilterEditor"', 'id="packageHistoryQuickFilter"',
        'id="packageQuickFilterEditorPlatform"', 'class="qf-tag-col',
        'id="exportPackageQuickFiltersBtn"', 'id="importPackageQuickFiltersBtn"', 'id="packageQuickFilterImportInput"',
        "fetch('/api/package-quick-filters/export'", "fetch('/api/package-quick-filters/import'",
        'mb-1">Filter 1:</label>', 'mb-1">Filter 2:</label>', 'border-b">Filter 1</th>', 'border-b">Filter 2</th>',
        "var packageQuickFilters = {{ package_quick_filters|tojson }};",
        "return (packageQuickFilters && packageQuickFilters[quickFilterPlatform()]) || [];",
        "if (quick.tag && !tagHaystack.includes(quick.tag)) return false;",
        "if (quick.terms.some(term => !logHaystack.includes(term))) return false;",
        "const norm = text => String(text || '').trim().toLowerCase();",
        "state.filterText && !logHaystack.includes(state.filterText)",
        "state.filterText2 && !logHaystack.includes(state.filterText2)",
        "body: JSON.stringify({ platform, filters }),",
        "qf: quickId,", "qf: filteredOnly ? quickId : '',",
        "document.getElementById('packageHistoryQuickFilter').value = quickId;",
        "['packageTagFilterInput', 'packageFilterInput', 'packageFilterInput2', 'packageFilterInput3'].forEach(id => {",
        "['packageHistoryFilterInput', 'packageHistoryFilterInput2', 'packageHistoryFilterInput3', 'packageHistoryFilterInput4'].forEach(id => {",
        'id="packageFilterInput3"', 'mb-1">Filter 3:</label>', 'id="packageHistoryFilterInput4"', 'border-b">Filter 3</th>',
        "state.filterText3 && !logHaystack.includes(state.filterText3)",
        "q4: keyword3,", "q4: filteredOnly ? keyword3 : '',",
        "document.getElementById('packageHistoryFilterInput4').value = message3;",
        'data-field="message3"',
    ):
        _assert(needle in source_text, f"quick filter contract missing: {needle}")


def test_package_log_switch_project_keeps_session() -> None:
    """Switching project in the Stream window = Stop + Start: the old session is closed and kept, a new one starts."""
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    for needle in (
        'id="packageStreamProjectSelect"',
        "function switchPackageStreamProject(pkg)",
        "if (!pkg || pkg === packageStreamTarget) return;",
        "socket.emit('start_package_log', {package_id: pkg});",
        "packageStreamTarget = payload.package_id || '';",
    ):
        _assert(needle in source_text, f"project switcher contract missing: {needle}")
    with _TempPackageDb() as db:
        lc._init_package_log_db()
        client = lc.socketio.test_client(lc.app)
        try:
            client.get_received()
            client.emit("start_package_log", {"package_id": "com.first.app"})
            first_session = lc.active_package_log_session_id
            client.emit("start_package_log", {"package_id": "com.second.app"})
            second_session = lc.active_package_log_session_id
            resets = [msg["args"][0] for msg in client.get_received() if msg["name"] == "package_log_reset"]
            _assert_equal([r["package_id"] for r in resets], ["com.first.app", "com.second.app"], "each start must announce its project")
            _assert(first_session and second_session and first_session != second_session, "switching must open a new session")
            _assert_equal(lc.target_package_name, "com.second.app", "the new project must be captured")
            conn = db.connect()
            rows = {row[0]: (row[1], row[2]) for row in conn.execute("SELECT id, package_id, status FROM package_log_sessions")}
            conn.close()
            _assert_equal(rows[first_session], ("com.first.app", "stopped"), "the previous project's session must be closed and kept")
            _assert_equal(rows[second_session], ("com.second.app", "running"), "the new project's session must be running")
            client.emit("start_package_log", {"package_id": ""})
        finally:
            client.disconnect()
            lc.target_package_name = ""
            lc.active_package_log_session_id = None


def test_package_history_split_by_platform() -> None:
    """Recorded Logs list Android and iOS sessions separately and filter each with its own quick filters."""
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    for needle in (
        'id="packageHistoryPlatformSwitch"',
        "fetch(`/api/package-log/sessions?limit=500&platform=${packageHistoryPlatform}`)",
        "qp: packageHistoryPlatform,",
        "historyQuickFilters().find(f => f.id === quickId)",
        "logDetailModal.style.zIndex = '80';",
    ):
        _assert(needle in source_text, f"Recorded platform split contract missing: {needle}")
    _assert("packageQuickFilters.find(" not in source_text, "packageQuickFilters is per platform; .find on it throws")
    original_data_dir = os.environ.get("EVENTINSPECTOR_DATA_DIR")
    original_presets = lc._load_package_log_presets
    original_platform = lc.active_platform
    with tempfile.TemporaryDirectory() as temp_dir, _TempPackageDb() as db:
        os.environ["EVENTINSPECTOR_DATA_DIR"] = temp_dir
        lc._load_package_log_presets = lambda: [{"code": "NG", "android": {"id": "com.indiez.nonogram"}, "ios": {"id": "PixelArt"}}]
        try:
            lc._init_package_log_db()
            conn = db.connect()
            columns = {row[1] for row in conn.execute("PRAGMA table_info(package_log_sessions)")}
            _assert("platform" in columns, "sessions need a platform column")
            # Sessions from before the column: resolved from the package id.
            old_android = conn.execute("INSERT INTO package_log_sessions (package_id, started_at, status) VALUES ('com.indiez.nonogram', 1, 'stopped')").lastrowid
            old_ios = conn.execute("INSERT INTO package_log_sessions (package_id, started_at, status) VALUES ('PixelArt', 2, 'stopped')").lastrowid
            old_unknown_ios = conn.execute("INSERT INTO package_log_sessions (package_id, started_at, status) VALUES ('SomeBundle', 3, 'stopped')").lastrowid
            for sid, tag, message in ((old_android, "Adjust", "event tracked"), (old_ios, "", "[Adjust]: event tracked")):
                conn.execute(
                    "INSERT INTO package_log_entries (session_id, created_at, time_display, device_id, device_name, level, tag, message, raw_log) VALUES (?, 1, 't', 'd', 'Dev', 'I', ?, ?, ?)",
                    (sid, tag, message, f"t I {tag}: {message}"),
                )
            conn.commit()
            conn.close()
            lc.active_platform = "ios"
            new_ios = lc._start_package_log_session("TrainIdle")
            lc._finish_package_log_session(new_ios)
            lc.active_platform = "android"
            new_android = lc._start_package_log_session("com.indiez.train.miner")
            lc._finish_package_log_session(new_android)
            client = lc.app.test_client()

            def ids(platform):
                data = client.get(f"/api/package-log/sessions?platform={platform}").get_json()
                _assert(data["ok"], f"sessions API failed for {platform}")
                _assert(all(s["platform"] == platform for s in data["sessions"]), f"{platform} list holds another platform")
                _assert_equal(data["total"], len(data["sessions"]), f"{platform} total must count only that platform")
                return sorted(s["id"] for s in data["sessions"])

            _assert_equal(ids("android"), sorted([old_android, new_android]), "Android sessions are wrong")
            _assert_equal(ids("ios"), sorted([old_ios, old_unknown_ios, new_ios]), "iOS sessions are wrong")
            _assert_equal(len(client.get("/api/package-log/sessions").get_json()["sessions"]), 5, "no platform must list every session")
            _assert_equal(client.get("/api/package-log/sessions?platform=web").status_code, 400, "unknown platform must be refused")

            def adjust(session_id, qp):
                data = client.get(f"/api/package-log/logs?session_id={session_id}&qf=adjust&qp={qp}").get_json()
                _assert(data["ok"], f"quick filter failed: {data}")
                return [row["message"] for row in data["rows"]]

            # The session's own platform decides which quick filter list applies,
            # whatever platform the app is on (qp).
            _assert_equal(adjust(old_ios, "android"), ["[Adjust]: event tracked"], "iOS session must use the iOS Adjust filter")
            _assert_equal(adjust(old_android, "ios"), ["event tracked"], "Android session must use the Android Adjust filter (tag)")
        finally:
            lc._load_package_log_presets = original_presets
            lc.active_platform = original_platform
            if original_data_dir is None:
                os.environ.pop("EVENTINSPECTOR_DATA_DIR", None)
            else:
                os.environ["EVENTINSPECTOR_DATA_DIR"] = original_data_dir


def test_package_quick_filters_api() -> None:
    """Quick filters are user settings in the data dir, one file per platform."""
    original_data_dir = os.environ.get("EVENTINSPECTOR_DATA_DIR")
    with tempfile.TemporaryDirectory() as temp_dir:
        os.environ["EVENTINSPECTOR_DATA_DIR"] = temp_dir
        try:
            client = lc.app.test_client()
            for platform in ("android", "ios"):
                data = client.get(f"/api/package-quick-filters?platform={platform}").get_json()
                _assert(data["ok"] and data["is_default"], f"{platform} must start with the default quick filters")
                _assert_equal(data["filters"], data["defaults"], f"{platform} defaults are not served")
            good = {"id": "iap", "name": "IAP", "tag": " Unity ", "message1": " [iap] ", "message2": ""}
            bad_cases = [
                {"filters": [good]},
                {"platform": "web", "filters": [good]},
                {"platform": "android", "filters": "x"},
                {"platform": "android", "filters": [dict(good, name=" ")]},
                {"platform": "android", "filters": [dict(good, tag="", message1="", message2=" ")]},
                {"platform": "android", "filters": [dict(good, id="bad id")]},
                {"platform": "android", "filters": [good, dict(good)]},
                {"platform": "android", "filters": [dict(good, message1="x" * 201)]},
                {"platform": "ios", "filters": [good]},
                {"platform": "android", "filters": [dict(good, id=f"f{i}") for i in range(lc.PACKAGE_QUICK_FILTER_MAX + 1)]},
            ]
            for payload in bad_cases:
                _assert_equal(client.post("/api/package-quick-filters", json=payload).status_code, 400, f"invalid quick filters were accepted: {payload}")
            saved = client.post("/api/package-quick-filters", json={"platform": "android", "filters": [good]}).get_json()
            _assert_equal((saved["filters"][0]["tag"], saved["filters"][0]["message1"]), ("Unity", "[iap]"), "filter text must be trimmed")
            legacy = client.post("/api/package-quick-filters", json={"platform": "android", "filters": [dict(good, case_sensitive=True)]}).get_json()
            _assert(legacy["ok"] and "case_sensitive" not in legacy["filters"][0], "an old case_sensitive flag must be ignored, not stored")
            android_path = os.path.join(temp_dir, "package_quick_filters_android.json")
            ios_path = os.path.join(temp_dir, "package_quick_filters_ios.json")
            _assert(os.path.exists(android_path) and not os.path.exists(ios_path), "Android filters must be saved in their own file only")
            _assert(client.get("/api/package-quick-filters?platform=ios").get_json()["is_default"], "saving Android must not change iOS")
            ios_filter = dict(good, id="ios_iap", name="iOS IAP", tag="")
            client.post("/api/package-quick-filters", json={"platform": "ios", "filters": [ios_filter]})
            _assert(os.path.exists(ios_path), "iOS filters must be saved in their own file")
            _assert_equal(client.get("/api/package-quick-filters?platform=android").get_json()["filters"], saved["filters"], "saving iOS must not change Android")
            html = client.get("/").get_data(as_text=True)
            _assert('"name": "IAP"' in html and '"name": "iOS IAP"' in html, "both saved lists must be rendered into the page")
            with open(android_path, "w", encoding="utf-8") as handle:
                handle.write("{broken")
            _assert(client.get("/api/package-quick-filters?platform=android").get_json()["is_default"], "a broken file must fall back to defaults")
            client.post("/api/package-quick-filters", json={"platform": "android", "filters": [good]})
            reset = client.post("/api/package-quick-filters", json={"platform": "ios", "reset": True}).get_json()
            _assert(reset["is_default"] and not os.path.exists(ios_path), "reset must remove only that platform's file")
            _assert(os.path.exists(android_path), "resetting iOS must not touch the Android file")
            _assert(not any(name.startswith("package_quick_filters") for name in os.listdir(ROOT)), "quick filters must not be written into the app folder")
        finally:
            if original_data_dir is None:
                os.environ.pop("EVENTINSPECTOR_DATA_DIR", None)
            else:
                os.environ["EVENTINSPECTOR_DATA_DIR"] = original_data_dir


def test_package_quick_filter_upgrade_backup_export_import() -> None:
    """User filters survive updates, new defaults are appended once, broken files are kept, export/import round-trips."""
    original_data_dir = os.environ.get("EVENTINSPECTOR_DATA_DIR")
    original_defaults = lc.DEFAULT_PACKAGE_QUICK_FILTERS
    with tempfile.TemporaryDirectory() as temp_dir:
        os.environ["EVENTINSPECTOR_DATA_DIR"] = temp_dir
        try:
            client = lc.app.test_client()
            path = os.path.join(temp_dir, "package_quick_filters_android.json")
            mine = {"id": "fmine", "name": "Mine", "tag": "", "message1": "my text", "message2": "", "message3": ""}
            kept = [dict(f) for f in original_defaults["android"] if f["id"] != "ads_event"] + [mine]
            # A file saved before seen_defaults existed: nothing is added back.
            with open(path, "w", encoding="utf-8") as handle:
                json.dump({"version": 2, "platform": "android", "filters": kept}, handle)
            filters, is_default = lc._load_package_quick_filters("android")
            _assert(not is_default and [f["id"] for f in filters] == [f["id"] for f in kept], "a legacy file must load unchanged (deleted defaults stay deleted)")
            with open(path, encoding="utf-8") as handle:
                _assert_equal(sorted(json.load(handle)["seen_defaults"]), sorted(f["id"] for f in original_defaults["android"]), "a legacy file must record the defaults it has seen")

            # A newer release ships a new default: it is appended once, nothing else changes.
            new_default = {"id": "iap", "name": "IAP", "tag": "", "message1": "[iap]", "message2": "", "message3": ""}
            lc.DEFAULT_PACKAGE_QUICK_FILTERS = {"android": tuple(original_defaults["android"]) + (new_default,), "ios": original_defaults["ios"]}
            filters, _ = lc._load_package_quick_filters("android")
            _assert_equal([f["id"] for f in filters], [f["id"] for f in kept] + ["iap"], "a new default must be appended after the user's filters")
            _assert_equal(filters[len(kept) - 1], mine, "the user's own filter must not change")
            with open(path, encoding="utf-8") as handle:
                _assert("iap" in json.load(handle)["seen_defaults"], "the added default must be remembered as seen")
            # The user deletes it: it must not come back.
            client.post("/api/package-quick-filters", json={"platform": "android", "filters": kept})
            filters, _ = lc._load_package_quick_filters("android")
            _assert("iap" not in [f["id"] for f in filters], "a default the user deleted must not come back")
            lc.DEFAULT_PACKAGE_QUICK_FILTERS = original_defaults

            # Broken files are kept as a backup, never silently overwritten.
            for broken in ("{broken", json.dumps({"filters": [{"id": "x"}]})):
                for name in os.listdir(temp_dir):
                    if ".broken-" in name:
                        os.remove(os.path.join(temp_dir, name))
                with open(path, "w", encoding="utf-8") as handle:
                    handle.write(broken)
                filters, is_default = lc._load_package_quick_filters("android")
                backups = [name for name in os.listdir(temp_dir) if name.startswith("package_quick_filters_android.json.broken-")]
                _assert(is_default and not os.path.exists(path), "a broken file must fall back to defaults and move aside")
                _assert_equal(len(backups), 1, "a broken file must be kept as one backup")
                with open(os.path.join(temp_dir, backups[0]), encoding="utf-8") as handle:
                    _assert_equal(handle.read(), broken, "the backup must hold the original bytes")

            # Export writes a file (no app window here, so into <data dir>/exports).
            _assert_equal(client.post("/api/package-quick-filters/export", json={"platform": "android", "filters": [dict(mine, name="")]}).status_code, 400, "invalid filters must not be exported")
            exported = client.post("/api/package-quick-filters/export", json={"platform": "android", "filters": kept}).get_json()
            _assert(exported["ok"] and exported["path"].startswith(os.path.join(temp_dir, "exports")), f"export failed: {exported}")
            with open(exported["path"], encoding="utf-8") as handle:
                data = json.load(handle)
            _assert_equal((data["kind"], data["platform"], data["filters"]), (lc.PACKAGE_QUICK_FILTER_EXPORT_KIND, "android", kept), "exported file content is wrong")

            # In a browser the page saves the file itself: the API only returns the content.
            exports_before = sorted(os.listdir(os.path.join(temp_dir, "exports")))
            download = client.post("/api/package-quick-filters/export", json={"platform": "android", "filters": kept, "download": True}).get_json()
            _assert(download["ok"] and download["filename"].startswith("quick_filters_android_") and download["filename"].endswith(".json"), f"download export failed: {download}")
            _assert_equal(json.loads(download["content"]), data, "download content must equal the exported file")
            _assert_equal(sorted(os.listdir(os.path.join(temp_dir, "exports"))), exports_before, "download mode must not write a file on the server")
            html = client.get("/").get_data(as_text=True)
            _assert("var packageNativeFileDialog = false;" in html, "without an app window the page must use the browser save path")
            _assert("window.showSaveFilePicker" in html and "download: true" in html, "browser export must let the user choose where to save")

            # Import validates and returns the list; it does not save by itself.
            _assert(not os.path.exists(path), "precondition: nothing saved")
            imported = client.post("/api/package-quick-filters/import", json={"platform": "android", "data": data}).get_json()
            _assert(imported["ok"] and imported["filters"] == kept, "import must return the exported filters")
            _assert(not os.path.exists(path), "import must not save until the user presses Save")
            wrong = client.post("/api/package-quick-filters/import", json={"platform": "ios", "data": data})
            _assert_equal((wrong.status_code, wrong.get_json()["error"]), (400, "wrong_platform"), "an Android file must not import into iOS")
            _assert_equal(client.post("/api/package-quick-filters/import", json={"platform": "android", "data": {"x": 1}}).status_code, 400, "a non-filter file must be refused")
            ios_with_tag = {"platform": "ios", "filters": [dict(mine, tag="Unity")]}
            _assert_equal(client.post("/api/package-quick-filters/import", json={"platform": "ios", "data": ios_with_tag}).status_code, 400, "iOS import must refuse a tag")
        finally:
            lc.DEFAULT_PACKAGE_QUICK_FILTERS = original_defaults
            if original_data_dir is None:
                os.environ.pop("EVENTINSPECTOR_DATA_DIR", None)
            else:
                os.environ["EVENTINSPECTOR_DATA_DIR"] = original_data_dir


def test_package_quick_filter_history_query() -> None:
    """Recorded logs, session search and export apply the same quick filter as the live view."""
    original_data_dir = os.environ.get("EVENTINSPECTOR_DATA_DIR")
    with tempfile.TemporaryDirectory() as temp_dir, _TempPackageDb() as db:
        os.environ["EVENTINSPECTOR_DATA_DIR"] = temp_dir
        try:
            lc._init_package_log_db()
            conn = db.connect()
            session_id = conn.execute(
                "INSERT INTO package_log_sessions (package_id, started_at, ended_at, status) VALUES ('qf', ?, ?, 'stopped')",
                (time.time(), time.time()),
            ).lastrowid
            rows = [
                ("Unity", "[Ad,RewardedBidding, MAX] Raise"),
                ("Unity", "[Ad,InterstitialBidding] Raise"),
                ("AppLovinSdk", "[TaskCollectSignals] bidding token"),
                ("Adjust", "event tracked"),
                ("Unity", "100%_off [IAP] Entitlement adjust"),
                ("", "[Adjust]: Event tracked"),
            ]
            for tag, message in rows:
                conn.execute(
                    "INSERT INTO package_log_entries (session_id, created_at, time_display, device_id, device_name, level, tag, message, raw_log) VALUES (?, ?, '10:00:00.000', 'd', 'Dev', 'I', ?, ?, ?)",
                    (session_id, time.time(), tag, message, f"10:00:00.000 I {tag}: {message}"),
                )
            conn.commit()
            conn.close()
            client = lc.app.test_client()

            def messages(qf, qp="android"):
                # The session's own platform picks the quick filter list.
                db_conn = db.connect()
                db_conn.execute("UPDATE package_log_sessions SET platform = ? WHERE id = ?", (qp, session_id))
                db_conn.commit()
                db_conn.close()
                data = client.get(f"/api/package-log/logs?session_id={session_id}&qf={qf}&qp={qp}").get_json()
                _assert(data["ok"], f"history query failed for {qf}: {data}")
                return sorted(row["message"] for row in data["rows"])

            _assert_equal(messages("bidding"), ["[Ad,InterstitialBidding] Raise", "[Ad,RewardedBidding, MAX] Raise"], "Bidding must need [ad, and bidding")
            _assert_equal(messages("adjust"), ["event tracked"], "Android Adjust must match the tag only")
            _assert_equal(messages("adjust", "ios"), ["[Adjust]: Event tracked"], "iOS Adjust must search the whole log for adjust]")
            _assert_equal(messages("applovinsdk"), ["[TaskCollectSignals] bidding token"], "Applovinsdk must search the whole log")

            def typed(**query):
                params = "&".join(f"{key}={value}" for key, value in query.items())
                data = client.get(f"/api/package-log/logs?session_id={session_id}&{params}").get_json()
                return sorted(row["message"] for row in data["rows"])

            _assert_equal(typed(q1="raise"), ["[Ad,InterstitialBidding] Raise", "[Ad,RewardedBidding, MAX] Raise"], "Filter 1 search is wrong")
            _assert_equal(typed(q1="raise", q2="bidding", q4="max"), ["[Ad,RewardedBidding, MAX] Raise"], "Filter 3 (q4) must narrow Filter 1 + 2")
            _assert_equal(typed(q4="tracked"), ["[Adjust]: Event tracked", "event tracked"], "Filter 3 alone must search the whole log")
            two_boxes = [
                {"id": "two", "name": "Two", "tag": "", "message1": "raise", "message2": "rewarded"},
                {"id": "three", "name": "Three", "tag": "", "message1": "raise", "message2": "bidding", "message3": "interstitial"},
            ]
            client.post("/api/package-quick-filters", json={"platform": "android", "filters": two_boxes})
            _assert_equal(messages("two"), ["[Ad,RewardedBidding, MAX] Raise"], "Filter 2 must narrow the Filter 1 result")
            _assert_equal(messages("three"), ["[Ad,InterstitialBidding] Raise"], "quick filter Filter 3 must narrow Filter 1 + 2")
            _assert_equal(client.get(f"/api/package-log/logs?session_id={session_id}&qf=missing").status_code, 400, "unknown quick filter must be refused")
            custom = [
                {"id": "pct", "name": "Pct", "tag": "", "message1": "100%_off", "message2": ""},
                {"id": "iap", "name": "IAP", "tag": "UNITY", "message1": "[IAP]", "message2": ""},
                {"id": "iap_lower", "name": "iap", "tag": "", "message1": "[iap]", "message2": "ENTITLEMENT"},
            ]
            _assert(client.post("/api/package-quick-filters", json={"platform": "android", "filters": custom}).get_json()["ok"], "custom filters were refused")
            _assert_equal(messages("pct"), ["100%_off [IAP] Entitlement adjust"], "filter text must be literal, not LIKE wildcards")
            _assert_equal(messages("iap"), ["100%_off [IAP] Entitlement adjust"], "tag + filter must ignore case")
            _assert_equal(messages("iap_lower"), ["100%_off [IAP] Entitlement adjust"], "filters must never be case sensitive")
            db_conn = db.connect()
            db_conn.execute("UPDATE package_log_sessions SET platform = 'ios' WHERE id = ?", (session_id,))
            db_conn.commit()
            db_conn.close()
            _assert_equal(client.get(f"/api/package-log/logs?session_id={session_id}&qf=pct&qp=android").status_code, 400, "an Android filter must not exist in an iOS session's list")
        finally:
            if original_data_dir is None:
                os.environ.pop("EVENTINSPECTOR_DATA_DIR", None)
            else:
                os.environ["EVENTINSPECTOR_DATA_DIR"] = original_data_dir


def test_price_rotation_exact_parser() -> None:
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    _assert('onclick="switchTab(\'PriceRotation\')">Bidding</button>' in source_text, "Price Rotation tab label must be Bidding")
    _assert(source_text.count('id="priceRotationTypeMax"') == 1, "MAX filter must exist exactly once")
    _assert('value="max"' in source_text, "MAX filter value is missing")
    _assert('id="priceRotationTypeAscendx"' not in source_text, "Ascendx Bidding filter must be hidden")
    _assert('id="priceRotationTypeCloudx"' not in source_text, "Cloudx Bidding filter must be hidden")
    _assert("rotationType === 'ascendx' || rotationType === 'cloudx'" in source_text, "hidden Bidding types must stay filtered in the UI")
    _assert(source_text.count('id="priceRotationTypeWaterfall"') == 1, "Waterfall filter must exist exactly once")
    _assert(source_text.count('id="priceRotationTypeInterstitialCap"') == 1, "InterstitialCap filter must exist exactly once")
    _assert(source_text.count('id="priceRotationAdTypeAll"') == 1, "Ad type All filter must exist exactly once")
    _assert(source_text.count('id="priceRotationAdTypeRewarded"') == 1, "Rewarded ad type filter must exist exactly once")
    _assert(source_text.count('id="priceRotationAdTypeInterstitial"') == 1, "Interstitial ad type filter must exist exactly once")
    _assert('value="rewarded"' in source_text and 'value="interstitial"' in source_text, "Rewarded/Interstitial filters are missing")
    _assert('value="waterfall"' in source_text, "Waterfall filter value is missing")
    valid = 'Unity : [Ad,RewardedBidding, CloudX] Raise: {"act":"demandFailedObserved","error":"WinnerBid not found"}'
    parsed = lc._parse_price_rotation_log(valid, "device-price")
    _assert(parsed is not None, "Price Rotation parser must accept the exact marker")
    _assert_equal(parsed["type"], "CloudX", "Price Rotation type should come from the marker")
    _assert('"act": "demandFailedObserved"' in parsed["details"], "Price Rotation details should format JSON")
    waterfall = lc._parse_price_rotation_log('Unity : [Ad,RewardedBidding, Waterfall] Raise: {"act":"bid"}', "device-price")
    _assert(waterfall is not None, "Price Rotation parser must accept the Waterfall marker")
    _assert_equal(waterfall["type"], "Waterfall", "Waterfall type should come from the marker")
    cap = lc._parse_price_rotation_log('Unity : [Ad,RewardedBidding] RewardedCapService: {"act":"cap"}', "device-price")
    _assert(cap is not None, "Price Rotation parser must accept the RewardedCapService marker")
    _assert_equal(cap["type"], "RewardedCap", "RewardedCapService must use the RewardedCap type")
    plain = lc._parse_price_rotation_log('Unity : [Ad,RewardedBidding] Raise: {"act":"bid"}', "device-price")
    _assert(plain is not None, "Price Rotation parser must accept a plain RewardedBidding marker")
    _assert_equal(plain["type"], "RewardedBidding", "A plain marker must use the RewardedBidding type")
    interstitial_raw = 'Unity : [Ad,InterstitialBidding, LevelPlay] Raise: {"act":"bid"}'
    interstitial = lc._parse_price_rotation_log(interstitial_raw, "device-price")
    _assert(interstitial is not None, "Price Rotation parser must accept the InterstitialBidding marker")
    _assert_equal(interstitial["type"], "LevelPlay", "Interstitial type should come from the marker")
    interstitial_cap = lc._parse_price_rotation_log('Unity : [Ad,InterstitialBidding] InterstitialCapService: {"act":"cap"}', "device-price")
    _assert(interstitial_cap is not None, "Price Rotation parser must accept the InterstitialCapService marker")
    _assert_equal(interstitial_cap["type"], "InterstitialCap", "InterstitialCapService must use the InterstitialCap type")
    interstitial_plain = lc._parse_price_rotation_log('Unity : [Ad,InterstitialBidding] Raise: {"act":"bid"}', "device-price")
    _assert(interstitial_plain is not None, "Price Rotation parser must accept a plain InterstitialBidding marker")
    _assert_equal(interstitial_plain["type"], "InterstitialBidding", "A plain marker must use the InterstitialBidding type")
    _assert(lc._parse_price_rotation_log(valid.replace("[Ad,RewardedBidding,", "[Ad,RewardedBiddingX,"), "device-price") is None, "Price Rotation marker must be exact")
    _assert(lc._parse_price_rotation_log(interstitial_raw.replace("[Ad,InterstitialBidding,", "[Ad,InterstitialBiddingX,"), "device-price") is None, "Interstitial marker must be exact")

    original_paused = lc.is_paused
    original_emit = lc.socketio.emit
    original_rows = list(lc.price_rotation_logs)
    try:
        lc.is_paused = False
        lc.price_rotation_logs.clear()
        lc.socketio.emit = lambda *_args, **_kwargs: None
        max_rewarded = lc._parse_price_rotation_log(
            'Unity : [Ad,RewardedBidding, MAX] Raise: {"act":"bid"}',
            "device-price",
        )
        max_rewarded_applovin = lc._parse_price_rotation_log(
            'Unity : [Ad,RewardedBidding, applovin] Raise: {"act":"bid"}',
            "device-price",
        )
        max_interstitial = lc._parse_price_rotation_log(
            'Unity : [Ad,InterstitialBidding, MAX] Raise: {"act":"bid"}',
            "device-price",
        )
        max_interstitial_applovin = lc._parse_price_rotation_log(
            'Unity : [Ad,InterstitialBidding, applovin] Raise: {"act":"bid"}',
            "device-price",
        )
        _assert(
            all(parsed and parsed["type"] in {"MAX", "applovin"} for parsed in (
                max_rewarded,
                max_rewarded_applovin,
                max_interstitial,
                max_interstitial_applovin,
            )),
            "MAX markers must parse for Rewarded and Interstitial bidding",
        )
        for raw_log in (
            'Unity : [Ad,RewardedBidding, MAX] Raise: {"act":"bid"}',
            'Unity : [Ad,RewardedBidding, applovin] Raise: {"act":"bid"}',
            'Unity : [Ad,InterstitialBidding, MAX] Raise: {"act":"bid"}',
            'Unity : [Ad,InterstitialBidding, applovin] Raise: {"act":"bid"}',
        ):
            lc.process_price_rotation_log(raw_log, "device-price")
        _assert_equal(len(lc.price_rotation_logs), 4, "MAX Bidding variants must remain visible")

        lc.process_price_rotation_log(valid, "device-price")
        _assert_equal(len(lc.price_rotation_logs), 4, "hidden Cloudx Bidding logs must not be stored")
        lc.process_price_rotation_log(
            'Unity : [Ad,InterstitialBidding, Ascendx] Raise: {"act":"bid"}',
            "device-price",
        )
        _assert_equal(len(lc.price_rotation_logs), 4, "hidden Ascendx Bidding logs must not be stored")
    finally:
        lc.is_paused = original_paused
        lc.socketio.emit = original_emit
        lc.price_rotation_logs.clear()
        lc.price_rotation_logs.extend(original_rows)


def test_load_ads_provider_contract() -> None:
    original_platform = lc.active_platform
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_emit = lc.socketio.emit
    try:
        lc.active_platform = "ios"
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": None})
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        emitted = []
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))

        ios_line = (
            'Appmetrica TrackAdRevenueEvent: '
            '{"AdRevenueValue":"0.01","Currency":"USD","AdNetwork":"unityads",'
            '"AdType":2,"Payload":{"ad_platform":"cloudx",'
            '"ad_network":"unityads","ad_format":"interstitial"},"Precision":null}'
        )
        lc.process_load_ads_ext_log(ios_line, "ios-load-ads")
        _assert_equal(len(lc.load_ads_ext_events), 1, "iOS AppMetrica Load Ads row was not recorded")
        _assert_equal(lc.load_ads_ext_events[0].get("provider"), "cloudx", "iOS provider was not extracted")

        # Provider is not unique: the same provider may appear for another network.
        second_ios_line = ios_line.replace('"AdNetwork":"unityads"', '"AdNetwork":"applovin"')
        second_ios_line = second_ios_line.replace('"ad_network":"unityads"', '"ad_network":"applovin"')
        lc.process_load_ads_ext_log(second_ios_line, "ios-load-ads")
        _assert_equal(len(lc.load_ads_ext_events), 2, "duplicate Provider must remain visible for another network")
        _assert(all(row.get("provider") == "cloudx" for row in lc.load_ads_ext_events), "provider values changed unexpectedly")

        # Same network/format from another provider is a distinct record.
        third_ios_line = ios_line.replace('"ad_platform":"cloudx"', '"ad_platform":"ascendx"')
        lc.process_load_ads_ext_log(third_ios_line, "ios-load-ads")
        _assert_equal(len(lc.load_ads_ext_events), 3, "different providers must not collide in dedup")
        _assert_equal(lc.load_ads_ext_events[-1].get("provider"), "ascendx", "second provider was not extracted")

        # Android's AdRevenue{...} parser must use the nested ad_platform too.
        lc.active_platform = "android"
        android_line = (
            'AdRevenue Received: AdRevenue{adType=rewarded,adNetwork=ironsource,'
            'payload={"ad_platform":"ironsource","ad_network":"ironsource",'
            '"ad_format":"rewarded"}}'
        )
        lc.process_load_ads_ext_log(android_line, "android-load-ads")
        android_rows = [row for row in lc.load_ads_ext_events if row.get("device_id") == "android-load-ads"]
        _assert_equal(len(android_rows), 1, "Android AppMetrica Load Ads row was not recorded")
        _assert_equal(android_rows[0].get("provider"), "ironsource", "Android provider was not extracted")

        rendered = lc.app.test_client().get("/").get_data(as_text=True)
        _assert('>Load Ads</button>' in rendered, "Load Ads tab label was not updated")
        _assert("Record Load Ads:" in rendered, "Load Ads recording label was not updated")
        _assert('>Provider</th>' in rendered, "Load Ads Provider column is missing")
        _assert('"provider": _normalize_load_ads_provider(provider)' in (ROOT / "Log_checker.py").read_text(encoding="utf-8"), "Sheet provider payload is missing")
    finally:
        lc.active_platform = original_platform
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.socketio.emit = original_emit


def test_load_ads_max_contract() -> None:
    """MAX Load Ads rows must come from the exact MAX revenue callback fields."""
    original_platform = lc.active_platform
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_emit = lc.socketio.emit
    original_send_to_sheet = lc.send_to_sheet
    original_rows = list(lc.load_ads_ext_events)
    original_unique = set(lc.unique_load_ads_ext)
    emitted = []
    sheet_rows = []
    try:
        lc.active_platform = "android"
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": "NG381"})
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()
        lc.send_to_sheet = lambda *args: sheet_rows.append(args)

        samples = [
            "09-03 11:39:34.462\tcom.indiez.nonogram\t13912\t-\tD\tAppLovinSdk\t [MaxInterstitialAd] MaxAdRevenueListener.onAdRevenuePaid(ad=MediatedAd{thirdPartyAdPlacementId=945322059, adUnitId=test_mode_interstitial, format=INTER, networkName='Pangle'}), listener=null",
            "09-03 11:40:05.062\tcom.indiez.nonogram\t13912\t-\tD\tAppLovinSdk\t [MaxRewardedAd] MaxAdRevenueListener.onAdRevenuePaid(ad=MediatedAd{thirdPartyAdPlacementId=945322060, adUnitId=test_mode_rewarded_TIKTOK_BIDDING, format=REWARDED, networkName='Pangle'}), listener=null",
            "09-03 11:41:13.494\tcom.indiez.nonogram\t13912\t-\tD\tAppLovinSdk\t [MaxAdView] MaxAdRevenueListener.onAdRevenuePaid(ad=MediatedAd{thirdPartyAdPlacementId=980040498, adUnitId=test_mode_banner, format=BANNER, networkName='Pangle'}), listener=null",
            "09-03 11:41:23.051\tcom.indiez.nonogram\t13912\t-\tD\tAppLovinSdk\t [MaxAdView] MaxAdRevenueListener.onAdRevenuePaid(ad=MediatedAd{thirdPartyAdPlacementId=980040498, adUnitId=test_mode_banner, format=BANNER, networkName='Pangle'}), listener=null",
            "09-03 11:41:35.454\tcom.indiez.nonogram\t13912\t-\tD\tAppLovinSdk\t [MaxAdView] MaxAdRevenueListener.onAdRevenuePaid(ad=MediatedAd{thirdPartyAdPlacementId=982160535, adUnitId=test_mode_mrec, format=MREC, networkName='Pangle'}), listener=null",
        ]
        for line in samples:
            lc.process_load_ads_max_log(line, "device-max")

        rows = list(lc.load_ads_ext_events)
        _assert_equal(len(rows), 4, "MAX Load Ads must keep one row per network/format")
        _assert_equal({row.get("ad_format") for row in rows}, {"INTER", "REWARDED", "BANNER", "MREC"}, "MAX formats were not parsed from MediatedAd")
        _assert(all(row.get("provider") == "MAX" for row in rows), "MAX Load Ads provider was not recorded")
        _assert(all(row.get("ad_network") == "Pangle" for row in rows), "MAX networkName was not parsed from the raw log")
        _assert(all(lc.MAX_LOAD_ADS_REVENUE_KEYWORD in row.get("raw_log", "") for row in rows), "MAX raw callback log was not preserved")
        _assert_equal(len(_live_rows_sent(emitted, "load_ads_ext")), 4, "MAX Load Ads rows were not emitted to the UI")
        _assert(not _live_rows_sent(emitted, "load_ads"), "MAX rows were emitted to the hidden Load Ads table")
        _assert_equal(len(sheet_rows), 4, "MAX Load Ads rows were not sent to the sheet")
        _assert(all(args[4:] == ("LoadAdsExt", "MAX") for args in sheet_rows), "MAX sheet payload used the wrong provider or tab")

        # Non-MAX diagnostics and disabled recording must never create rows.
        lc.process_load_ads_max_log("AppLovinSdk unrelated diagnostic format=BANNER networkName='Pangle'", "device-max")
        _assert_equal(len(lc.load_ads_ext_events), 4, "unrelated AppLovin diagnostics must be ignored")
        lc.recording_states["LoadAdsExt"]["is_recording"] = False
        lc.process_load_ads_max_log(samples[0], "device-max-2")
        _assert_equal(len(lc.load_ads_ext_events), 4, "MAX rows must respect Load Ads recording state")

        rendered = lc.app.test_client().get("/").get_data(as_text=True)
        _assert(rendered.count(">Provider</th>") >= 2, "Provider column must be present in both Load Ads tables")
    finally:
        lc.active_platform = original_platform
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.load_ads_ext_events.clear()
        lc.load_ads_ext_events.extend(original_rows)
        lc.unique_load_ads_ext.clear()
        lc.unique_load_ads_ext.update(original_unique)
        lc.socketio.emit = original_emit
        lc.send_to_sheet = original_send_to_sheet


def test_ios_load_ads_max_viewability_contract() -> None:
    """iOS MAX viewability dictionaries must become visible Load Ads rows."""
    original_platform = lc.active_platform
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_emit = lc.socketio.emit
    original_send_to_sheet = lc.send_to_sheet
    original_rows = list(lc.load_ads_ext_events)
    original_unique = set(lc.unique_load_ads_ext)
    original_buffer = dict(lc.incomplete_ios_max_load_ads_logs)
    emitted = []
    sheet_rows = []
    try:
        lc.active_platform = "ios"
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": "NG381"})
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()
        lc.send_to_sheet = lambda *args: sheet_rows.append(args)

        lines = [
            "17:17:10 PixelArt [AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting mediated_ad_viewability_impression_called with extra parameters {",
            '    "ad_event_id" = "3a4fb284-4236-406f-a774-c763f28c3e5b";',
            '    "ad_format" = BANNER;',
            '    "ad_unit_id" = "test_mode_banner";',
            '    "adapter_class" = ALChartboostMediationAdapter;',
            '    "adapter_version" = "9.14.0.1";',
            '    "network_name" = "CHARTBOOST_NETWORK";',
            "}",
        ]
        for line in lines:
            lc.process_load_ads_max_log(line, "ios-max")

        rows = list(lc.load_ads_ext_events)
        _assert_equal(len(rows), 1, "iOS MAX viewability log did not create a Load Ads row")
        _assert_equal(rows[0].get("provider"), "MAX", "iOS MAX provider was not recorded")
        _assert_equal(rows[0].get("ad_network"), "CHARTBOOST_NETWORK", "iOS MAX network_name was not parsed")
        _assert_equal(rows[0].get("ad_format"), "BANNER", "iOS MAX ad_format was not parsed")
        _assert_equal(rows[0].get("adapter_class"), "ALChartboostMediationAdapter", "iOS MAX adapter class was not parsed")
        _assert_equal(rows[0].get("adapter_version"), "9.14.0.1", "iOS MAX adapter version was not parsed")
        _assert(lc.MAX_LOAD_ADS_VIEWABILITY_KEYWORD in rows[0].get("raw_log", ""), "iOS MAX raw multiline log was not preserved")
        _assert(_live_rows_sent(emitted, "load_ads_ext"), "iOS MAX row was not emitted to Load Ads")
        _assert_equal(sheet_rows[0][4:], ("LoadAdsExt", "MAX"), "iOS MAX sheet payload used the wrong tab or provider")
        _assert_equal(lc.incomplete_ios_max_load_ads_logs, {}, "completed iOS MAX log left a stale buffer")
    finally:
        lc.active_platform = original_platform
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.load_ads_ext_events.clear()
        lc.load_ads_ext_events.extend(original_rows)
        lc.unique_load_ads_ext.clear()
        lc.unique_load_ads_ext.update(original_unique)
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.incomplete_ios_max_load_ads_logs.update(original_buffer)
        lc.socketio.emit = original_emit
        lc.send_to_sheet = original_send_to_sheet


def test_ios_load_ads_max_viewability_record_boundary_contract() -> None:
    """A following iOS syslog record must not be appended to a MAX row."""
    original_platform = lc.active_platform
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_emit = lc.socketio.emit
    original_send_to_sheet = lc.send_to_sheet
    original_rows = list(lc.load_ads_ext_events)
    original_unique = set(lc.unique_load_ads_ext)
    original_buffer = dict(lc.incomplete_ios_max_load_ads_logs)
    try:
        lc.active_platform = "ios"
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": "NG381"})
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.socketio.emit = lambda *_args: None
        lc.send_to_sheet = lambda *_args: None

        combined = (
            "Sep 23 14:00:48 iPhone-11-pro PixelArt(AppLovinSDK)[2113] <Notice>: "
            "[AppLovinSdk] DEBUG [ALHealthEventsReporter] Reporting "
            "mediated_ad_viewability_impression_called with extra parameters {\n"
            '    "ad_format" = BANNER;\n'
            '    "adapter_class" = ALAppLovinMediationAdapter;\n'
            '    "adapter_version" = "13.6.4";\n'
            '    "network_name" = "APPLOVIN_EXCHANGE";\n'
            "}\n"
            "Sep 23 14:00:48 iPhone-11-pro PixelArt(AppLovinSDK)[2113] <Notice>: "
            "[ALPersistentPostbackManager] Enqueued postback: <ALPersistentPostback: "
            "uniqueIdentifier = D5D00408-C901-43E1-A632-D7B4B308E82E>\n"
            "Sep 23 14:00:48 iPhone-11-pro PixelArt(UIKitCore)[2113] <Notice>: "
            "Will add background task with taskName: <private>"
        )
        lc.process_load_ads_max_log(combined, "ios-max-boundary")

        rows = list(lc.load_ads_ext_events)
        _assert_equal(len(rows), 1, "combined iOS MAX chunk created the wrong row count")
        raw_log = rows[0].get("raw_log", "")
        _assert("ALPersistentPostbackManager" not in raw_log, "postback record leaked into MAX Load Ads")
        _assert("UIKitCore" not in raw_log, "UIKitCore record leaked into MAX Load Ads")
        _assert(raw_log.rstrip().endswith("}"), "MAX Load Ads row did not stop at its closing dictionary")
        _assert_equal(lc.incomplete_ios_max_load_ads_logs, {}, "record boundary test left a stale MAX buffer")
    finally:
        lc.active_platform = original_platform
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.load_ads_ext_events.clear()
        lc.load_ads_ext_events.extend(original_rows)
        lc.unique_load_ads_ext.clear()
        lc.unique_load_ads_ext.update(original_unique)
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.incomplete_ios_max_load_ads_logs.update(original_buffer)
        lc.socketio.emit = original_emit
        lc.send_to_sheet = original_send_to_sheet


def test_ios_load_ads_max_delegate_contract() -> None:
    """iOS MAX didDisplayAd/didPayRevenueForAd callbacks must become Load Ads rows."""
    original_platform = lc.active_platform
    original_recording_state = dict(lc.recording_states["LoadAdsExt"])
    original_emit = lc.socketio.emit
    original_send_to_sheet = lc.send_to_sheet
    original_rows = list(lc.load_ads_ext_events)
    original_unique = set(lc.unique_load_ads_ext)
    original_buffer = dict(lc.incomplete_ios_max_load_ads_logs)
    emitted = []
    sheet_rows = []
    try:
        lc.active_platform = "ios"
        lc.recording_states["LoadAdsExt"].update({"is_recording": True, "current_sheet": "NG381"})
        lc.load_ads_ext_events.clear()
        lc.unique_load_ads_ext.clear()
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()
        lc.send_to_sheet = lambda *args: sheet_rows.append(args)

        samples = [
            "16 17:12 Galaxy-iphone PixelArt(AppLovinSDK)[42329] <Notice>: [AppLovinSdk] DEBUG [MAAdView] -[MAAdDelegate didDisplayAd: [ALMediatedAdViewAd thirdPartyAdPlacementIdentifier=mrec_regular, adUnitIdentifier=043278b2d7a2125b, format=MREC, networkName=APPLOVIN_EXCHANGE]], delegate: <SMLDelegate: 0x15f5b0190>",
            "16 19:22 Galaxy-iphone PixelArt(AppLovinSDK)[42329] <Notice>: [AppLovinSdk] DEBUG [MAInterstitialAd] -[MAAdDelegate didDisplayAd: [ALMediatedFullscreenAd thirdPartyAdPlacementIdentifier=inter_regular, adUnitIdentifier=5a0e30d201088d93, format=INTER, networkName=AppLovin]], delegate: <SMLDelegate: 0x166e5ec10",
            "16 21:39 Galaxy-iphone PixelArt(AppLovinSDK)[42329] <Notice>: [AppLovinSdk] DEBUG [MARewardedAd] -[MAAdDelegate didDisplayAd: [ALMediatedFullscreenAd thirdPartyAdPlacementIdentifier=inter_videoa, adUnitIdentifier=8738b82fc7f2f9b5, format=REWARDED, networkName=AppLovin]], delegate: <SMLDelegate: 0x166e5caf0>",
            "16 27:54 Galaxy-iphone PixelArt(AppLovinSDK)[42329] <Notice>: [AppLovinSdk] DEBUG [MAAdView] -[MAAdRevenueDelegate didPayRevenueForAd: [ALMediatedAdViewAd thirdPartyAdPlacementIdentifier=StR7aygvGBRCdKBc, adUnitIdentifier=4ffa27b3e6066b41, format=BANNER, networkName=Moloco]], delegate: <MAUnityAdManager: 0x14f40d180>",
        ]
        for line in samples:
            lc.process_load_ads_max_log(line, "ios-max-delegate")

        rows = list(lc.load_ads_ext_events)
        _assert_equal(len(rows), 4, "iOS MAX delegate callbacks did not create four Load Ads rows")
        _assert_equal(
            {(row.get("ad_network"), row.get("ad_format")) for row in rows},
            {
                ("APPLOVIN_EXCHANGE", "MREC"),
                ("AppLovin", "INTER"),
                ("AppLovin", "REWARDED"),
                ("Moloco", "BANNER"),
            },
            "iOS MAX delegate fields were not parsed",
        )
        _assert(all(row.get("provider") == "MAX" for row in rows), "iOS MAX delegate provider was not recorded")
        _assert(
            all(any(keyword in row.get("raw_log", "") for keyword in lc.MAX_LOAD_ADS_IOS_DELEGATE_KEYWORDS) for row in rows),
            "iOS MAX delegate raw logs were not preserved",
        )
        _assert_equal(len(_live_rows_sent(emitted, "load_ads_ext")), 4, "iOS MAX delegate rows were not emitted to the UI")
        _assert(not _live_rows_sent(emitted, "load_ads"), "iOS MAX delegate rows used the wrong UI event")
        _assert_equal(len(sheet_rows), 4, "iOS MAX delegate rows were not sent to the sheet")
        _assert(all(args[4:] == ("LoadAdsExt", "MAX") for args in sheet_rows), "iOS MAX delegate sheet target was wrong")

        duplicate = samples[0].replace("17:12", "17:13")
        lc.process_load_ads_max_log(duplicate, "ios-max-delegate")
        _assert_equal(len(lc.load_ads_ext_events), 4, "duplicate iOS MAX delegate callback was recorded twice")
    finally:
        lc.active_platform = original_platform
        lc.recording_states["LoadAdsExt"].clear()
        lc.recording_states["LoadAdsExt"].update(original_recording_state)
        lc.load_ads_ext_events.clear()
        lc.load_ads_ext_events.extend(original_rows)
        lc.unique_load_ads_ext.clear()
        lc.unique_load_ads_ext.update(original_unique)
        lc.incomplete_ios_max_load_ads_logs.clear()
        lc.incomplete_ios_max_load_ads_logs.update(original_buffer)
        lc.socketio.emit = original_emit
        lc.send_to_sheet = original_send_to_sheet


def test_levelplay_impression_data_callback_contract() -> None:
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    for marker in (
        "LevelPlayAdImpressionEventMapper->Map",
        "LevelPlayImpressionData - {ad_unit_name}",
        'return "Banner"',
        'return "Rewarded"',
        'return "Interstitial"',
    ):
        _assert(marker in source_text, f"LevelPlay callback contract is missing: {marker}")

    original_paused = lc.is_paused
    original_emit = lc.socketio.emit
    original_rows = list(lc.callback_ad_logs)
    emitted = []
    try:
        lc.is_paused = False
        lc.callback_ad_logs.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()

        ad_units = {
            "Banner": {"mediationAdUnitName": "Banner", "adUnit": "banner", "adFormat": "banner"},
            "Rewarded": {"mediationAdUnitName": "Rewarded", "adUnit": "rewarded_video", "adFormat": "rewarded_video"},
            "Interstitial": {"mediationAdUnitName": "Interstitial", "adUnit": "interstitial", "adFormat": "interstitial"},
        }
        for expected_label, impression_data in ad_units.items():
            nested_json = json.dumps(impression_data, separators=(",", ":"))
            outer_json = json.dumps({"LevelPlayImpressionData": nested_json}, separators=(",", ":"))
            line = f"Unity LevelPlayAdImpressionEventMapper->Map: {outer_json}"
            lc.process_callback_and_ad_event_log(line, "device-levelplay")

        rows = list(lc.callback_ad_logs)
        _assert_equal(len(rows), 3, "LevelPlay impression mapper should add one callback row per ad unit")
        _assert_equal(
            [row["event_name"] for row in rows],
            ["LevelPlayImpressionData - Banner", "LevelPlayImpressionData - Rewarded", "LevelPlayImpressionData - Interstitial"],
            "LevelPlay impression callback names do not include the ad unit",
        )
        _assert(all(row["type"] == "Callback Levelplay" for row in rows), "LevelPlay impression mapper rows must use the LevelPlay callback filter")
        _assert(all("LevelPlayAdImpressionEventMapper->Map" in row["raw_log"] for row in rows), "LevelPlay impression raw logs were not preserved")
        _assert_equal(len(_live_rows_sent(emitted, "callback")), 3, "LevelPlay impression callbacks were not emitted to the Callback tab")

        # Existing LevelPlay listener callbacks must keep their old display name.
        lc.process_callback_and_ad_event_log("Unity LevelPlayInterstitialAdListener", "device-levelplay")
        _assert_equal(lc.callback_ad_logs[-1]["event_name"], "Interstitial", "existing LevelPlay listener name changed")
    finally:
        lc.is_paused = original_paused
        lc.callback_ad_logs.clear()
        lc.callback_ad_logs.extend(original_rows)
        lc.socketio.emit = original_emit


def test_ascendx_cloudx_callback_contract() -> None:
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    required_markers = (
        "_OnAdLoadedEvent",
        "_OnAdImpressionEvent",
        "_OnAdDisplayedEvent",
        "_OnAdClickedEvent",
        "_OnAdClosedEvent",
        "_OnAdReceivedRewardEvent",
        "OnAdLoadFailed",
        "OnAdLoadSuccess",
        "OnAdRevenuePaid",
        "OnAdClicked",
        "OnAdShowSuccess",
        "OnAdRewarded",
        "OnAdHidden",
        'value="ascendx_callback"',
        'value="cloudx_callback"',
        "Callback Ascendx",
        "Callback Cloudx",
        "Callback Levelplay",
        'id="callbackTypeAll" name="callbackType" type="checkbox"',
        'id="callbackTypeCallback" name="callbackType" type="checkbox"',
        'id="callbackTypeAscendx" name="callbackType" type="checkbox"',
        'id="callbackTypeCloudx" name="callbackType" type="checkbox"',
        'id="callbackTypeGadsme" name="callbackType" type="checkbox"',
        'id="callbackTypeAdverty5" name="callbackType" type="checkbox"',
        "selectedTypeFilters",
        "bindCallbackCheckboxGroup",
        'details-cell text-sm align-top"><div class="max-h-64 overflow-auto">${res.details}',
        'log-cell text-xs font-normal text-gray-600 align-top"><div class="max-h-64 overflow-auto whitespace-pre-wrap break-all">',
    )
    for marker in required_markers:
        _assert(marker in source_text, f"AscendX/CloudX callback contract is missing: {marker}")

    original_paused = lc.is_paused
    original_emit = lc.socketio.emit
    original_rows = list(lc.callback_ad_logs)
    emitted = []
    try:
        lc.is_paused = False
        lc.callback_ad_logs.clear()
        lc.socketio.emit = lambda event, payload: emitted.append((event, payload))
        lc._clear_live_table_pending()

        for callback_event in lc.ASCENDX_CALLBACK_EVENTS:
            ad_unit = "Reward" if callback_event == "_OnAdReceivedRewardEvent" else "Interstitial"
            line = f"Unity [Ad,AscendX,{ad_unit}] {callback_event}"
            lc.process_callback_and_ad_event_log(line, "device-callback")

        for callback_event in lc.CLOUDX_CALLBACK_EVENTS:
            ad_unit = "Reward" if callback_event == "OnAdRewarded" else "Interstitial"
            line = f"Unity [Ad,CloudX,{ad_unit}] {callback_event}"
            lc.process_callback_and_ad_event_log(line, "device-callback")

        rows = list(lc.callback_ad_logs)
        _assert_equal(len(rows), len(lc.ASCENDX_CALLBACK_EVENTS) + len(lc.CLOUDX_CALLBACK_EVENTS), "provider callback row count changed")
        _assert_equal(
            [row["type"] for row in rows[:len(lc.ASCENDX_CALLBACK_EVENTS)]],
            ["Callback Ascendx"] * len(lc.ASCENDX_CALLBACK_EVENTS),
            "AscendX callback rows use the wrong filter type",
        )
        _assert_equal(
            [row["type"] for row in rows[len(lc.ASCENDX_CALLBACK_EVENTS):]],
            ["Callback Cloudx"] * len(lc.CLOUDX_CALLBACK_EVENTS),
            "CloudX callback rows use the wrong filter type",
        )
        _assert_equal(rows[0]["event_name"], "_OnAdLoadedEvent - Interstitial", "AscendX ad unit name was not added")
        _assert_equal(rows[5]["event_name"], "_OnAdReceivedRewardEvent - Reward", "AscendX reward name was not added")
        _assert_equal(rows[6]["event_name"], "OnAdLoadFailed - Interstitial", "CloudX ad unit name was not added")
        _assert_equal(rows[11]["event_name"], "OnAdRewarded - Reward", "CloudX reward name was not added")
        _assert(all(row["raw_log"].startswith("Unity [Ad,") for row in rows), "provider callback raw logs were not preserved")
        _assert_equal(len(_live_rows_sent(emitted, "callback")), len(rows), "provider callbacks were not emitted to the Callback tab")

        rendered = lc.app.test_client().get("/").get_data(as_text=True)
        for marker in (
            'id="callbackTypeAscendx"',
            'id="callbackTypeCloudx"',
            'value="ascendx_callback"',
            'value="cloudx_callback"',
        ):
            _assert(marker in rendered, f"rendered callback filter is missing: {marker}")
        ordered_filters = [
            rendered.index('id="callbackTypeCallback"'),
            rendered.index('id="callbackTypeAscendx"'),
            rendered.index('id="callbackTypeCloudx"'),
            rendered.index('id="callbackTypeGadsme"'),
            rendered.index('id="callbackTypeAdverty5"'),
        ]
        _assert_equal(ordered_filters, sorted(ordered_filters), "callback filters are not grouped in the requested order")
    finally:
        lc.is_paused = original_paused
        lc.callback_ad_logs.clear()
        lc.callback_ad_logs.extend(original_rows)
        lc.socketio.emit = original_emit


def test_bright_skad_ipa_inspection_contract() -> None:
    """Bright/SKAds reads SKAdNetworkItems only from the app Info.plist."""
    rendered = lc.app.test_client().get("/").get_data(as_text=True)
    _assert(">Bright/SKAds<" in rendered, "Bright tab label was not renamed")
    _assert('id="skAdNetworkList"' not in rendered, "standalone SKAdNetwork result list must be removed")
    _assert("renderSkAdNetworkIdentifiers" not in rendered, "standalone SKAdNetwork renderer must be removed")
    _assert('id="brightSkadPresetSelect"' in rendered, "Bright/SKAds preset selector is missing")
    _assert('id="brightSkadCompareBody"' in rendered, "Bright/SKAds comparison table is missing")
    _assert('id="brightSkadExtraWarning"' in rendered, "Bright/SKAds extra identifier warning is missing")
    _assert(
        rendered.index('id="brightSkadExtraWarning"') < rendered.index('id="brightSkadCompareBody"'),
        "Bright/SKAds extra identifier warning must appear before the comparison rows",
    )
    _assert("buildBrightSkadComparison" in rendered, "Bright/SKAds comparison renderer is missing")
    _assert("PASSED" in rendered, "Bright/SKAds comparison must use PASSED status")

    with tempfile.TemporaryDirectory(prefix="eventinspector_bright_skad_") as temp_dir:
        ipa_path = Path(temp_dir) / "Harness.ipa"
        app_plist = {
            "CFBundleDisplayName": "HarnessApp",
            "CFBundleName": "HarnessApp",
            "SKAdNetworkItems": [
                {"SKAdNetworkIdentifier": "app-one.skadnetwork"},
                {"SKAdNetworkIdentifier": "app-two.skadnetwork"},
            ],
        }
        bright_plist = {
            "CFBundleShortVersionString": "9.9.9",
            "SKAdNetworkItems": [
                {"SKAdNetworkIdentifier": "framework-only.skadnetwork"},
            ],
        }
        with zipfile.ZipFile(ipa_path, "w") as archive:
            archive.writestr("Payload/Harness.app/Info.plist", plistlib.dumps(app_plist))
            archive.writestr(
                "Payload/Harness.app/Frameworks/brdsdk.framework/Info.plist",
                plistlib.dumps(bright_plist),
            )

        result = lc._extract_brightsdk_version_from_ipa(str(ipa_path))
        _assert_equal(result.get("app_name"), "HarnessApp", "Bright/SKAds app name changed")
        _assert_equal(result.get("brightsdk_version"), "9.9.9", "BrightSDK version was not read")
        _assert_equal(
            result.get("skadnetwork_identifiers"),
            ["app-one.skadnetwork", "app-two.skadnetwork"],
            "SKAdNetwork list must come from the app Info.plist only",
        )
        _assert("framework-only.skadnetwork" not in result.get("skadnetwork_identifiers", []), "BrightSDK plist leaked into SKAdNetwork results")
        c192_ios = lc._load_sdk_check_presets()["C-192-iOS"]
        compare = lc._compare_bright_skad_preset(
            result.get("brightsdk_version"),
            result.get("skadnetwork_identifiers"),
            preset_name="C-192-iOS",
            preset=c192_ios,
        )
        _assert_equal(compare.get("rows", [{}])[0].get("status"), "FAILED", "BrightData mismatch must fail")
        _assert_equal(compare.get("extra_identifiers"), ["app-one.skadnetwork", "app-two.skadnetwork"], "IPA-only SKAdNetwork identifiers must be warned separately")
        with ipa_path.open("rb") as handle:
            response = lc.app.test_client().post(
                "/api/ipa/inspect",
                data={
                    "ipa_file": (handle, "Harness.ipa"),
                    "preset_name": "C-192-iOS",
                },
                content_type="multipart/form-data",
            )
        _assert_equal(response.status_code, 200, "Bright/SKAds IPA inspect endpoint failed")
        api_payload = response.get_json() or {}
        _assert_equal(api_payload.get("bright_compare", {}).get("preset_name"), "C-192-iOS", "IPA inspect did not apply the selected Bright/SKAds preset")
        _assert_equal(api_payload.get("bright_compare", {}).get("total_count"), 291, "IPA inspect Bright/SKAds compare row count changed")


def test_release_payload_sync() -> None:
    source_text = (ROOT / "Log_checker.py").read_text(encoding="utf-8", errors="ignore")
    manifest_path = ROOT / "Updates_2_5" / "remote_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    log_item = next(item for item in manifest["files"] if item.get("path") == "Log_checker.py")
    _assert_equal(_sha256_file(ROOT / "Log_checker.py"), log_item["sha256"], "source/release Log_checker.py drift detected")
    _assert("compat_sha256" not in log_item, "canonical v2.5 payload must not select the stale compatibility source")
    remote_item = next(item for item in manifest["files"] if item.get("path") == "remote_update.py")
    _assert("compat_sha256" not in remote_item, "canonical updater payload must not select the stale compatibility source")
    _assert_equal(manifest["version"], CURRENT_RELEASE_VERSION, "v2.5 release manifest version changed")

    markers = {
        "release_badge": r"v2\.5\.0\((\d+)\)",
        "html_title": rf"<title>Event Inspector v2\.5\.0\({RELEASE_SOURCE_BUILD}\)</title>",
        "socket_fallback": r"typeof window\.io === 'function'",
        "brightsdk_tab": r"switchTab\('BrightSDK'\)",
        "package_quick_select": r'id="packageQuickSelectList"',
        "check_update_call": r"result = remote_update\.check_for_updates\(\)",
    }

    for label, pattern in markers.items():
        source_match = re.search(pattern, source_text)
        _assert(source_match is not None, f"source missing {label}")

    for remote_contract in (
        "def _fetch_sdk_check_presets(force_remote=False):",
        "refresh_requested = request.args.get(\"refresh\"",
        "force_remote=refresh_requested",
        "const refreshQuery = force ? '&refresh=1'",
    ):
        _assert(remote_contract in source_text, f"canonical SDK preset refresh is missing: {remote_contract}")


def test_macos_architecture_contract() -> None:
    spec = (ROOT / "EventInspector.spec").read_text(encoding="utf-8", errors="ignore")
    workflow = (ROOT / ".github" / "workflows" / "macos-build.yml").read_text(encoding="utf-8", errors="ignore")
    build_script = (ROOT / "build" / "macos" / "build_macos.sh").read_text(encoding="utf-8", errors="ignore")
    _assert("target_arch='arm64'" not in spec, "macOS spec must not force Apple Silicon")
    _assert("MACOS_TARGET_ARCH" in spec, "macOS spec must accept an explicit target architecture")
    _assert("arch: arm64" in workflow and "arch: x86_64" in workflow, "macOS workflow must publish both architectures")
    _assert(
        "runner: macos-14" in workflow and "runner: macos-15-intel" in workflow,
        "macOS workflow must use native Apple Silicon and Intel runners",
    )
    _assert('x86_64) MACOS_TARGET_ARCH="x86_64"' in build_script, "macOS build script lost Intel host support")


def test_update_candidate_does_not_downgrade() -> None:
    candidates = [
        {"update_dir": "/tmp/build52", "build": 52, "source": "older"},
        {"update_dir": "/tmp/build53", "build": 53, "source": "current"},
        {"update_dir": "/tmp/build56", "build": 56, "source": "release"},
        {"update_dir": "/tmp/build54", "build": 54, "source": "older-release"},
    ]
    selected = desktop._select_prepared_update_candidate(candidates, bundled_build=53)
    _assert_equal(selected["build"], 56, "bundled build 53 must select the newest prepared update")
    _assert_equal(
        desktop._select_prepared_update_candidate(candidates[:2], bundled_build=53),
        None,
        "a prepared payload at or below the bundled build must be ignored",
    )
    _assert_equal(
        desktop._select_prepared_update_candidate(candidates[:1], bundled_build=None)["build"],
        52,
        "a client without a detected bundled build must keep update compatibility",
    )
    equal_build_candidate = [{"update_dir": "/tmp/build53-cache", "build": 53, "source": "same_build_cache"}]
    _assert_equal(
        desktop._select_prepared_update_candidate(equal_build_candidate, bundled_build=53),
        None,
        "a same-build prepared payload must not override the bundled source",
    )
    compatibility_candidate = [{"update_dir": "/tmp/build59", "build": CURRENT_RELEASE_BUILD, "source": "channel_state"}]
    _assert_equal(
        desktop._select_prepared_update_candidate(compatibility_candidate, bundled_build=53)["build"],
        CURRENT_RELEASE_BUILD,
        "a newer prepared payload must be accepted over the bundled build",
    )


def test_services_checker_gradle_mapping_contract() -> None:
    service_source = (ROOT / "services_checker" / "app.py").read_text(encoding="utf-8")
    gradle_mapping = json.loads((ROOT / "services_checker" / "gradle_lib_mapping.json").read_text(encoding="utf-8"))
    gradle_presets = json.loads((ROOT / "services_checker" / "gradle_check_presets.json").read_text(encoding="utf-8"))
    podfile_presets = json.loads((ROOT / "services_checker" / "podfile_check_presets.json").read_text(encoding="utf-8"))
    podfile_mapping = json.loads((ROOT / "services_checker" / "podfile_lib_mapping.json").read_text(encoding="utf-8"))
    c192_gradle_lines = gradle_presets.get("C-192-Android", {}).get("lines") or []
    _assert(c192_gradle_lines, "C-192 Android Gradle preset is missing")
    expected_c191_max_lines = {
        "Mobilefuse - MAX\t1.12.0.0",
        "Ascendx - MAX\t1.12.0",
        "Yeahmobi/ Maticoo - MAX\t2.0.7.1",
        "TaurusX - MAX\t1.16.3.1",
    }
    _assert(
        expected_c191_max_lines.issubset(set(c192_gradle_lines)),
        "C-192 MAX Gradle preset entries are incomplete",
    )
    expected_c191_max_mapping = {
        "Mobilefuse - MAX": "com.applovin.mediation:mobilefuse-adapter",
        "Ascendx - MAX": "com.knorex:ascendx-mobile-sdk-max-custom-adapter",
        "Yeahmobi/ Maticoo - MAX": "io.github.maticooads:maticoo-adapter-max",
        "TaurusX - MAX": "com.applovin.mediation:taurusXAdapters",
    }
    for library_name, artifact_id in expected_c191_max_mapping.items():
        _assert_equal(
            gradle_mapping.get(library_name),
            artifact_id,
            f"C-192 MAX Gradle mapping changed for {library_name}",
        )
    _assert_equal(
        podfile_presets.get("C-180-iOS", {}).get("lines") or [],
        [
            "Kidoz Adapter\t2.2.0",
            "Kidoz SDK\t10.1.5",
            "Yeahmobi/ Maticoo Adapter\t2.2.0",
            "Yeahmobi/ Maticoo SDK\t2.2.0",
            "TaurusX SDK\t1.18.1",
            "TaurusX Adapter\t1.18.1",
            "Odeeo SDK\t3.10.0",
            "AppMetrica Analytics\t6.4.0",
            "Google UMP SDK\t3.1.0",
            "Ascendx adapter\t0.9.0",
            "Voodoo Adapter\t5.2.0.0",
            "Adjust/AdjustGoogleOd\t5.6.2",
        ],
        "C-180 iOS Podfile preset changed",
    )
    _assert_equal(
        podfile_mapping.get("TaurusX Adapter"),
        "TaurusxAdsSDK/IronSourceAdapter",
        "Podfile TaurusX mapping changed",
    )
    _assert_equal(
        podfile_mapping.get("Adjust/AdjustGoogleOd"),
        "Adjust/AdjustGoogleOdm",
        "Podfile Adjust mapping changed",
    )
    _assert("gradle_lib_mapping.json" in service_source, "Services Checker Gradle mapping file is missing")
    _assert("podfile_lib_mapping.json" in service_source, "Services Checker Podfile mapping file is missing")
    _assert("def _load_gradle_lib_mapping(" in service_source, "Services Checker Gradle mapping loader is missing")
    _assert("def _load_podfile_lib_mapping(" in service_source, "Services Checker Podfile mapping loader is missing")
    _assert("GRADLE_LIB_MAPPING =" not in service_source, "Gradle mapping must not remain hardcoded in app.py")
    _assert("PODFILE_LIB_MAPPING =" not in service_source, "Podfile mapping must not remain hardcoded in app.py")
    _assert_equal(
        service_source.count("def scan_gradle_for_versions("),
        1,
        "Services Checker must have one authoritative Gradle scanner",
    )
    _assert("import json" in service_source.split("# --- Flask Routes ---", 1)[0], "Services Checker preset JSON import is missing")
    _assert_equal(
        service_source.count("def _load_build_check_presets("),
        1,
        "Services Checker preset loader must be defined in the active source block",
    )
    for loader_name in ("_load_gradle_check_presets", "_load_podfile_check_presets"):
        _assert_equal(service_source.count(f"def {loader_name}("), 1, f"Services Checker loader missing: {loader_name}")
    for preset_file in (
        "apk_check_presets.json",
        "gradle_check_presets.json",
        "podfile_check_presets.json",
        "manifest_check_presets.json",
    ):
        _assert(preset_file in service_source, f"Services Checker preset file is missing: {preset_file}")
    for remote_contract in (
        "import requests",
        "SERVICES_CHECKER_PRESET_BRANCHES",
        "SERVICES_CHECKER_GITHUB_API_URL",
        "def _fetch_services_checker_revision(",
        "SERVICES_CHECKER_COMMIT_SHA_PATTERN",
        "SERVICES_CHECKER_PRESET_FILENAMES",
        "SERVICES_CHECKER_REMOTE_DATA_FILENAMES",
        "def _remote_preset_urls(",
        "requests.get(",
        "response.raise_for_status()",
        "eventinspector_refresh={cache_bust}",
        '"Accept-Encoding": "identity"',
        "def _refresh_remote_preset_files(",
        "_remote_preset_refresh_lock = threading.Lock()",
        "with _remote_preset_refresh_lock:",
        "explicit_refresh = request.args.get(\"refresh\"",
        "refresh_requested = True",
        "forceRemote = false",
        "refreshQuery = forceRemote ? '&refresh=1'",
        "await loadBuildCheckPresets(true)",
        "function presetMatchesControl(control, preset)",
        "function syncBuildCheckPreset(control, select, clearManual = false)",
        "function restoreSelectedBuildCheckPreset(selectId)",
        "area.replaceChildren(messageDiv)",
        "let buildCheckPresetRequestSerial = 0;",
        "const requestSerial = ++buildCheckPresetRequestSerial;",
        "if (requestSerial !== buildCheckPresetRequestSerial) return;",
        "syncBuildCheckPreset(control, select)",
        "select.onchange = () => syncBuildCheckPreset(control, select, true)",
        "if (!response.ok) throw new Error('preset_request_failed:' + response.status)",
    ):
        _assert(remote_contract in service_source, f"Live Services Checker preset refresh is missing: {remote_contract}")
    _assert("'gradle_presets': gradle_presets" in service_source, "Gradle preset API group is missing")
    _assert("'podfile_presets': podfile_presets" in service_source, "Podfile preset API group is missing")
    required_mappings = {
        "Adjust Meta Referrer": "com.adjust.sdk:adjust-android-meta-referrer",
        "Adjust Samsung Referrer": "com.adjust.sdk:adjust-android-samsung-referrer",
        "Adjust Vivo Referrer": "com.adjust.sdk:adjust-android-vivo-referrer",
        "Adjust Xiaomi Referrer": "com.adjust.sdk:adjust-android-xiaomi-referrer",
        "Xiaomi Install Referrer": "com.miui.referrer:homereferrer",
        "Samsung Install Referrer": "store.galaxy.samsung.installreferrer:samsung_galaxystore_install_referrer",
    }
    for mapping_name, mapping_artifact in required_mappings.items():
        _assert_equal(
            gradle_mapping.get(mapping_name),
            mapping_artifact,
            f"Services Checker mapping changed or missing: {mapping_name}",
        )
    _assert("LINE Ads adapter" not in gradle_mapping, "LINE Ads mapping must remain removed")
    for control_id in (
        'id="apk-build-check-preset"',
        'id="gradle-build-check-preset"',
        'id="podfile-build-check-preset"',
        'id="manifest-check-preset"',
        'id="reload-apk-build-check-presets"',
        'id="reload-gradle-build-check-presets"',
        'id="reload-podfile-build-check-presets"',
        'id="reload-manifest-check-presets"',
    ):
        _assert(service_source.count(control_id) == 1, f"Services Checker preset control must exist once: {control_id}")
    _assert('platform: \'android\'' in service_source, "Android build preset filtering is missing")
    _assert('platform: \'ios\'' in service_source, "iOS build preset filtering is missing")
    _assert('presetMatchesControl(control, preset)' in service_source, "build preset platform filtering is missing")
    _assert('restoreSelectedBuildCheckPreset(\'apk-build-check-preset\')' in service_source, "APK preset must survive tab reset")
    _assert('restoreSelectedBuildCheckPreset(\'gradle-build-check-preset\')' in service_source, "Gradle preset must survive tab reset")
    _assert('restoreSelectedBuildCheckPreset(\'podfile-build-check-preset\')' in service_source, "Podfile preset must survive tab reset")
    _assert('SERVICES_CHECKER_PRESET_BRANCH = "main"' in service_source, "main must be the canonical preset branch")
    _assert(
        'SERVICES_CHECKER_PRESET_BRANCHES = (SERVICES_CHECKER_PRESET_BRANCH,)' in service_source,
        "preset refresh must not fall back to a stale release branch",
    )
    _assert("'refreshed_sources': refreshed_sources" in service_source, "preset reload must report the GitHub branch used")
    _assert('id="build-check-preset"' not in service_source, "obsolete global build preset selector must be removed")


def test_services_checker_live_preset_refresh_after_restart() -> None:
    """Catch the stale-cache regression with two fresh service imports."""
    import requests as service_requests

    service_path = ROOT / "services_checker" / "app.py"
    original_get = service_requests.get
    original_cache_dir = os.environ.get("EVENTINSPECTOR_PRESET_CACHE_DIR")
    revision = {"value": "old"}
    calls = []

    class FakeResponse:
        def __init__(self, payload):
            self.content = json.dumps(payload, separators=(",", ":")).encode("utf-8")

        def raise_for_status(self):
            return None

        def json(self):
            return json.loads(self.content.decode("utf-8"))

    def fake_get(url, **_kwargs):
        calls.append(url)
        if "api.github.com/repos/" in url:
            return FakeResponse({"sha": "a" * 40})
        _assert(
            "/services_checker/" in url
            and (
                "/main/services_checker/" in url
                or "/" + ("a" * 40) + "/services_checker/" in url
            ),
            f"preset refresh used a non-main source: {url}",
        )
        filename = url.split("/services_checker/", 1)[1].split("?", 1)[0]
        if filename == "gradle_lib_mapping.json":
            return FakeResponse({"Harness Library": f"com.example:harness-{revision['value']}"})
        if filename == "podfile_lib_mapping.json":
            return FakeResponse({"Harness Pod": f"PodHarness{revision['value']}"})
        payload = {
            "C-191-Android": {
                "platform": "ios" if filename == "podfile_check_presets.json" else "android",
                "lines": [f"{filename}:{revision['value']}"],
            }
        }
        return FakeResponse(payload)

    module_names = []
    try:
        with tempfile.TemporaryDirectory(prefix="eventinspector_services_presets_") as cache_dir:
            os.environ["EVENTINSPECTOR_PRESET_CACHE_DIR"] = cache_dir
            service_requests.get = fake_get

            def load_service_module(name):
                spec = importlib.util.spec_from_file_location(name, service_path)
                _assert(spec is not None and spec.loader is not None, "Services Checker module could not be loaded")
                module = importlib.util.module_from_spec(spec)
                sys.modules[name] = module
                module_names.append(name)
                spec.loader.exec_module(module)
                return module

            first = load_service_module("eventinspector_services_checker_refresh_one")
            first_response = first.app.test_client().get("/api/build-check-presets?ts=1")
            _assert_equal(first_response.status_code, 200, "first Services Checker preset request failed")
            first_data = first_response.get_json()
            _assert_equal(
                first_data["manifest_presets"]["C-191-Android"]["lines"],
                ["manifest_check_presets.json:old"],
                "first Service Checker start did not load the remote preset",
            )
            _assert_equal(
                first._load_gradle_lib_mapping().get("Harness Library"),
                "com.example:harness-old",
                "first Services Checker start did not load the remote Gradle mapping",
            )
            _assert_equal(
                first._load_podfile_lib_mapping().get("Harness Pod"),
                "PodHarnessold",
                "first Services Checker start did not load the remote Podfile mapping",
            )

            revision["value"] = "new"
            second = load_service_module("eventinspector_services_checker_refresh_two")
            second_response = second.app.test_client().get("/api/build-check-presets?ts=2")
            _assert_equal(second_response.status_code, 200, "restart Services Checker preset request failed")
            second_data = second_response.get_json()
            _assert_equal(
                second_data["manifest_presets"]["C-191-Android"]["lines"],
                ["manifest_check_presets.json:new"],
                "Services Checker restart kept the stale per-user preset cache",
            )
            _assert_equal(
                second._load_gradle_lib_mapping().get("Harness Library"),
                "com.example:harness-new",
                "Services Checker restart kept the stale Gradle mapping cache",
            )
            _assert_equal(
                second._load_podfile_lib_mapping().get("Harness Pod"),
                "PodHarnessnew",
                "Services Checker restart kept the stale Podfile mapping cache",
            )
            _assert_equal(
                set(second_data["refreshed_files"]),
                {
                    "apk_check_presets.json",
                    "gradle_check_presets.json",
                    "gradle_lib_mapping.json",
                    "podfile_check_presets.json",
                    "podfile_lib_mapping.json",
                    "manifest_check_presets.json",
                },
                "restart did not refresh all Services Checker preset files",
            )
            preset_calls = [url for url in calls if "/services_checker/" in url]
            _assert(
                preset_calls
                and all("/" + ("a" * 40) + "/services_checker/" in url for url in preset_calls),
                "preset source was not immutable main commit",
            )
    finally:
        service_requests.get = original_get
        for name in module_names:
            sys.modules.pop(name, None)
        if original_cache_dir is None:
            os.environ.pop("EVENTINSPECTOR_PRESET_CACHE_DIR", None)
        else:
            os.environ["EVENTINSPECTOR_PRESET_CACHE_DIR"] = original_cache_dir


def test_services_checker_git_value_reload_from_real_commit() -> None:
    """Require a real Git commit change to reach the same client process.

    A mocked requests.get can prove only that code paths exist. This test
    creates a temporary Git repository, serves its checked-in files over HTTP,
    changes a committed preset value, and verifies that the already imported
    Services Checker returns the new value after Reload.
    """
    service_path = ROOT / "services_checker" / "app.py"
    preset_filenames = (
        "apk_check_presets.json",
        "gradle_check_presets.json",
        "gradle_lib_mapping.json",
        "podfile_check_presets.json",
        "podfile_lib_mapping.json",
        "manifest_check_presets.json",
    )
    original_cache_dir = os.environ.get("EVENTINSPECTOR_PRESET_CACHE_DIR")
    module_name = "eventinspector_services_checker_real_git_reload"

    def run_git(repo, *args):
        return subprocess.run(
            ["git", *args],
            cwd=repo,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()

    class QuietHandler(http.server.SimpleHTTPRequestHandler):
        def log_message(self, _format, *_args):
            return

    server = None
    try:
        with tempfile.TemporaryDirectory(prefix="eventinspector_git_presets_") as temp_dir:
            repo = Path(temp_dir) / "remote"
            cache_dir = Path(temp_dir) / "client-cache"
            (repo / "services_checker").mkdir(parents=True)
            run_git(repo, "init", "-q", "-b", "main")
            run_git(repo, "config", "user.email", "harness@example.invalid")
            run_git(repo, "config", "user.name", "EventInspector Harness")

            def write_presets(revision):
                for filename in preset_filenames:
                    if filename == "gradle_lib_mapping.json":
                        payload = {"Harness Library": f"com.example:harness-{revision}"}
                    elif filename == "podfile_lib_mapping.json":
                        payload = {"Harness Pod": f"PodHarness{revision}"}
                    else:
                        payload = {
                            "C-191-Android": {
                                "platform": "ios" if filename == "podfile_check_presets.json" else "android",
                                "lines": [f"{filename}:{revision}"],
                                "harness_marker": revision,
                            }
                        }
                    (repo / "services_checker" / filename).write_text(
                        json.dumps(payload, sort_keys=True),
                        encoding="utf-8",
                    )

            write_presets("commit-one")
            run_git(repo, "add", "services_checker")
            run_git(repo, "commit", "-qm", "preset commit one")
            first_commit = run_git(repo, "rev-parse", "HEAD")

            server = http.server.ThreadingHTTPServer(
                ("127.0.0.1", 0),
                functools.partial(QuietHandler, directory=str(repo)),
            )
            threading.Thread(target=server.serve_forever, daemon=True).start()
            base_url = f"http://127.0.0.1:{server.server_port}/services_checker"

            os.environ["EVENTINSPECTOR_PRESET_CACHE_DIR"] = str(cache_dir)
            spec = importlib.util.spec_from_file_location(module_name, service_path)
            _assert(spec is not None and spec.loader is not None, "Services Checker module could not be loaded")
            module = importlib.util.module_from_spec(spec)
            sys.modules[module_name] = module
            spec.loader.exec_module(module)
            module._fetch_services_checker_revision = lambda: None
            module._remote_preset_urls = lambda filename, _revision=None: (
                ("git-fixture", f"{base_url}/{filename}")
                for _ in (0,)
            )

            client = module.app.test_client()
            first_response = client.get("/api/build-check-presets?refresh=1&ts=commit-one")
            _assert_equal(first_response.status_code, 200, "real Git preset request failed")
            first_data = first_response.get_json()
            _assert_equal(
                first_data["manifest_presets"]["C-191-Android"]["harness_marker"],
                "commit-one",
                "client did not receive the first committed Git value",
            )
            _assert_equal(
                module._load_gradle_lib_mapping().get("Harness Library"),
                "com.example:harness-commit-one",
                "client did not receive the first committed Gradle mapping",
            )
            _assert_equal(
                module._load_podfile_lib_mapping().get("Harness Pod"),
                "PodHarnesscommit-one",
                "client did not receive the first committed Podfile mapping",
            )
            first_digest = first_data["loaded_preset_files"]["manifest_check_presets.json"]["sha256"]
            _assert_equal(
                first_data["refreshed_sources"]["manifest_check_presets.json"],
                "git-fixture",
                "wrong Git source",
            )

            write_presets("commit-two")
            run_git(repo, "add", "services_checker")
            run_git(repo, "commit", "-qm", "preset commit two")
            second_commit = run_git(repo, "rev-parse", "HEAD")
            _assert(first_commit != second_commit, "harness Git fixture did not create a new commit")

            second_response = client.get("/api/build-check-presets?refresh=1&ts=commit-two")
            _assert_equal(second_response.status_code, 200, "real Git reload request failed")
            second_data = second_response.get_json()
            _assert_equal(
                second_data["manifest_presets"]["C-191-Android"]["harness_marker"],
                "commit-two",
                "client kept the old preset value after a real Git commit changed",
            )
            _assert_equal(
                module._load_gradle_lib_mapping().get("Harness Library"),
                "com.example:harness-commit-two",
                "client kept the old Gradle mapping after a real Git commit changed",
            )
            _assert_equal(
                module._load_podfile_lib_mapping().get("Harness Pod"),
                "PodHarnesscommit-two",
                "client kept the old Podfile mapping after a real Git commit changed",
            )
            second_digest = second_data["loaded_preset_files"]["manifest_check_presets.json"]["sha256"]
            _assert(first_digest != second_digest, "client preset hash did not change after Git commit")
            _assert_equal(
                set(second_data["refreshed_files"]),
                set(preset_filenames),
                "Reload did not download every Git-backed Service Checker preset",
            )
            _assert_equal(
                second_data["preset_revision"] != first_data["preset_revision"],
                True,
                "client revision marker did not change after Git commit",
            )
    finally:
        if server is not None:
            server.shutdown()
            server.server_close()
        sys.modules.pop(module_name, None)
        if original_cache_dir is None:
            os.environ.pop("EVENTINSPECTOR_PRESET_CACHE_DIR", None)
        else:
            os.environ["EVENTINSPECTOR_PRESET_CACHE_DIR"] = original_cache_dir


def test_sdk_preset_git_value_reload_from_real_commit() -> None:
    """Require a real Git commit change to reach the SDK preset client.

    This intentionally uses the real SDK endpoint and HTTP transport. The
    temporary repository gives the test two actual committed revisions without
    changing the production repository.
    """
    preset_filename = "sdk_check_presets.json"
    original_urls = list(lc.SDK_CHECK_PRESETS_REMOTE_URLS)

    def run_git(repo, *args):
        return subprocess.run(
            ["git", *args],
            cwd=repo,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()

    class QuietHandler(http.server.SimpleHTTPRequestHandler):
        def log_message(self, _format, *_args):
            return

    server = None
    try:
        with tempfile.TemporaryDirectory(prefix="eventinspector_sdk_git_presets_") as temp_dir:
            repo = Path(temp_dir) / "remote"
            repo.mkdir()
            run_git(repo, "init", "-q", "-b", "main")
            run_git(repo, "config", "user.email", "harness@example.invalid")
            run_git(repo, "config", "user.name", "EventInspector Harness")

            def write_preset(revision):
                payload = {
                    "C-191-Android": {
                        "platform": "android",
                        "lines": [f"Ads Network\tHarness-{revision}"],
                        "harness_marker": revision,
                    }
                }
                (repo / preset_filename).write_text(
                    json.dumps(payload, sort_keys=True),
                    encoding="utf-8",
                )

            write_preset("commit-one")
            run_git(repo, "add", preset_filename)
            run_git(repo, "commit", "-qm", "SDK preset commit one")
            first_commit = run_git(repo, "rev-parse", "HEAD")

            server = http.server.ThreadingHTTPServer(
                ("127.0.0.1", 0),
                functools.partial(QuietHandler, directory=str(repo)),
            )
            threading.Thread(target=server.serve_forever, daemon=True).start()
            lc.SDK_CHECK_PRESETS_REMOTE_URLS[:] = [
                f"http://127.0.0.1:{server.server_port}/{preset_filename}"
            ]

            client = lc.app.test_client()
            first_response = client.get("/api/sdk-check-presets?refresh=1&ts=commit-one")
            _assert_equal(first_response.status_code, 200, "real SDK Git preset request failed")
            first_data = first_response.get_json()
            _assert_equal(first_data.get("source"), "github", "SDK preset did not come from Git HTTP")
            _assert_equal(
                first_data["presets"]["C-191-Android"]["lines"],
                ["Ads Network\tHarness-commit-one"],
                "client did not receive the first committed SDK preset value",
            )

            write_preset("commit-two")
            run_git(repo, "add", preset_filename)
            run_git(repo, "commit", "-qm", "SDK preset commit two")
            second_commit = run_git(repo, "rev-parse", "HEAD")
            _assert(first_commit != second_commit, "harness SDK Git fixture did not create a new commit")

            second_response = client.get("/api/sdk-check-presets?refresh=1&ts=commit-two")
            _assert_equal(second_response.status_code, 200, "real SDK Git reload request failed")
            second_data = second_response.get_json()
            _assert_equal(
                second_data["presets"]["C-191-Android"]["lines"],
                ["Ads Network\tHarness-commit-two"],
                "client kept the old SDK preset value after a real Git commit changed",
            )
    finally:
        if server is not None:
            server.shutdown()
            server.server_close()
        lc.SDK_CHECK_PRESETS_REMOTE_URLS[:] = original_urls


def test_canonical_manifest_build_contract() -> None:
    manifest_path = ROOT / "Updates_2_5" / "remote_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    _assert_equal(manifest.get("version"), CURRENT_RELEASE_VERSION, "canonical manifest must target the current build")
    payload_ref = str(manifest.get("payload_ref") or "").strip().lower()
    _assert(re.fullmatch(r"[0-9a-f]{40}", payload_ref) is not None, "canonical manifest payload_ref must be an immutable Git SHA")
    for item in manifest.get("files") or []:
        _assert("compat_sha256" not in item, "canonical manifest must not carry a compatibility payload hash")
        for url in [item.get("url"), *(item.get("urls") or [])]:
            _assert(
                payload_ref in str(url) and "/main/" not in str(url) and "@main/" not in str(url),
                f"canonical manifest URL must use immutable payload ref {payload_ref}: {url}",
            )
    protected_preset_paths = {
        "sdk_check_presets.json",
        "package_log_presets.json",
        "services_checker/apk_check_presets.json",
        "services_checker/gradle_check_presets.json",
        "services_checker/gradle_lib_mapping.json",
        "services_checker/podfile_check_presets.json",
        "services_checker/podfile_lib_mapping.json",
        "services_checker/manifest_check_presets.json",
    }
    manifest_paths = {str(item.get("path") or "").strip() for item in manifest.get("files") or []}
    _assert(not (manifest_paths & protected_preset_paths), "release manifest must not ship preset data")


def test_update_flow_canonical_v25() -> None:
    import remote_update as updater

    manifest_path = ROOT / "Updates_2_5" / "remote_manifest.json"
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes)
    payloads = _immutable_manifest_payloads(manifest)
    original_home = os.environ.get("HOME")
    original_bundle_build = os.environ.get("EVENTINSPECTOR_BUNDLED_BUILD")
    original_bundle_source = os.environ.get("EVENTINSPECTOR_BUNDLED_BUILD_SOURCE")
    original_download_first = updater._download_first
    original_download_verified = updater._download_verified
    try:
        with tempfile.TemporaryDirectory(prefix="eventinspector_harness_") as temp_home:
            os.environ["HOME"] = temp_home
            os.environ["EVENTINSPECTOR_BUNDLED_BUILD"] = "1"
            os.environ["EVENTINSPECTOR_BUNDLED_BUILD_SOURCE"] = "detected"

            def fake_download_first(_urls, _timeout):
                return manifest_bytes, "harness://remote_manifest.json"

            def fake_download_verified(urls, _timeout, _expected_sha256=""):
                rel_path = next((path for path in payloads if any(url.endswith("/" + path) for url in urls)), None)
                _assert(rel_path is not None, f"unexpected update payload URL list: {urls}")
                return payloads[rel_path], f"harness://{rel_path}"

            updater._download_first = fake_download_first
            updater._download_verified = fake_download_verified

            first = updater.check_for_updates(force_refresh=True)
            _assert_equal(first.get("status"), "updated", "v2.5 client must prepare the release payload")
            _assert_equal(first.get("version"), CURRENT_RELEASE_VERSION, "prepared v2.5 payload version mismatch")
            prepared = updater.get_prepared_update_info()
            _assert_equal(prepared.get("build"), CURRENT_RELEASE_BUILD, "prepared v2.5 payload build mismatch")
            _assert(os.path.exists(os.path.join(prepared["update_dir"], "Log_checker.py")), "prepared Log_checker.py missing")
            _assert(
                os.path.exists(os.path.join(prepared["update_dir"], "services_checker", "bundletool-all-1.18.1.jar")),
                "prepared bundletool asset missing from the remote update payload",
            )
            _assert(
                os.path.exists(os.path.join(prepared["update_dir"], "services_checker", "axml_fallback.py")),
                "prepared dependency-free AXML fallback payload missing",
            )
            for preset_path in (
                "sdk_check_presets.json",
                "package_log_presets.json",
                "services_checker/apk_check_presets.json",
                "services_checker/gradle_check_presets.json",
                "services_checker/gradle_lib_mapping.json",
                "services_checker/podfile_check_presets.json",
                "services_checker/podfile_lib_mapping.json",
                "services_checker/manifest_check_presets.json",
            ):
                _assert(
                    not os.path.exists(os.path.join(prepared["update_dir"], *preset_path.split("/"))),
                    f"preset data must stay outside the prepared app update: {preset_path}",
                )

            selected = desktop._select_prepared_update_candidate(
                [{
                    "update_dir": prepared["update_dir"],
                    "build": prepared["build"],
                    "source": "harness-prepared-release",
                }],
                bundled_build=50,
            )
            _assert(selected is not None, f"bundle build 50 must accept the prepared build {CURRENT_RELEASE_BUILD} payload")
            _assert_equal(selected.get("build"), CURRENT_RELEASE_BUILD, f"prepared build {CURRENT_RELEASE_BUILD} was not selected")
            updated_source = Path(prepared["update_dir"], "Log_checker.py").read_text(
                encoding="utf-8", errors="ignore"
            )
            _assert("Default Ad Events" in updated_source, "prepared payload does not contain Default Ad Events")
            _assert(
                f"v2.5.0({CURRENT_RELEASE_BUILD})" in updated_source,
                "prepared payload carries the wrong UI build marker",
            )

            second = updater.check_for_updates()
            _assert_equal(second.get("status"), "up_to_date", "same v2.5 payload must not download repeatedly")

    finally:
        updater._download_first = original_download_first
        updater._download_verified = original_download_verified
        if original_home is None:
            os.environ.pop("HOME", None)
        else:
            os.environ["HOME"] = original_home
        if original_bundle_build is None:
            os.environ.pop("EVENTINSPECTOR_BUNDLED_BUILD", None)
        else:
            os.environ["EVENTINSPECTOR_BUNDLED_BUILD"] = original_bundle_build
        if original_bundle_source is None:
            os.environ.pop("EVENTINSPECTOR_BUNDLED_BUILD_SOURCE", None)
        else:
            os.environ["EVENTINSPECTOR_BUNDLED_BUILD_SOURCE"] = original_bundle_source


def test_windows_portable_update_staging() -> None:
    """Windows must stage a new payload beside, not over, the active payload."""
    import remote_update as updater

    manifest_path = ROOT / "Updates_2_5" / "remote_manifest.json"
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes)
    payloads = _immutable_manifest_payloads(manifest)
    original_os_name = updater.os.name
    original_user_data_dir = updater._user_data_dir
    original_download_first = updater._download_first
    original_download_verified = updater._download_verified
    try:
        with tempfile.TemporaryDirectory(prefix="eventinspector_windows_harness_") as temp_support:
            updater.os.name = "nt"
            updater._user_data_dir = lambda: temp_support

            active_dir = os.path.join(temp_support, "updates_v250")
            os.makedirs(active_dir, exist_ok=True)
            with open(os.path.join(active_dir, "active_payload.locked"), "w", encoding="utf-8") as handle:
                handle.write("keep")
            updater._save_state({
                "last_check": 0,
                "version": "2026-08-27-1-2.5.0-52",
                "update_dir": str(active_dir),
                "files": ["active_payload.locked"],
            })

            def fake_download_first(_urls, _timeout):
                return manifest_bytes, "harness://remote_manifest.json"

            def fake_download_verified(urls, _timeout, _expected_sha256=""):
                rel_path = next((path for path in payloads if any(url.endswith("/" + path) for url in urls)), None)
                _assert(rel_path is not None, f"unexpected update payload URL list: {urls}")
                return payloads[rel_path], f"harness://{rel_path}"

            updater._download_first = fake_download_first
            updater._download_verified = fake_download_verified

            result = updater.check_for_updates()
            _assert_equal(result.get("status"), "updated", "Windows portable update must complete without replacing the active directory")
            prepared = updater.get_prepared_update_info()
            prepared_dir = prepared["update_dir"]
            _assert(prepared_dir != active_dir, "Windows update must use a new payload directory")
            _assert(os.path.basename(prepared_dir).startswith("updates_v250_"), "Windows staged payload name must identify the v2.5 channel")
            _assert(os.path.exists(os.path.join(active_dir, "active_payload.locked")), "active Windows payload was modified during update")
            _assert(os.path.exists(os.path.join(prepared_dir, "Log_checker.py")), "Windows staged Log_checker.py is missing")
            _assert_equal(prepared.get("build"), CURRENT_RELEASE_BUILD, "Windows staged payload build mismatch")
            selected = desktop._select_prepared_update_candidate(
                [{
                    "update_dir": prepared_dir,
                    "build": prepared.get("build"),
                    "source": "windows-channel-state",
                }],
                bundled_build=53,
            )
            _assert(selected is not None, "Windows restart must see the staged build as a usable update")
            _assert_equal(selected.get("update_dir"), prepared_dir, "Windows restart selected the wrong payload directory")

            second = updater.check_for_updates()
            _assert_equal(second.get("status"), "up_to_date", "Windows staged payload must not download repeatedly")
    finally:
        updater.os.name = original_os_name
        updater._user_data_dir = original_user_data_dir
        updater._download_first = original_download_first
        updater._download_verified = original_download_verified


def test_build_scripts_clean_outputs() -> None:
    attributes = (ROOT / ".gitattributes").read_text(encoding="utf-8", errors="ignore")
    _assert("*.py text eol=lf" in attributes, "Git checkout must keep Python payloads on LF")
    _assert("*.json text eol=lf" in attributes, "Git checkout must keep JSON payloads on LF")
    _assert("*.jar binary" in attributes, "Git checkout must keep bundletool binary")
    _assert("*.keystore binary" in attributes, "Git checkout must keep keystore binary")
    mac_script = (ROOT / "build" / "macos" / "build_macos.sh").read_text(encoding="utf-8", errors="ignore")
    win_portable_script = (ROOT / "build" / "windows" / "build_portable.bat").read_text(encoding="utf-8", errors="ignore")
    win_installer_script = (ROOT / "build" / "windows" / "build_windows.bat").read_text(encoding="utf-8", errors="ignore")
    _assert("verify_update_manifest.py" in mac_script, "macOS build must verify update manifest hashes")
    _assert("verify_update_manifest.py" in win_portable_script, "Windows portable build must verify update manifest hashes")
    _assert("verify_update_manifest.py" in win_installer_script, "Windows installer build must verify update manifest hashes")
    _assert("python -m venv .venv" in win_portable_script, "Windows portable build must use the configured Python interpreter")
    _assert("python -m venv .venv" in win_installer_script, "Windows installer build must use the configured Python interpreter")
    _assert("python -m PyInstaller" in win_portable_script, "Windows portable build must invoke PyInstaller from the active environment")
    _assert("python -m PyInstaller" in win_installer_script, "Windows installer build must invoke PyInstaller from the active environment")
    _assert("python -m pip install -r requirements.txt" in win_portable_script, "Windows portable build must install dependencies into the active environment")
    _assert("python -m pip install -r requirements.txt" in win_installer_script, "Windows installer build must install dependencies into the active environment")

    mac_expected = [
        'rm -rf "dist/EventInspector.app"',
        'rm -f "dist/EventInspector.dmg"',
        'rm -rf "build/EventInspector"',
    ]
    for needle in mac_expected:
        _assert(needle in mac_script, f"build_macos.sh must clean stale artifact: {needle}")

    _assert('if [ -z "${MACOS_TARGET_ARCH:-}" ]; then' in mac_script, "macOS build must resolve a target architecture")
    _assert('arm64) MACOS_TARGET_ARCH="arm64"' in mac_script, "macOS build must support Apple Silicon")
    _assert('x86_64) MACOS_TARGET_ARCH="x86_64"' in mac_script, "macOS build must support Intel")
    _assert('--target-arch "$MACOS_TARGET_ARCH"' in mac_script, "macOS build must pass the target architecture")
    _assert('--exclude-module "markupsafe._speedups"' in mac_script, "macOS universal build must avoid the arm64-only MarkupSafe speedup")
    _assert('--add-data "Log_checker.py:."' in mac_script, "macOS build must package the bundled release marker")
    _assert('--add-data "sdk_check_presets.json:."' not in mac_script, "macOS build must keep SDK presets outside the app bundle")
    _assert('--add-data "package_log_presets.json:."' not in mac_script, "macOS build must keep Package Log presets outside the app bundle")
    _assert('--collect-submodules "androguard.core"' in mac_script, "macOS build must package manifest parser submodules")
    _assert('--collect-data "androguard"' in mac_script, "macOS build must package manifest parser data")
    _assert('--hidden-import "androguard.core.axml"' in mac_script, "macOS build must include AXMLPrinter explicitly")
    for service_asset in (
        "services_checker/app.py",
        "services_checker/axml_fallback.py",
        "services_checker/bundletool-all-1.18.1.jar",
        "services_checker/my-key.keystore",
    ):
        _assert(service_asset in mac_script, f"macOS build must package {service_asset}")
    _assert('--add-data "sdk_check_presets.json;."' not in win_portable_script, "Windows portable build must keep SDK presets outside the app bundle")
    _assert('--add-data "sdk_check_presets.json;."' not in win_installer_script, "Windows installer build must keep SDK presets outside the app bundle")
    _assert('--add-data "package_log_presets.json;."' not in win_portable_script, "Windows portable build must keep Package Log presets outside the app bundle")
    _assert('--add-data "package_log_presets.json;."' not in win_installer_script, "Windows installer build must keep Package Log presets outside the app bundle")
    _assert('--add-data "Log_checker.py;."' in win_portable_script, "Windows portable build must package the release source")
    _assert('--add-data "Log_checker.py;."' in win_installer_script, "Windows installer build must package the release source")
    _assert('--add-data "remote_update.py;."' in win_portable_script, "Windows portable build must package the updater source")
    _assert('--add-data "remote_update.py;."' in win_installer_script, "Windows installer build must package the updater source")
    _assert('verify_release_source.ps1' in win_portable_script, "Windows portable build must validate the release source")
    _assert('verify_release_source.ps1' in win_installer_script, "Windows installer build must validate the release source")
    _assert('verify_bundle.ps1' in win_portable_script, "Windows portable build must validate the built bundle")
    _assert('verify_bundle.ps1' in win_installer_script, "Windows installer build must validate the built bundle")
    _assert('--collect-submodules "androguard.core"' in win_portable_script, "Windows portable build must package manifest parser submodules")
    _assert('--collect-submodules "androguard.core"' in win_installer_script, "Windows installer build must package manifest parser submodules")
    _assert('--collect-data "androguard"' in win_portable_script, "Windows portable build must package manifest parser data")
    _assert('--collect-data "androguard"' in win_installer_script, "Windows installer build must package manifest parser data")
    _assert('--hidden-import "androguard.core.axml"' in win_portable_script, "Windows portable build must include AXMLPrinter explicitly")
    _assert('--hidden-import "androguard.core.axml"' in win_installer_script, "Windows installer build must include AXMLPrinter explicitly")
    for service_asset in (
        "services_checker\\app.py",
        "services_checker\\axml_fallback.py",
        "services_checker\\bundletool-all-1.18.1.jar",
        "services_checker\\my-key.keystore",
    ):
        _assert(service_asset in win_portable_script, f"Windows portable build must package {service_asset}")
        _assert(service_asset in win_installer_script, f"Windows installer build must package {service_asset}")
    spec_text = (ROOT / "EventInspector.spec").read_text(encoding="utf-8", errors="ignore")
    _assert("('sdk_check_presets.json', '.')" not in spec_text, "PyInstaller spec must keep SDK presets outside the app bundle")
    _assert("('package_log_presets.json', '.')" not in spec_text, "PyInstaller spec must keep Package Log presets outside the app bundle")
    _assert("('remote_update.py', '.')" in spec_text, "PyInstaller spec must package the updater source")
    _assert("collect_data_files('androguard')" in spec_text, "PyInstaller spec must collect manifest parser data")
    _assert("collect_submodules('androguard.core')" in spec_text, "PyInstaller spec must include manifest parser submodules")
    _assert("'androguard.core.axml'" in spec_text, "PyInstaller spec must include AXMLPrinter explicitly")
    _assert("('services_checker', 'services_checker')" not in spec_text, "PyInstaller spec must not package runtime upload files")
    for service_asset in (
        "services_checker/app.py",
        "services_checker/axml_fallback.py",
        "services_checker/bundletool-all-1.18.1.jar",
        "services_checker/my-key.keystore",
    ):
        _assert(service_asset in spec_text, f"PyInstaller spec must package {service_asset}")

    win_expected = [
        'if exist "dist\\EventInspector" rmdir /s /q "dist\\EventInspector"',
        'if exist "build\\EventInspector" rmdir /s /q "build\\EventInspector"',
    ]
    for needle in win_expected:
        _assert(needle in win_portable_script, f"build_portable.bat must clean stale artifact: {needle}")
        _assert(needle in win_installer_script, f"build_windows.bat must clean stale artifact: {needle}")
    _assert(
        'if exist "build\\windows\\Output" rmdir /s /q "build\\windows\\Output"' in win_installer_script,
        "build_windows.bat must clean stale installer output",
    )


def test_windows_release_build_version_contract() -> None:
    workflow = (ROOT / ".github" / "workflows" / "windows-build.yml").read_text(encoding="utf-8", errors="ignore")
    source_guard = (ROOT / "build" / "windows" / "verify_release_source.ps1").read_text(encoding="utf-8", errors="ignore")
    bundle_guard = (ROOT / "build" / "windows" / "verify_bundle.ps1").read_text(encoding="utf-8", errors="ignore")
    installer_script = (ROOT / "build" / "windows" / "build_windows.bat").read_text(encoding="utf-8", errors="ignore")
    iss_script = (ROOT / "build" / "windows" / "EventChecker.iss").read_text(encoding="utf-8", errors="ignore")
    smoke_source = (ROOT / "tools" / "windows_update_smoke.py").read_text(encoding="utf-8", errors="ignore")
    manifest_guard = (ROOT / "tools" / "verify_update_manifest.py").read_text(encoding="utf-8", errors="ignore")

    _assert('default: "main"' in workflow, "Windows workflow must default to the main source ref")
    _assert(
        "ref: ${{ github.event_name == 'workflow_dispatch' && inputs.source_ref || github.sha }}" in workflow
        or "ref: ${{ inputs.source_ref || 'main' }}" in workflow
        or "ref: ${{ inputs.source_ref || github.sha }}" in workflow,
        "Windows workflow must checkout the requested source ref",
    )
    _assert("Verify Portable ZIP version" in workflow, "Windows workflow must verify the portable ZIP")
    _assert("Verify Windows portable updater" in workflow, "Windows workflow must run the portable updater smoke test")
    _assert("python tools\\windows_update_smoke.py" in workflow, "Windows workflow must execute the Windows updater smoke test")
    _assert("sys.path.insert(0, str(ROOT))" in smoke_source, "Windows updater smoke must import the root updater from CI")
    _assert("EVENTINSPECTOR_WINDOWS_SMOKE_SIMULATE" in smoke_source, "Windows updater smoke must support local code-path testing")
    _assert("sha256" in manifest_guard and "update manifest: PASS" in manifest_guard, "release manifest hash guard is incomplete")
    _assert("verify_update_manifest.py" in smoke_source, "Windows updater smoke must verify manifest hashes")
    _assert("EventInspector-Windows-v${{ steps.release.outputs.release_version }}" in workflow, "Windows artifacts must be versioned")
    _assert('[switch]$PrintVersion' in source_guard, "Windows source guard must expose the resolved version")
    _assert('ExpectedSeries = "2.5.0"' in source_guard, "Windows source guard must enforce the v2.5 series")
    _assert("$match.Groups['build'].Value" in source_guard, "Windows source guard must derive the build number from source")
    _assert("bundleMarker -ne $sourceMarker" in bundle_guard, "Windows bundle guard must reject stale bundled source")
    _assert('Resolve-BundleFile $BundleRoot "remote_update.py"' in bundle_guard, "Windows bundle guard must verify the updater source")
    _assert("EventInspector.exe" in bundle_guard, "Windows bundle guard must require the executable")
    _assert("-PrintVersion" in installer_script, "Windows installer must derive its version from the source")
    _assert("/DMyAppVersion=%EVENTINSPECTOR_RELEASE_VERSION%" in installer_script, "Inno Setup must receive the source version")
    _assert(
        f'#define MyAppVersion "2.5.0.{RELEASE_SOURCE_BUILD}"' in iss_script,
        "Inno Setup fallback must match the current v2.5 release",
    )


def test_windows_update_http_smoke() -> None:
    """Run the same real-HTTP updater smoke used by the Windows release workflow."""
    smoke_path = ROOT / "tools" / "windows_update_smoke.py"
    environment = os.environ.copy()
    if os.name != "nt":
        environment["EVENTINSPECTOR_WINDOWS_SMOKE_SIMULATE"] = "1"
    completed = subprocess.run(
        [sys.executable, str(smoke_path)],
        cwd=str(ROOT),
        env=environment,
        capture_output=True,
        text=True,
        timeout=120,
    )
    output = (completed.stdout or "") + (completed.stderr or "")
    _assert_equal(completed.returncode, 0, f"Windows updater HTTP smoke failed:\n{output}")
    _assert("windows_update_smoke: PASS" in output, "Windows updater HTTP smoke did not report PASS")


def test_windows_update_recovery_script() -> None:
    script_path = ROOT / "tools" / "reset_update_state_windows.bat"
    text = script_path.read_text(encoding="utf-8", errors="ignore")
    _assert(
        f"TARGET_VERSION={CURRENT_RELEASE_VERSION}" in text,
        "windows recovery script must target the current release",
    )
    _assert('updates_%%C' in text and 'v250' in text, "windows recovery script must clear every update channel")
    _assert('Updates_2_5/remote_manifest.json' in text, "windows recovery script must target the v2.5 manifest")
    _assert('services_checker/bundletool-all-1.18.1.jar' in text, "windows recovery script must preserve the Services Checker payload")
    stale_scripts = sorted((ROOT / "tools").glob("bootstrap_windows_to_v*.bat"))
    _assert(not stale_scripts, f"remove stale Windows bootstrap scripts: {[p.name for p in stale_scripts]}")


def test_services_checker_bridge_contract() -> None:
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8")
    for payload_source in (source,):
        _assert("EVENTINSPECTOR_ALLOW_EXTERNAL_SERVICES_CHECKER" in payload_source, "bundled Services Checker source switch missing")
        _assert("_services_checker_external_sources_enabled" in payload_source, "bundled Services Checker source policy missing")
        _assert("EVENTINSPECTOR_SERVICES_COMMAND" in payload_source, "Services Checker command override missing")
        _assert("AndroidTool.command" in payload_source, "Drive launcher fallback missing")
        _assert("Androidchecker.cmd" in payload_source, "Windows Drive launcher fallback missing")
        _assert("SERVICES_CHECKER_APP_RELATIVE_PATH" in payload_source, "shared Drive app path resolver missing")
        _assert("_services_checker_uses_shared_keystore" in payload_source, "Drive keystore source check missing")
        _assert("force_drive_import" in payload_source, "stale Services Checker source must be bypassed")
        _assert("subprocess.Popen" in payload_source, "Services Checker launcher missing")
        _assert("_services_checker_ready" in payload_source, "Services Checker readiness check missing")
        _assert("_prepare_services_checker_runtime" in payload_source, "Services Checker bundle import preparation missing")
        _assert("_services_checker_resource_roots" in payload_source, "Services Checker resource layout resolver missing")
        _assert("_services_checker_bundled_resource_root" in payload_source, "Services Checker complete resource contract missing")
        _assert("EVENTINSPECTOR_BUNDLETOOL_PATH" in payload_source, "Services Checker bundletool path pin missing")
        _assert("SERVICES_CHECKER_REQUIRED_RESOURCES" in payload_source, "Services Checker required resource list missing")
        _assert('importlib.import_module("androguard.core.axml")' in payload_source, "Services Checker androguard preload missing")
        _assert("_services_checker_saved_host_path" in payload_source, "saved host override missing")
        _assert("/api/services-checker/host" in payload_source, "host replacement API missing")
        _assert("/api/services-checker/reload" in payload_source, "Services Checker reload API missing")
        _assert("servicesCheckerReloadBtn" not in payload_source, "obsolete Services Checker reload button must stay removed")
        _assert('id="servicesCheckerStatus"' in payload_source, "Services Checker status UI missing")
        _assert('id="servicesCheckerError"' in payload_source, "Services Checker error UI missing")
        _assert("Source: bundled with Event Inspector" not in payload_source, "obsolete Services Checker source label must stay removed")
        _assert("servicesCheckerReplaceHostBtn" not in payload_source, "host replacement UI must not be exposed in bundled mode")
        _assert('"services_checker", "app.py"' in payload_source, "bundled Services Checker fallback missing")

        file_candidates_start = payload_source.index("def _services_checker_file_candidates")
        command_candidates_start = payload_source.index("def _services_checker_command_candidates")
        file_candidates_source = payload_source[file_candidates_start:command_candidates_start]
        bundled_marker = 'candidates.append(os.path.join(SCRIPT_DIR, "services_checker", "app.py"))'
        external_marker = "if _services_checker_external_sources_enabled():"
        _assert(bundled_marker in file_candidates_source, "bundled Services Checker candidate missing")
        _assert(external_marker in file_candidates_source, "external Services Checker fallback must be opt-in")
        _assert(
            file_candidates_source.index(bundled_marker) < file_candidates_source.index(external_marker),
            "bundled Services Checker must be preferred over external sources",
        )

    windows_launcher = Path.home() / "Downloads" / "Androidchecker.cmd"
    if windows_launcher.exists():
        launcher_source = windows_launcher.read_text(encoding="utf-8", errors="ignore")
        _assert("Shared drives\\IndieZ - Tester" in launcher_source, "Windows launcher must resolve the shared Drive app")
        _assert("SERVICE_APP_PATH" in launcher_source, "Windows launcher override missing")

    service_source = (ROOT / "services_checker" / "app.py").read_text(encoding="utf-8")
    fallback_source = (ROOT / "services_checker" / "axml_fallback.py").read_text(encoding="utf-8")
    _assert("def _load_axml_printer()" in service_source, "Services Checker AXML retry loader is missing")
    _assert("def _load_fallback_axml_printer()" in service_source, "Services Checker dependency-free fallback loader is missing")
    _assert('importlib.import_module("androguard.core.axml")' in service_source, "Services Checker AXML import is not explicit")
    _assert("_live_remote_preset_payloads" in service_source, "Services Checker must keep the latest downloaded payload in memory")
    _assert("refreshed_digests" in service_source, "Services Checker must expose downloaded preset hashes")
    _assert("loaded_preset_files" in service_source, "Services Checker must expose the preset copy used by the response")
    _assert("preset_revision" in service_source, "Services Checker must expose a client-visible preset revision")
    _assert("class AXMLPrinter" in fallback_source, "Services Checker dependency-free AXML fallback is missing")
    _assert((ROOT / "services_checker" / "axml_fallback.py").is_file(), "Services Checker fallback file is missing")
    _assert(".container h2.text-xl" in service_source, "Services Checker typography override is missing")
    _assert(".results-section h3, .results-section h4 { font-size: 0.78rem" in service_source, "Services Checker result headings are too large")
    _assert(
        ".results-section p, .results-section li, .results-section .comparison-item { font-size: 0.68rem"
        in service_source,
        "Services Checker comparison rows are too large",
    )
    _assert("value.join('\\\\n')" in service_source, "bundled Services Checker newline escape is invalid")
    key_path = ROOT / "services_checker" / "my-key.keystore"
    _assert(key_path.is_file() and key_path.stat().st_size > 0, "bundled test keystore is missing")


def test_services_checker_resource_contract_is_platform_neutral() -> None:
    source = (ROOT / "Log_checker.py").read_text(encoding="utf-8")
    helper_start = source.index("def _services_checker_resource_roots")
    helper_end = source.index("def _prepare_services_checker_runtime")
    helper_source = source[helper_start:helper_end].lower()
    for forbidden in ("sys.platform", "platform.system", "os.name", '"darwin"', '"windows"'):
        _assert(forbidden not in helper_source, f"resource resolver must not infer platform from {forbidden}")

    previous_root = os.environ.get("EVENTINSPECTOR_BUNDLED_SERVICES_ROOT")
    previous_bundletool = os.environ.get("EVENTINSPECTOR_BUNDLETOOL_PATH")
    try:
        for layout in (
            Path("mac-artifact") / "Contents" / "Resources" / "services_checker",
            Path("windows-artifact") / "_internal" / "services_checker",
        ):
            with tempfile.TemporaryDirectory(prefix="eventchecker-resource-contract-") as temp_dir:
                resource_root = Path(temp_dir) / layout
                resource_root.mkdir(parents=True)
                (resource_root / "bundletool-all-1.18.1.jar").write_bytes(b"bundletool")
                (resource_root / "my-key.keystore").write_bytes(b"keystore")
                app_path = resource_root / "app.py"
                app_path.write_text("# test app\n", encoding="utf-8")
                os.environ["EVENTINSPECTOR_BUNDLED_SERVICES_ROOT"] = str(resource_root)
                os.environ.pop("EVENTINSPECTOR_BUNDLETOOL_PATH", None)

                resolved_root = lc._services_checker_bundled_resource_root(str(app_path))
                _assert_equal(
                    Path(resolved_root).resolve(),
                    resource_root.resolve(),
                    f"resource resolver must use the declared {layout} layout",
                )
                lc._prepare_services_checker_runtime(str(app_path))
                _assert_equal(
                    Path(os.environ["EVENTINSPECTOR_BUNDLED_SERVICES_ROOT"]).resolve(),
                    resource_root.resolve(),
                    "runtime must pin the complete resource root",
                )
                _assert_equal(
                    Path(os.environ["EVENTINSPECTOR_BUNDLETOOL_PATH"]).resolve(),
                    (resource_root / "bundletool-all-1.18.1.jar").resolve(),
                    "runtime must pin bundletool inside the same resource root",
                )
    finally:
        if previous_root is None:
            os.environ.pop("EVENTINSPECTOR_BUNDLED_SERVICES_ROOT", None)
        else:
            os.environ["EVENTINSPECTOR_BUNDLED_SERVICES_ROOT"] = previous_root
        if previous_bundletool is None:
            os.environ.pop("EVENTINSPECTOR_BUNDLETOOL_PATH", None)
        else:
            os.environ["EVENTINSPECTOR_BUNDLETOOL_PATH"] = previous_bundletool


def test_native_download_contract() -> None:
    source = (ROOT / "desktop_app.py").read_text(encoding="utf-8")
    configure_marker = "webview.settings['ALLOW_DOWNLOADS'] = True"
    _assert(configure_marker in source, "native WebView downloads must be enabled")
    configure_index = source.index("_configure_webview_downloads()")
    window_index = source.index("webview.create_window(")
    _assert(configure_index < window_index, "download setting must be applied before creating the WebView")

    service_path = ROOT / "services_checker" / "app.py"
    service_source = service_path.read_text(encoding="utf-8", errors="ignore")
    _assert("def _resolve_bundletool_path()" in service_source, "AAB converter bundletool resolver is missing")
    _assert("def _resolve_keystore_path()" in service_source, "AAB converter keystore resolver is missing")
    _assert(
        "bundletool_path = _resolve_bundletool_path()" in service_source,
        "AAB conversion must resolve bundletool at request time",
    )
    _assert(
        "final_ks_path = _resolve_keystore_path()" in service_source,
        "AAB conversion must resolve the keystore at request time",
    )
    _assert(
        "BUNDLETOOL_PATH" not in service_source[service_source.index("def convert_aab_to_apk"):service_source.index("@app.route('/download/<filename>')")],
        "AAB conversion route must not use the import-time bundletool path",
    )
    _assert(
        "def _cleanup_stale_conversion_artifacts()" in service_source,
        "AAB conversion must clean abandoned generated artefacts",
    )
    _assert(
        "_cleanup_stale_conversion_artifacts()" in service_source[service_source.index("def convert_aab_to_apk"):service_source.index("@app.route('/download/<filename>')")],
        "AAB conversion must clean abandoned artefacts before bundletool runs",
    )
    _assert(
        "_remove_generated_file(source_path)" in service_source,
        "Native APK save must remove the generated source after copying",
    )
    _assert("@app.route('/download/<filename>')" in service_source, "Services Checker download route is missing")
    _assert("as_attachment=True" in service_source, "Services Checker download route must return an attachment")
    _assert(
        "@app.route('/save_download/<filename>', methods=['POST'])" in service_source,
        "Services Checker native save route is missing",
    )
    _assert(
        "def _copy_generated_file_to_downloads" in service_source,
        "AAB conversion must provide a platform-native Downloads save helper",
    )
    conversion_start = service_source.index("def convert_aab_to_apk()")
    download_route_start = service_source.index("@app.route('/download/<filename>')")
    conversion_source = service_source[conversion_start:download_route_start]
    _assert(
        "def _write_zip_member_to_downloads" in service_source,
        "AAB conversion must have a direct Downloads writer",
    )
    _assert(
        "download_apk_path, saved_filename = _write_zip_member_to_downloads(" in conversion_source,
        "AAB conversion must write universal.apk directly to Downloads",
    )
    _assert(
        "'saved_to_downloads': True" in conversion_source,
        "AAB conversion response must confirm the direct Downloads save",
    )
    _assert(
        "session.setdefault('downloadable_files', {})[final_apk_filename]" not in conversion_source,
        "Final APK must not be staged in the Flask session download cache",
    )
    _assert(
        "final_apk_savelocation" not in conversion_source,
        "Final APK must not be staged in UPLOAD_FOLDER",
    )
    _assert(
        "'apk_download_url'" not in conversion_source,
        "AAB conversion must not fall back to a temporary browser download URL",
    )
    _assert(
        "const saved_to_downloads" not in service_source,
        "AAB conversion UI must require the direct Downloads result",
    )
    _assert(
        "APK saved directly to Downloads:" in service_source,
        "AAB conversion UI must show the direct Downloads result",
    )


def test_gradle_comparison_hides_unlisted_preset_rows() -> None:
    service_source = (ROOT / "services_checker" / "app.py").read_text(encoding="utf-8", errors="ignore")
    _assert(
        'if (archiveType !== "Gradle") {' in service_source,
        "Gradle comparison must hide mapped libraries outside the selected preset",
    )
    _assert(
        "INFO (Found in ' + archiveType + ', not in expected list)" in service_source,
        "non-Gradle comparisons must retain informational unlisted-library rows",
    )
    _assert(
        "gradle_mapping = _load_gradle_lib_mapping()" in service_source
        and "found_versions = scan_gradle_for_versions(gradle_content, gradle_mapping)" in service_source,
        "Gradle scanner must continue scanning the complete loaded mapping",
    )


TESTS: List[Callable[[], None]] = [
    test_manifest_contract,
    test_manifest_payload_integrity,
    test_canonical_update_channel_contract,
    test_package_code_mapping,
    test_package_log_presets_contract,
    test_installation_id_state_machine,
    test_installation_id_log_parsing,
    test_sdk_exact_contracts,
    test_cloudx_sdk_adapter_metadata,
    test_max_sdk_logs,
    test_ios_max_sdk_search_only,
    test_audiomob_open_measurement_sdk_log,
    test_sdk_preset_switch_does_not_race_stop_start,
    test_clear_all_resets_sdk_check_state,
    test_platform_reconnect_contract,
    test_platform_bootstrap_sdk_state_contract,
    test_device_filter_reset_contract,
    test_platform_switch_drops_stale_device_status_contract,
    test_ios_reader_singleton_contract,
    test_ios_transport_fallback_contract,
    test_ios_discovery_grace_contract,
    test_ios_reader_liveness_contract,
    test_android_logcat_reconnect_cursor_contract,
    test_package_consumer_replacement_guard,
    test_clear_does_not_reopen_replay_dedup_contract,
    test_ios_package_frame_buffer_has_memory_guard,
    test_android_tracking_stream_reassembles_wrapped_records,
    test_ios_tracking_stream_reassembles_repeated_header_wrap,
    test_package_log_delta_sends_only_new_rows,
    test_package_log_reset_and_sync_protocol,
    test_package_log_raw_and_db_seq_contract,
    test_package_log_db_writer_survives_write_errors,
    test_package_log_older_rows_and_db_status,
    test_package_log_stuck_sessions_and_row_counts,
    test_package_log_sessions_api_pages,
    test_package_log_retention_policy,
    test_package_log_compact_api,
    test_specific_validator_delta_matches_full_rebuild,
    test_specific_validator_filters_clear_and_snapshot,
    test_specific_validator_delta_ui_contract,
    test_validator_delta_matches_server_table,
    test_validator_delta_ui_contract,
    test_revenue_delta_matches_full_rebuild,
    test_live_tables_delta_snapshot_and_clear,
    test_live_tables_ui_contract,
    test_ui_zoom_setting_contract,
    test_specific_filter_boxes_compact,
    test_full_window_layout,
    test_sdk_check_results_text_size,
    test_repository_moved_to_indiezvn,
    test_package_log_delta_ui_contract,
    test_main_script_has_unique_top_level_names,
    test_android_package_log_ignores_global_pause,
    test_ios_package_log_preserves_multiline_records,
    test_ios_package_log_tidevice_nul_framing,
    test_ios_package_log_tidevice_newline_stream_does_not_merge_kernel,
    test_ios_max_load_ads_from_package_frame_contract,
    test_sdk_base_name_matching,
    test_sdk_check_preset_contract,
    test_rendered_sdk_preset_javascript_contract,
    test_default_ad_event_contract,
    test_ios_default_ad_event_tracking_record,
    test_exact_duplicate_record_dedup_across_ios_tabs,
    test_installation_id_copy_contract,
    test_sdk_failed_groups_sort_first,
    test_release_build_marker,
    test_bright_skad_ipa_inspection_contract,
    test_preset_revert_protection_contract,
    test_rewarded_bidding_filter_contract,
    test_package_quick_filters_api,
    test_package_history_split_by_platform,
    test_package_log_switch_project_keeps_session,
    test_package_quick_filter_history_query,
    test_package_quick_filter_upgrade_backup_export_import,
    test_price_rotation_exact_parser,
    test_load_ads_provider_contract,
    test_load_ads_max_contract,
    test_ios_load_ads_max_viewability_contract,
    test_ios_load_ads_max_viewability_record_boundary_contract,
    test_ios_load_ads_max_delegate_contract,
    test_levelplay_impression_data_callback_contract,
    test_ascendx_cloudx_callback_contract,
    test_release_payload_sync,
    test_macos_architecture_contract,
    test_update_candidate_does_not_downgrade,
    test_services_checker_gradle_mapping_contract,
    test_services_checker_live_preset_refresh_after_restart,
    test_services_checker_git_value_reload_from_real_commit,
    test_sdk_preset_git_value_reload_from_real_commit,
    test_canonical_manifest_build_contract,
    test_update_flow_canonical_v25,
    test_windows_portable_update_staging,
    test_build_scripts_clean_outputs,
    test_windows_release_build_version_contract,
    test_windows_update_http_smoke,
    test_windows_update_recovery_script,
    test_services_checker_bridge_contract,
    test_services_checker_resource_contract_is_platform_neutral,
    test_native_download_contract,
    test_gradle_comparison_hides_unlisted_preset_rows,
]


RELEASE_ONLY_TESTS = {
    # These checks intentionally compare the current worktree with the
    # immutable release manifest.  They belong to the release gate, not to a
    # normal local logic run while a new change is still uncommitted.
    "test_release_payload_sync",
    "test_windows_update_http_smoke",
}


def run_tests(selected: Iterable[Callable[[], None]]) -> List[HarnessResult]:
    results: List[HarnessResult] = []
    for test in selected:
        try:
            test()
            results.append(HarnessResult(test.__name__, True, "ok"))
        except Exception as exc:  # pragma: no cover - harness should surface the failure directly
            results.append(HarnessResult(test.__name__, False, str(exc)))
    return results


def main() -> int:
    parser = argparse.ArgumentParser(description="EventInspector local harness")
    parser.add_argument("--json", action="store_true", help="Emit JSON results")
    parser.add_argument("--only", nargs="*", default=[], help="Run only named tests")
    parser.add_argument(
        "--profile",
        choices=("local", "release"),
        default="local",
        help="local skips worktree-vs-release gates; release runs every test (default: local)",
    )
    args = parser.parse_args()

    selected = TESTS
    if args.only:
        wanted = set(args.only)
        selected = [test for test in TESTS if test.__name__ in wanted]
        missing = wanted - {test.__name__ for test in selected}
        if missing:
            print(f"Unknown test name(s): {', '.join(sorted(missing))}", file=sys.stderr)
            return 2
    elif args.profile == "local":
        selected = [test for test in TESTS if test.__name__ not in RELEASE_ONLY_TESTS]

    results = run_tests(selected)
    failures = [result for result in results if not result.passed]

    if args.json:
        print(json.dumps([result.__dict__ for result in results], indent=2, ensure_ascii=False))
    else:
        for result in results:
            status = "PASS" if result.passed else "FAIL"
            print(f"{status} {result.name}: {result.message}")

    if failures:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
